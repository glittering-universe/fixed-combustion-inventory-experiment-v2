#!/usr/bin/env python3
"""Read-only reconciliation of exported formula workbooks against engine snapshots."""

from __future__ import annotations

import json
import math
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


POLLUTANTS = ("SO2", "NOX", "CO", "VOCS", "PM10", "PM25", "BC", "OC")


def close_enough(actual, expected, tolerance=1e-8):
    if actual in (None, "") and expected in (None, ""):
        return True
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return math.isclose(float(actual), float(expected), rel_tol=tolerance, abs_tol=tolerance)
    return actual == expected


def header_map(sheet):
    return {cell.value: index for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1)}


def verify(workbook_path: Path, payload_path: Path):
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    values_book = load_workbook(workbook_path, read_only=True, data_only=True)
    formula_book = load_workbook(workbook_path, read_only=True, data_only=False)
    mismatches = []
    error_values = []

    calculation_sheet = values_book["03_燃料污染物计算"]
    calculation_headers = header_map(calculation_sheet)
    formula_rows = {}
    source_aggregates = defaultdict(lambda: defaultdict(lambda: {"withheld": False, "generation": 0.0, "emission": 0.0}))
    calculation_checks = 0
    for row_number, row in enumerate(calculation_sheet.iter_rows(min_row=2), start=2):
        source_id = row[calculation_headers["源ID"] - 1].value
        raw_amount = row[calculation_headers["原始消耗量"] - 1].value
        raw_unit = row[calculation_headers["原始单位"] - 1].value
        expected_activity = raw_amount * 1000 if raw_unit == "吨" and raw_amount is not None else raw_amount * 10000 if raw_unit == "万立方米" and raw_amount is not None else None
        actual_activity = row[calculation_headers["标准活动水平(公式)"] - 1].value
        calculation_checks += 1
        if not close_enough(actual_activity, expected_activity):
            mismatches.append({"sheet": calculation_sheet.title, "row": row_number, "field": "activity", "actual": actual_activity, "expected": expected_activity})
        formula_rows[row_number] = source_id
        for pollutant in POLLUTANTS:
            action = row[calculation_headers[f"{pollutant}_准入动作"] - 1].value
            comparisons = (
                (f"{pollutant}_有效因子(公式)", f"{pollutant}_因子快照"),
                (f"{pollutant}_产生量(公式,t)", f"{pollutant}_产生量快照(t)"),
                (f"{pollutant}_排放量(公式,t)", f"{pollutant}_排放量快照(t)"),
            )
            for formula_field, snapshot_field in comparisons:
                actual = row[calculation_headers[formula_field] - 1].value
                expected = row[calculation_headers[snapshot_field] - 1].value
                calculation_checks += 1
                if not close_enough(actual, expected):
                    mismatches.append({"sheet": calculation_sheet.title, "row": row_number, "field": formula_field, "actual": actual, "expected": expected})
            aggregate = source_aggregates[source_id][pollutant]
            if action == "withhold":
                aggregate["withheld"] = True
            generation = row[calculation_headers[f"{pollutant}_产生量(公式,t)"] - 1].value
            emission = row[calculation_headers[f"{pollutant}_排放量(公式,t)"] - 1].value
            if isinstance(generation, (int, float)):
                aggregate["generation"] += float(generation)
            if isinstance(emission, (int, float)):
                aggregate["emission"] += float(emission)

    nh3_sheet = values_book["04_氨逃逸"]
    nh3_headers = header_map(nh3_sheet)
    nh3_by_source = {}
    nh3_checks = 0
    positive_nh3 = 0
    for row_number, row in enumerate(nh3_sheet.iter_rows(min_row=2), start=2):
        source_id = row[nh3_headers["源ID"] - 1].value
        fields = (
            ("煤活动水平(公式,kg)", "煤活动水平快照(kg)"),
            ("有效因子(公式,g/kg煤)", "因子快照(g/kg煤)"),
            ("产生量(公式,t)", "产生量快照(t)"),
            ("排放量(公式,t)", "排放量快照(t)"),
        )
        for formula_field, snapshot_field in fields:
            actual = row[nh3_headers[formula_field] - 1].value
            expected = row[nh3_headers[snapshot_field] - 1].value
            nh3_checks += 1
            if not close_enough(actual, expected):
                mismatches.append({"sheet": nh3_sheet.title, "row": row_number, "field": formula_field, "actual": actual, "expected": expected})
        action = row[nh3_headers["准入动作"] - 1].value
        generation = row[nh3_headers["产生量(公式,t)"] - 1].value
        emission = row[nh3_headers["排放量(公式,t)"] - 1].value
        nh3_by_source[source_id] = {"withheld": action == "withhold", "generation": generation, "emission": emission}
        if isinstance(generation, (int, float)) and generation > 0:
            positive_nh3 += 1

    result_sheet = values_book["05_源结果"]
    result_headers = header_map(result_sheet)
    result_checks = 0
    for row_number, row in enumerate(result_sheet.iter_rows(min_row=2), start=2):
        source_id = row[result_headers["源ID"] - 1].value
        for pollutant in POLLUTANTS + ("NH3",):
            if pollutant == "NH3":
                aggregate = nh3_by_source[source_id]
            else:
                aggregate = source_aggregates[source_id][pollutant]
            expected_generation = None if aggregate["withheld"] else aggregate["generation"]
            expected_emission = None if aggregate["withheld"] else aggregate["emission"]
            actual_generation = row[result_headers[f"{pollutant}_产生量汇总(t)"] - 1].value
            actual_emission = row[result_headers[f"{pollutant}_排放量汇总(t)"] - 1].value
            for field, actual, expected in (
                (f"{pollutant}_产生量汇总(t)", actual_generation, expected_generation),
                (f"{pollutant}_排放量汇总(t)", actual_emission, expected_emission),
            ):
                result_checks += 1
                if not close_enough(actual, expected):
                    mismatches.append({"sheet": result_sheet.title, "row": row_number, "field": field, "actual": actual, "expected": expected})

    formula_count = 0
    formula_texts = []
    for sheet in formula_book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "f":
                    formula_count += 1
                    formula_texts.append(str(value))
                elif isinstance(value, str) and value.startswith("#"):
                    error_values.append({"sheet": sheet.title, "cell": cell.coordinate, "value": value})
    data_error_count = 0
    for sheet in values_book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                    data_error_count += 1
                    if len(error_values) < 30:
                        error_values.append({"sheet": sheet.title, "cell": cell.coordinate, "value": cell.value})
    old_logic_hits = sum("/5" in formula or "/2" in formula for formula in formula_texts)
    cems_formula_hits = sum("CEMS" in formula.upper() for formula in formula_texts)
    external_links = len(getattr(formula_book, "_external_links", []))
    values_book.close()
    formula_book.close()
    return {
        "workbook": str(workbook_path),
        "payload": str(payload_path),
        "source_count": len(payload["sources"]),
        "fuel_row_count": len(payload["calculations"]),
        "calculation_item_count": payload["calculation_item_count"],
        "calculation_checks": calculation_checks,
        "nh3_checks": nh3_checks,
        "result_checks": result_checks,
        "mismatch_count": len(mismatches),
        "mismatch_examples": mismatches[:30],
        "formula_count": formula_count,
        "formula_error_count": data_error_count,
        "formula_error_examples": error_values[:30],
        "external_link_count": external_links,
        "old_arbitrary_division_formula_hits": old_logic_hits,
        "cems_formula_hits": cems_formula_hits,
        "positive_nh3_source_count": positive_nh3,
    }


def main():
    if len(sys.argv) < 4 or (len(sys.argv) - 2) % 2:
        raise SystemExit("usage: verify_exported_workbooks.py output.json workbook payload [workbook payload ...]")
    output = Path(sys.argv[1])
    reports = []
    for index in range(2, len(sys.argv), 2):
        reports.append(verify(Path(sys.argv[index]), Path(sys.argv[index + 1])))
    result = {"success": all(report["mismatch_count"] == 0 and report["formula_error_count"] == 0 for report in reports), "reports": reports}
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))
    if not result["success"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
