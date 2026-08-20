#!/usr/bin/env python3
"""Prepare three privacy-minimized, raw-structure environmental base tables."""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import subprocess
from dataclasses import dataclass
from math import ceil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from source_identity import identity_from_base102_row, stable_pseudonym


ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parent
DEFAULT_SOURCE_ROOT = WORKSPACE / "01 环境统计数据"
DEFAULT_OUTPUT = ROOT / "inputs" / "prepared"
WRITER = ROOT / "experiment_control" / "write_deidentified_inputs.mjs"


@dataclass(frozen=True)
class RawSource:
    source_tag: str
    role: str
    filename: str
    output_name: str
    sheet: str = "Sheet1"
    shard_count: int = 1


SOURCES = (
    RawSource(
        "B102-2002",
        "device_fuel_base",
        "2022年基表查询工业企业锅炉_燃气轮机污染物和温室气体排放及治理情况(基102表)2023102002.xlsx",
        "基102-2002_原始结构脱敏输入.xlsx",
    ),
    RawSource(
        "B101-2001",
        "enterprise_master",
        "2022年基表查询工业企业污染物和温室气体排放及治理情况(基101表)2023102001.xlsx",
        "基101-2001_原始结构脱敏输入.xlsx",
        "Sheet1",
        4,
    ),
    RawSource(
        "B101-2002",
        "control_facility_base",
        "2022年基表查询工业企业污染物和温室气体排放及治理情况(基101表)2023102002.xlsx",
        "基101-2002_原始结构脱敏输入.xlsx",
    ),
)


PSEUDONYM_HEADERS = {
    "组织机构代码": "ORG",
    "统一社会信用代码": "CREDIT",
    "填报单位详细名称": "ENT",
    "曾用名": "FORMER",
}
REMOVE_TOKENS = (
    "联系人",
    "联系电话",
    "电话号码",
    "手机",
    "通讯地址",
    "详细地址乡(镇)",
    "详细地址乡（镇）",
    "详细地址街(村)",
    "详细地址街（村）",
    "门牌号",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_value(value: object) -> object:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def is_reported_pollutant_result(header: str) -> bool:
    """Block reported outcomes while retaining activity and control inputs."""
    return "产生量" in header or ("排放量" in header and "排放口" not in header)


def deidentify_value(header: object, value: object) -> object:
    name = "" if header is None else str(header).strip()
    if is_reported_pollutant_result(name):
        return None
    if any(token in name for token in REMOVE_TOKENS):
        return None
    if name in PSEUDONYM_HEADERS:
        return stable_pseudonym(PSEUDONYM_HEADERS[name], value)
    return json_value(value)


def output_name(source: RawSource, shard_index: int) -> str:
    if source.shard_count == 1:
        return source.output_name
    stem = Path(source.output_name).stem
    return f"{stem}_分片{shard_index:02d}-of{source.shard_count:02d}.xlsx"


def write_payload(payload_path: Path) -> dict[str, Any]:
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    environment = dict(os.environ)
    environment["NODE_OPTIONS"] = "--max-old-space-size=12288"
    completed = subprocess.run(
        ["node", str(WRITER), str(payload_path)],
        text=True,
        capture_output=True,
        check=False,
        env=environment,
    )
    if completed.returncode:
        raise RuntimeError(completed.stderr.strip() or completed.stdout.strip())
    destination = Path(payload["output_path"])
    if not destination.is_file():
        raise RuntimeError(f"artifact-tool did not create {destination}")
    result: dict[str, Any] | None = None
    for line in reversed(completed.stdout.splitlines()):
        try:
            candidate = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(candidate, dict) and candidate.get("output"):
            result = candidate
            break
    inspect_sidecar = destination.with_name(destination.name + ".inspect.ndjson")
    if inspect_sidecar.is_file():
        inspect_sidecar.unlink()
    if result is not None:
        return result
    return {
        "output": str(destination),
        "source_tag": payload["source_tag"],
        "schema_mode": payload["schema_mode"],
        "sheet": payload["sheet_name"],
        "rows": len(payload["values"]) - 1,
        "columns": len(payload["values"][0]),
    }


def prepare_raw_source(
    source: RawSource,
    source_root: Path,
    output_root: Path,
    payload_dir: Path,
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    source_path = source_root / source.filename
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    if source.sheet not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"missing source sheet {source.sheet}: {source_path}")
    sheet = workbook[source.sheet]
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = list(next(iterator))
    except StopIteration as exc:
        workbook.close()
        raise ValueError(f"empty workbook: {source_path}") from exc
    identities: list[dict[str, object]] = []
    workbook_rows = max(0, int(sheet.max_row or 1) - 1)
    rows_per_shard = max(1, ceil(workbook_rows / source.shard_count))
    shard_values: list[list[object]] = [headers]
    shard_index = 1
    workbook_entries: list[dict[str, object]] = []

    def flush_shard() -> None:
        nonlocal shard_values, shard_index
        destination = output_root / output_name(source, shard_index)
        payload = {
            "source_tag": source.source_tag,
            "role": source.role,
            "schema_mode": "raw_structure",
            "sheet_name": source.sheet,
            "output_path": str(destination),
            "values": shard_values,
            "shard_index": shard_index,
            "shard_count": source.shard_count,
        }
        payload_path = payload_dir / f"{source.source_tag.lower()}-s{shard_index:02d}_raw_deidentified.json"
        payload_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        writer_result = write_payload(payload_path)
        payload_path.unlink()
        workbook_entries.append({
            "source_tag": source.source_tag,
            "role": source.role,
            "sheet": source.sheet,
            "output_path": str(destination),
            "output_sha256": sha256(destination),
            "row_count": len(shard_values) - 1,
            "column_count": len(headers),
            "schema_mode": "raw_structure",
            "shard_index": shard_index,
            "shard_count": source.shard_count,
            "writer": "artifact-tool",
            "writer_result": writer_result,
        })
        shard_values = [headers]
        shard_index += 1

    for excel_row, raw_values in enumerate(iterator, start=2):
        padded = list(raw_values[: len(headers)]) + [None] * max(0, len(headers) - len(raw_values))
        row = {str(headers[index]).strip(): padded[index] for index in range(len(headers)) if headers[index] is not None}
        if source.source_tag == "B102-2002":
            identity = identity_from_base102_row(row)
            identities.append({
                **identity,
                "source_tag": source.source_tag,
                "source_sheet": source.sheet,
                "source_row": excel_row,
                "sequence": row.get("序号"),
                "inventory_year": row.get("统计年份"),
            })
        shard_values.append([deidentify_value(headers[index], padded[index]) for index in range(len(headers))])
        if len(shard_values) - 1 >= rows_per_shard and shard_index < source.shard_count:
            flush_shard()
    workbook.close()
    flush_shard()
    while shard_index <= source.shard_count:
        flush_shard()
    source_metadata = {
        "source_path": str(source_path.resolve()),
        "source_sha256": sha256(source_path),
        "source_rows": sum(int(entry["row_count"]) for entry in workbook_entries),
        "source_columns": len(headers),
    }
    for entry in workbook_entries:
        entry.update(source_metadata)
    return identities, workbook_entries


def prepare(source_root: Path, output_root: Path) -> dict[str, Any]:
    source_root = Path(source_root).resolve()
    output_root = Path(output_root).resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    payload_dir = output_root / "payloads"
    payload_dir.mkdir(parents=True, exist_ok=True)
    for stale in payload_dir.glob("*_raw_deidentified.json"):
        stale.unlink()
    stale_unsplit = output_root / "基101-2001_原始结构脱敏输入.xlsx"
    if stale_unsplit.is_file():
        stale_unsplit.unlink()
    workbook_manifest: list[dict[str, object]] = []
    identity_records: list[dict[str, object]] = []
    for source in SOURCES:
        identities, entries = prepare_raw_source(source, source_root, output_root, payload_dir)
        workbook_manifest.extend(entries)
        identity_records.extend(identities)
    if not any(payload_dir.iterdir()):
        payload_dir.rmdir()

    candidate_ids = [str(record["source_id"]) for record in identity_records]
    if len(candidate_ids) != len(set(candidate_ids)):
        duplicates = sorted(value for value in set(candidate_ids) if candidate_ids.count(value) > 1)
        raise ValueError(f"duplicate neutral source ids: {duplicates[:5]}")
    identity_index = {
        "version": "1.0.0",
        "identity_contract": "source_identity_v1",
        "candidate_source": "B102-2002",
        "candidate_ids": candidate_ids,
        "records": identity_records,
    }
    identity_path = output_root / "source_identity_index.json"
    identity_path.write_text(json.dumps(identity_index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = {
        "version": "2.0.0",
        "status": "frozen",
        "input_contract": "raw_base_tables_v2",
        "input_form": "three_raw_environmental_base_tables",
        "workbooks": workbook_manifest,
        "source_identity_index_path": str(identity_path),
        "source_identity_index_sha256": sha256(identity_path),
        "candidate_count": len(candidate_ids),
    }
    manifest_path = output_root / "raw_input_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=DEFAULT_SOURCE_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = prepare(args.source_root, args.output)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
