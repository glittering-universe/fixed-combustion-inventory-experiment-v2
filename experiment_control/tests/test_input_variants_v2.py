from __future__ import annotations

import json
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
CONTROL = ROOT / "experiment_control"
if str(CONTROL) not in sys.path:
    sys.path.insert(0, str(CONTROL))

import build_experiment_matrix
import generate_b2_injections
import generate_b_variants
import generate_scale_inputs
import validate_input_variants
from source_identity import pseudonymous_entity_id


class StructuralVariantTests(unittest.TestCase):
    def test_all_b1_variants_preserve_canonical_cell_values(self) -> None:
        original = [
            ["填报单位详细名称", "工业锅炉类型", "燃料一消耗量", "燃料一消耗量单位"],
            ["ENT-001", "燃气锅炉", 2.5, "万立方米"],
            ["ENT-002", "燃煤锅炉", 8.0, "吨"],
        ]
        for variant in ("S1", "S2", "S3", "S4"):
            perturbed = generate_b_variants.build_variant_values(
                "B102-2002", original, variant
            )
            self.assertEqual(
                validate_input_variants.canonical_rows("B102-2002", perturbed),
                validate_input_variants.canonical_rows("B102-2002", original),
                variant,
            )


class InjectionManifestTests(unittest.TestCase):
    def test_b2_manifest_has_frozen_20_6_4_target_neutral_design(self) -> None:
        records = []
        for index in range(20):
            records.append({
                "plan_class": "positive_injection",
                "root_cause": "ACTIVITY_MISSING",
                "source_id": f"SRC-RAW-P{index:019d}",
                "changes": [{"field": "燃料一消耗量", "before": 1, "after": None}],
            })
        for index in range(6):
            records.append({
                "plan_class": "negative_control",
                "root_cause": "NO_INJECTION_CONTROL",
                "source_id": f"SRC-RAW-N{index:019d}",
                "changes": [],
            })
        for index in range(4):
            records.append({
                "plan_class": "preexisting_problem",
                "root_cause": "PREEXISTING_INPUT_PROBLEM",
                "source_id": f"SRC-RAW-E{index:019d}",
                "changes": [],
                "existing_issue_reason": "missing combustion technology",
            })
        manifest = generate_b2_injections.build_manifest(records)
        validate_input_variants.validate_injection_manifest(manifest)
        self.assertEqual(manifest["planned_record_count"], 30)
        self.assertEqual(manifest["class_counts"], {
            "positive_injection": 20,
            "negative_control": 6,
            "preexisting_problem": 4,
        })
        self.assertNotIn("target", json.dumps(manifest, ensure_ascii=False).lower())

    def test_b2_evaluation_audit_is_balanced_without_out_of_scope_records(self) -> None:
        records = []
        class_by_id = {}
        for source_class, prefix in (("INDUSTRIAL", "I"), ("POWER", "P")):
            classes = ["positive_injection"] * 10 + ["negative_control"] * 3 + ["preexisting_problem"] * 2
            for index, plan_class in enumerate(classes):
                source_id = f"SRC-RAW-{prefix}{index:019d}"
                records.append({"source_id": source_id, "plan_class": plan_class, "root_cause": "TEST"})
                class_by_id[source_id] = source_class
        audit = generate_b2_injections.build_class_audit(records, class_by_id)
        expected = {"positive_injection": 10, "negative_control": 3, "preexisting_problem": 2}
        self.assertEqual(audit["class_counts"], {"INDUSTRIAL": expected, "POWER": expected})
        self.assertNotIn("EXCLUDE", json.dumps(audit))


class MatrixV2Tests(unittest.TestCase):
    def test_matrix_v2_keeps_136_machine_and_14_human_rows(self) -> None:
        payload = build_experiment_matrix.build_matrix()
        self.assertEqual(payload["matrix_version"], "2.0.0")
        self.assertEqual(payload["machine_runs"], 136)
        self.assertEqual(payload["human_runs"], 14)
        self.assertEqual(len(payload["rows"]), 150)
        self.assertEqual(payload["input_contract"], "raw_base_tables_v2")


class ScaleInputTests(unittest.TestCase):
    def test_scale_counts_are_nested_over_neutral_candidates(self) -> None:
        candidates = [f"SRC-RAW-{index:020d}" for index in range(1, 9)]
        scopes = generate_scale_inputs.nested_scale_ids(candidates)
        self.assertEqual({key: len(value) for key, value in scopes.items()}, {25: 2, 50: 4, 75: 6, 100: 8})
        self.assertTrue(set(scopes[25]).issubset(scopes[50]))
        self.assertTrue(set(scopes[50]).issubset(scopes[75]))
        self.assertTrue(set(scopes[75]).issubset(scopes[100]))

    def test_generated_scale_manifests_have_physical_candidate_rows_and_full_100_hashes(self) -> None:
        expected = {25: 1481, 50: 2961, 75: 4442, 100: 5922}
        for scale, count in expected.items():
            manifest = json.loads((ROOT / "inputs" / "scales" / str(scale) / "scale_input_manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["candidate_count"], count)
            self.assertEqual(
                sum(entry["row_count"] for entry in manifest["workbooks"] if entry["source_tag"] == "B102-2002"),
                count,
            )
        prepared = json.loads((ROOT / "inputs" / "prepared" / "raw_input_manifest.json").read_text(encoding="utf-8"))
        full = json.loads((ROOT / "inputs" / "scales" / "100" / "scale_input_manifest.json").read_text(encoding="utf-8"))
        prepared_hashes = {Path(entry["output_path"]).name: entry["output_sha256"] for entry in prepared["workbooks"]}
        full_hashes = {Path(entry["output_path"]).name: entry["output_sha256"] for entry in full["workbooks"]}
        self.assertEqual(full_hashes, prepared_hashes)

    def test_scaled_identity_records_point_to_their_physical_b102_rows(self) -> None:
        for scale in (25, 50, 75):
            scale_dir = ROOT / "inputs" / "scales" / str(scale)
            index = json.loads((scale_dir / "source_identity_index.json").read_text(encoding="utf-8"))
            manifest = json.loads((scale_dir / "scale_input_manifest.json").read_text(encoding="utf-8"))
            b102 = next(entry for entry in manifest["workbooks"] if entry["source_tag"] == "B102-2002")
            workbook = load_workbook(b102["output_path"], read_only=True, data_only=True)
            rows = list(workbook["Sheet1"].iter_rows(values_only=True))
            workbook.close()
            headers = list(rows[0])
            for expected_row, record in enumerate(index["records"], start=2):
                self.assertEqual(record["source_row"], expected_row)
                values = list(rows[expected_row - 1])
                row = dict(zip(headers, values))
                entity = pseudonymous_entity_id(
                    credit=row.get("统一社会信用代码"),
                    organization=row.get("组织机构代码"),
                    company=row.get("填报单位详细名称"),
                )
                self.assertEqual(record["pseudonymous_entity_id"], entity)
                self.assertEqual(str(record["sequence"]), str(row["序号"]))
                self.assertEqual(str(record["inventory_year"]), str(row["统计年份"]))


if __name__ == "__main__":
    unittest.main()
