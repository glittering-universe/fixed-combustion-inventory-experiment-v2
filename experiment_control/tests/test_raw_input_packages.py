from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]
CONTROL = ROOT / "experiment_control"
if str(CONTROL) not in sys.path:
    sys.path.insert(0, str(CONTROL))

import prepare_raw_inputs
import prepare_run
from source_identity import pseudonymous_entity_id, source_id_for_base102


BASE102 = "2022年基表查询工业企业锅炉_燃气轮机污染物和温室气体排放及治理情况(基102表)2023102002.xlsx"
BASE101_2001 = "2022年基表查询工业企业污染物和温室气体排放及治理情况(基101表)2023102001.xlsx"
BASE101_2002 = "2022年基表查询工业企业污染物和温室气体排放及治理情况(基101表)2023102002.xlsx"


def write_fixture(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class SourceIdentityContractTests(unittest.TestCase):
    def test_source_id_is_stable_and_target_neutral(self) -> None:
        entity = pseudonymous_entity_id(
            credit="91440101MA5EXAMPLE",
            organization="ORG-IGNORED",
            company="COMPANY-IGNORED",
        )
        self.assertEqual(entity, "CREDIT-0A99950FC832")
        source_id = source_id_for_base102(entity, sequence=12, inventory_year=2022)
        self.assertEqual(source_id, "SRC-RAW-9CDB7F8132B96188E7C2")
        self.assertNotIn("IND", source_id)
        self.assertNotIn("PWR", source_id)


class RawInputPreparationTests(unittest.TestCase):
    def test_geographic_levels_are_retained_but_street_address_is_removed(self) -> None:
        self.assertEqual(
            prepare_raw_inputs.deidentify_value("详细地址地区(市、州、盟)", "广州市"),
            "广州市",
        )
        self.assertEqual(
            prepare_raw_inputs.deidentify_value("详细地址县(区、市、旗)", "黄埔区"),
            "黄埔区",
        )
        self.assertIsNone(
            prepare_raw_inputs.deidentify_value("详细地址街(村)、门牌号", "某路1号")
        )

    def test_preparation_preserves_three_raw_schemas_and_writes_neutral_index(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            work = Path(temporary)
            source_root = work / "source"
            source_root.mkdir()
            output_one = work / "prepared-one"
            b102_headers = [
                "统计年份", "组织机构代码", "统一社会信用代码", "填报单位详细名称", "序号",
                "工业锅炉类型", "二氧化硫产生量（吨）", "联系人",
            ]
            write_fixture(
                source_root / BASE102,
                b102_headers,
                [[2022, "ORG-01", "91440101MA5EXAMPLE", "某企业", 12, "燃气锅炉", 88.5, "张三"]],
            )
            write_fixture(
                source_root / BASE101_2001,
                ["统计年份", "统一社会信用代码", "填报单位详细名称", "详细地址地区(市、州、盟)", "详细地址县(区、市、旗)", "详细地址街(村)、门牌号", "天然气消费量（万立方米）"],
                [[2022, "91440101MA5EXAMPLE", "某企业", "广州市", "黄埔区", "某路1号", 2.5]],
            )
            write_fixture(
                source_root / BASE101_2002,
                ["统计年份", "统一社会信用代码", "填报单位详细名称", "编号", "处理工艺名称", "去除效率（%）"],
                [[2022, "91440101MA5EXAMPLE", "某企业", "CTRL-1", "SCR", 80]],
            )

            first = prepare_raw_inputs.prepare(source_root, output_one)

            self.assertEqual(len(first["workbooks"]), 6)
            self.assertEqual(
                {entry["source_tag"] for entry in first["workbooks"]},
                {"B102-2002", "B101-2001", "B101-2002"},
            )
            for entry in first["workbooks"]:
                self.assertTrue(Path(entry["output_path"]).is_file())
                self.assertEqual(entry["output_sha256"], file_hash(Path(entry["output_path"])))

            index_one = json.loads((output_one / "source_identity_index.json").read_text(encoding="utf-8"))
            self.assertEqual(index_one["candidate_ids"], ["SRC-RAW-9CDB7F8132B96188E7C2"])
            self.assertNotIn("target", json.dumps(index_one, ensure_ascii=False).lower())

            b102_out = output_one / "基102-2002_原始结构脱敏输入.xlsx"
            workbook = load_workbook(b102_out, read_only=True, data_only=True)
            sheet = workbook["Sheet1"]
            actual_headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            actual_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
            actual_row += [None] * (len(actual_headers) - len(actual_row))
            workbook.close()
            self.assertEqual(actual_headers, b102_headers)
            self.assertEqual(actual_row[2], "CREDIT-0A99950FC832")
            self.assertTrue(str(actual_row[3]).startswith("ENT-"))
            self.assertIsNone(actual_row[6])
            self.assertIsNone(actual_row[7])

            enterprise_shards = sorted(output_one.glob("基101-2001_原始结构脱敏输入_分片*-of04.xlsx"))
            self.assertEqual(len(enterprise_shards), 4)
            self.assertEqual(
                sum(entry["row_count"] for entry in first["workbooks"] if entry["source_tag"] == "B101-2001"),
                1,
            )
            enterprise_out = enterprise_shards[0]
            workbook = load_workbook(enterprise_out, read_only=True, data_only=True)
            enterprise_row = [cell.value for cell in next(workbook["Sheet1"].iter_rows(min_row=2, max_row=2))]
            workbook.close()
            self.assertEqual(enterprise_row[3], "广州市")
            self.assertEqual(enterprise_row[4], "黄埔区")
            self.assertIsNone(enterprise_row[5])


class RunPackageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.work = Path(self.temporary.name)
        self.root = self.work / "repo"
        (self.root / "method_package").mkdir(parents=True)
        (self.root / "rules" / "frozen" / "v1.0.1").mkdir(parents=True)
        (self.root / "baseline" / "simple_deterministic").mkdir(parents=True)
        prepared = self.root / "inputs" / "prepared"
        prepared.mkdir(parents=True)
        self.prepared = prepared
        for name in (
            "基102-2002_原始结构脱敏输入.xlsx",
            "基101-2002_原始结构脱敏输入.xlsx",
        ):
            (prepared / name).write_bytes(name.encode("utf-8"))
        for shard in range(1, 5):
            name = f"基101-2001_原始结构脱敏输入_分片{shard:02d}-of04.xlsx"
            (prepared / name).write_bytes(name.encode("utf-8"))
        identity = {
            "version": "1.0.0",
            "candidate_ids": [f"SRC-RAW-{index:020d}" for index in range(1, 9)],
            "records": [],
        }
        (prepared / "source_identity_index.json").write_text(json.dumps(identity), encoding="utf-8")
        for scale, count in ((25, 2), (50, 4)):
            scale_dir = self.root / "inputs" / "scales" / str(scale)
            scale_dir.mkdir(parents=True)
            for path in prepared.glob("*.xlsx"):
                (scale_dir / path.name).write_bytes(path.read_bytes())
            subset = {**identity, "candidate_ids": identity["candidate_ids"][:count]}
            (scale_dir / "source_identity_index.json").write_text(json.dumps(subset), encoding="utf-8")
        (self.root / "method_package" / "frozen_manifest.json").write_text(
            json.dumps({"method_hashes": {"full": "full-hash"}}), encoding="utf-8"
        )
        (self.root / "method_package" / "input_adapter.json").write_text("{}", encoding="utf-8")
        (self.root / "rules" / "frozen" / "v1.0.1" / "rule_package_lock.json").write_text(
            json.dumps({"package_hash": "rule-hash"}), encoding="utf-8"
        )

        self.originals = {
            name: getattr(prepare_run, name)
            for name in (
                "ROOT", "METHOD_MANIFEST", "RULE_PATH", "RULE_LOCK", "INPUT_ADAPTER",
                "DETERMINISTIC_BASELINE",
            )
        }
        prepare_run.ROOT = self.root
        prepare_run.METHOD_MANIFEST = self.root / "method_package" / "frozen_manifest.json"
        prepare_run.RULE_PATH = self.root / "rules" / "frozen" / "v1.0.1"
        prepare_run.RULE_LOCK = prepare_run.RULE_PATH / "rule_package_lock.json"
        prepare_run.INPUT_ADAPTER = self.root / "method_package" / "input_adapter.json"
        prepare_run.DETERMINISTIC_BASELINE = self.root / "baseline" / "simple_deterministic"

    def tearDown(self) -> None:
        for name, value in self.originals.items():
            setattr(prepare_run, name, value)
        self.temporary.cleanup()

    def args(self, run_id: str, scale: str = "100") -> Namespace:
        return Namespace(
            run_id=run_id,
            method="full",
            scenario="full",
            target="INDUSTRIAL",
            scale=scale,
            variant="B0",
            export_mode="machine",
            output=self.work / run_id,
        )

    def test_a_b_d_packages_do_not_disclose_scope_or_candidate_membership(self) -> None:
        for run_id in ("A-IND-FULL-R1", "B-B0-IND-FULL", "D-IND-FULL"):
            manifest = prepare_run.prepare_run_package(self.args(run_id))
            self.assertNotIn("scope_ids", manifest)
            self.assertNotIn("candidate_ids", manifest)
            self.assertEqual(len(manifest["inputs"]), 6)
            self.assertEqual(
                {entry["tag"] for entry in manifest["inputs"]},
                {"B102-2002", "B101-2001", "B101-2002"},
            )
            for entry in manifest["inputs"]:
                destination = Path(entry["path"])
                source = self.prepared / destination.name
                self.assertNotEqual(source.stat().st_ino, destination.stat().st_ino)
                self.assertEqual(destination.stat().st_mode & 0o222, 0)

    def test_c_package_receives_only_target_neutral_nested_candidate_ids(self) -> None:
        quarter = prepare_run.prepare_run_package(self.args("C-25-IND-FULL-R1", "25"))
        half = prepare_run.prepare_run_package(self.args("C-50-IND-FULL-R1", "50"))
        self.assertNotIn("scope_ids", quarter)
        self.assertEqual(len(quarter["candidate_ids"]), 2)
        self.assertEqual(len(half["candidate_ids"]), 4)
        self.assertTrue(set(quarter["candidate_ids"]).issubset(half["candidate_ids"]))
        self.assertTrue(all(value.startswith("SRC-RAW-") for value in half["candidate_ids"]))
        self.assertNotIn("target", json.dumps(half["candidate_ids"]).lower())


class InputAdapterContractTests(unittest.TestCase):
    def test_adapter_registers_three_raw_roles_without_target_membership(self) -> None:
        adapter = json.loads((ROOT / "method_package" / "input_adapter.json").read_text(encoding="utf-8"))
        self.assertEqual(adapter["version"], "2.0.0")
        self.assertEqual(len(adapter["registered_sources"]), 6)
        self.assertEqual(
            {entry["tag"] for entry in adapter["registered_sources"]},
            {"B102-2002", "B101-2001", "B101-2002"},
        )
        master = [entry for entry in adapter["registered_sources"] if entry["tag"] == "B101-2001"]
        self.assertEqual([entry["shard_index"] for entry in master], [1, 2, 3, 4])
        self.assertEqual(adapter["source_identity"]["prefix"], "SRC-RAW-")
        self.assertEqual(adapter["source_identity"]["entity_priority"], ["credit", "organization", "company"])
        serialized = json.dumps(adapter, ensure_ascii=False).lower()
        self.assertNotIn("scope_ids", serialized)
        self.assertNotIn("scope_catalog", serialized)
        self.assertNotIn("target_membership", serialized)


if __name__ == "__main__":
    unittest.main()
