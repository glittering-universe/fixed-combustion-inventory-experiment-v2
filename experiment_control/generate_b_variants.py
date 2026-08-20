#!/usr/bin/env python3
"""Generate full-table, value-equivalent structural variants for Experiment B1."""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import shutil
from pathlib import Path

from openpyxl import load_workbook

from prepare_raw_inputs import sha256, write_payload


ROOT = Path(__file__).resolve().parents[1]
PREPARED = ROOT / "inputs" / "prepared"
OUT = ROOT / "inputs" / "experiment_b"
ADAPTER = json.loads((ROOT / "method_package" / "input_adapter.json").read_text(encoding="utf-8"))
VARIANTS = ("S1", "S2", "S3", "S4")


def form_perturb(header: object, index: int) -> object:
    if not isinstance(header, str) or not header:
        return header
    value = header.replace("（", "(").replace("）", ")")
    if index % 3 == 0:
        value = value.replace("(", "（").replace(")", "）")
    if index % 3 == 1:
        value = "  " + value + "  "
    if index % 3 == 2 and len(value) > 6:
        value = value[: len(value) // 2] + "\n" + value[len(value) // 2 :]
    return value


def alias_headers(tag: str, headers: list[object]) -> list[object]:
    reverse = {target: alias for alias, target in ADAPTER["header_aliases"].get(tag, {}).items()}
    return [reverse.get(header, header) for header in headers]


def column_order(tag: str, headers: list[object]) -> list[int]:
    groups: dict[str, list[int]] = {}
    for index, header in enumerate(headers):
        normalized = "" if header is None else "".join(str(header).split()).replace("（", "(").replace("）", ")")
        groups.setdefault(normalized, []).append(index)
    order = list(groups)
    seed = int(hashlib.sha256(f"B1-v2-column-order|{tag}".encode()).hexdigest()[:16], 16)
    random.Random(seed).shuffle(order)
    # Duplicate raw headers remain in their original relative order.  This
    # preserves occurrence semantics while the header groups as a whole move.
    return [index for header in order for index in groups[header]]


def build_variant_values(tag: str, rows: list[list[object]], variant: str) -> list[list[object]]:
    if variant not in VARIANTS:
        raise ValueError(f"unknown B1 variant: {variant}")
    values = [list(row) for row in rows]
    if not values:
        raise ValueError("variant input is empty")
    width = len(values[0])
    if any(len(row) != width for row in values):
        raise ValueError("variant input is not rectangular")
    if variant in {"S2", "S4"}:
        values[0] = alias_headers(tag, values[0])
    if variant in {"S1", "S4"}:
        values[0] = [form_perturb(header, index) for index, header in enumerate(values[0])]
    if variant in {"S3", "S4"}:
        order = column_order(tag, values[0])
        values = [[row[index] for index in order] for row in values]
    return values


def workbook_values(path: Path, sheet_name: str) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        raise ValueError(f"empty workbook: {path}")
    width = len(rows[0])
    return [row[:width] + [None] * max(0, width - len(row)) for row in rows]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--variants", nargs="+", choices=VARIANTS, default=list(VARIANTS))
    args = parser.parse_args()
    prepared_manifest = json.loads((PREPARED / "raw_input_manifest.json").read_text(encoding="utf-8"))
    for variant in args.variants:
        variant_dir = OUT / variant
        variant_dir.mkdir(parents=True, exist_ok=True)
        identity_destination = variant_dir / "source_identity_index.json"
        shutil.copy2(PREPARED / "source_identity_index.json", identity_destination)
        entries: list[dict[str, object]] = []
        for source in prepared_manifest["workbooks"]:
            input_path = Path(source["output_path"])
            values = build_variant_values(source["source_tag"], workbook_values(input_path, source["sheet"]), variant)
            destination = variant_dir / input_path.name
            payload = {
                "source_tag": source["source_tag"],
                "role": source["role"],
                "schema_mode": "equivalent_structural_perturbation",
                "sheet_name": source["sheet"],
                "output_path": str(destination),
                "values": values,
                "variant": variant,
                "cover_sheet": variant in {"S3", "S4"},
                "shard_index": source["shard_index"],
                "shard_count": source["shard_count"],
            }
            payload_path = variant_dir / f".{input_path.stem}.{variant}.payload.json"
            payload_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
            writer_result = write_payload(payload_path)
            payload_path.unlink()
            entry = {
                **{key: source[key] for key in (
                    "source_tag", "role", "sheet", "row_count", "column_count", "shard_index", "shard_count"
                )},
                "variant": variant,
                "output_path": str(destination),
                "output_sha256": sha256(destination),
                "writer": "artifact-tool",
                "writer_result": writer_result,
            }
            entries.append(entry)
        variant_manifest = {
            "version": "2.0.0",
            "variant": variant,
            "equivalence_contract": "same rows and values after registered header normalization and column-order recovery",
            "workbooks": entries,
            "source_identity_index_path": str(identity_destination),
            "source_identity_index_sha256": sha256(identity_destination),
        }
        (variant_dir / "variant_manifest.json").write_text(
            json.dumps(variant_manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    generated = []
    for variant in VARIANTS:
        manifest_path = OUT / variant / "variant_manifest.json"
        if manifest_path.is_file():
            generated.extend(json.loads(manifest_path.read_text(encoding="utf-8"))["workbooks"])
    manifest = {"version": "2.0.0", "variants": list(VARIANTS), "workbook_count": len(generated), "workbooks": generated}
    (OUT / "b1_generation_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"manifest": str(OUT / "b1_generation_manifest.json"), "workbook_count": len(generated)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
