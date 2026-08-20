#!/usr/bin/env python3
"""Validate B1 value equivalence and the frozen, target-neutral B2 design."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PREPARED = ROOT / "inputs" / "prepared"
EXPERIMENT_B = ROOT / "inputs" / "experiment_b"
ADAPTER = json.loads((ROOT / "method_package" / "input_adapter.json").read_text(encoding="utf-8"))
VARIANTS = ("S1", "S2", "S3", "S4")


def normalized_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.strip().replace("（", "(").replace("）", ")")
    return re.sub(r"\s+", "", text)


def canonical_header(tag: str, value: object) -> str:
    normalized = normalized_header(value)
    aliases = ADAPTER["header_aliases"].get(tag, {})
    for alias, target in aliases.items():
        if normalized == normalized_header(alias):
            return normalized_header(target)
    return normalized


def canonical_rows(tag: str, values: list[list[object]]) -> list[dict[str, object]]:
    if not values:
        raise ValueError("empty values")
    headers = [canonical_header(tag, value) for value in values[0]]
    totals = Counter(headers)
    seen: Counter[str] = Counter()
    keys = []
    for header in headers:
        seen[header] += 1
        keys.append(f"{header}__occurrence_{seen[header]}" if totals[header] > 1 else header)
    result = []
    for row in values[1:]:
        padded = row[: len(headers)] + [None] * max(0, len(headers) - len(row))
        result.append({keys[index]: padded[index] for index in range(len(headers))})
    return result


def workbook_values(path: Path, sheet_name: str = "Sheet1") -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    workbook.close()
    return rows


def validate_injection_manifest(manifest: dict[str, object]) -> None:
    records = list(manifest.get("records", []))
    expected = {"positive_injection": 20, "negative_control": 6, "preexisting_problem": 4}
    counts = {key: sum(record.get("plan_class") == key for record in records) for key in expected}
    if len(records) != 30 or counts != expected:
        raise ValueError(f"invalid B2 20/6/4 design: total={len(records)}, counts={counts}")
    if manifest.get("planned_record_count") != 30 or manifest.get("class_counts") != expected:
        raise ValueError("B2 manifest summary does not match record classes")
    if "target" in json.dumps(manifest, ensure_ascii=False).lower():
        raise ValueError("B2 manifest discloses target membership")
    source_ids = [str(record.get("source_id", "")) for record in records]
    if len(source_ids) != len(set(source_ids)) or not all(value.startswith("SRC-RAW-") for value in source_ids):
        raise ValueError("B2 source ids are not unique target-neutral source identities")
    for record in records:
        changes = record.get("changes") or []
        if record["plan_class"] == "positive_injection" and not changes:
            raise ValueError("positive B2 injection has no declared change")
        if record["plan_class"] != "positive_injection" and changes:
            raise ValueError("B2 control/preexisting record must not be modified")


def validate_class_audit(audit: dict[str, object]) -> None:
    expected = {
        "INDUSTRIAL": {"positive_injection": 10, "negative_control": 3, "preexisting_problem": 2},
        "POWER": {"positive_injection": 10, "negative_control": 3, "preexisting_problem": 2},
    }
    if audit.get("class_counts") != expected:
        raise ValueError(f"B2 source-class balance mismatch: {audit.get('class_counts')}")
    if any(record.get("source_class") not in {"INDUSTRIAL", "POWER"} for record in audit.get("records", [])):
        raise ValueError("B2 class audit contains an out-of-scope record")


def validate_b1() -> list[dict[str, object]]:
    base_manifest = json.loads((PREPARED / "raw_input_manifest.json").read_text(encoding="utf-8"))
    checks = []
    for variant in VARIANTS:
        variant_dir = EXPERIMENT_B / variant
        for source in base_manifest["workbooks"]:
            base_path = Path(source["output_path"])
            variant_path = variant_dir / base_path.name
            baseline = canonical_rows(source["source_tag"], workbook_values(base_path, source["sheet"]))
            observed = canonical_rows(source["source_tag"], workbook_values(variant_path, source["sheet"]))
            if observed != baseline:
                raise ValueError(f"B1 value equivalence failed: {variant_path}")
            checks.append({"variant": variant, "source_tag": source["source_tag"], "file": variant_path.name,
                           "rows": len(observed), "equivalent": True})
    return checks


def main() -> None:
    b1 = validate_b1()
    b2_manifest_path = EXPERIMENT_B / "B2" / "injection_manifest.json"
    b2 = json.loads(b2_manifest_path.read_text(encoding="utf-8"))
    validate_injection_manifest(b2)
    class_audit = json.loads((EXPERIMENT_B / "B2" / "evaluation_class_audit.json").read_text(encoding="utf-8"))
    validate_class_audit(class_audit)
    report = {"version": "2.0.0", "success": True, "b1_checks": b1, "b2_manifest": {
        "planned_record_count": b2["planned_record_count"], "class_counts": b2["class_counts"],
        "source_class_counts": class_audit["class_counts"],
    }}
    destination = EXPERIMENT_B / "input_variant_validation_report.json"
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
