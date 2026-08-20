#!/usr/bin/env python3
"""Generate physical 25/50/75/100% raw-input subsets for Experiment C."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from math import ceil
from pathlib import Path

from openpyxl import load_workbook

from prepare_raw_inputs import SOURCES, output_name, sha256, write_payload
from source_identity import pseudonymous_entity_id


ROOT = Path(__file__).resolve().parents[1]
PREPARED = ROOT / "inputs" / "prepared"
OUT = ROOT / "inputs" / "scales"
SCALES = (25, 50, 75, 100)


def nested_scale_ids(candidates: list[str]) -> dict[int, list[str]]:
    if len(candidates) != len(set(candidates)):
        raise ValueError("duplicate neutral candidate ids")
    ordered = sorted(
        candidates,
        key=lambda value: hashlib.sha256(f"fixed-combustion-neutral-scale-v2|{value}".encode()).hexdigest(),
    )
    return {scale: ordered[: ceil(len(ordered) * scale / 100)] for scale in SCALES}


def read_values(path: Path, sheet_name: str = "Sheet1") -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    workbook.close()
    if not values:
        raise ValueError(f"empty workbook: {path}")
    width = len(values[0])
    return [row[:width] + [None] * max(0, width - len(row)) for row in values]


def entity_for(headers: list[object], values: list[object]) -> str | None:
    row = {str(headers[index]).strip(): values[index] for index in range(len(headers)) if headers[index] is not None}
    try:
        return pseudonymous_entity_id(
            credit=row.get("统一社会信用代码"),
            organization=row.get("组织机构代码"),
            company=row.get("填报单位详细名称"),
        )
    except ValueError:
        return None


def shard_rows(rows: list[list[object]], count: int) -> list[list[list[object]]]:
    per_shard = max(1, ceil(len(rows) / count))
    return [rows[index * per_shard : (index + 1) * per_shard] for index in range(count)]


def write_artifact(
    destination: Path,
    *,
    tag: str,
    role: str,
    headers: list[object],
    rows: list[list[object]],
    shard_index: int,
    shard_count: int,
) -> dict[str, object]:
    payload = {
        "source_tag": tag,
        "role": role,
        "schema_mode": "physical_scale_subset",
        "sheet_name": "Sheet1",
        "output_path": str(destination),
        "values": [headers, *rows],
        "shard_index": shard_index,
        "shard_count": shard_count,
    }
    payload_path = destination.with_name(f".{destination.stem}.payload.json")
    payload_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    writer_result = write_payload(payload_path)
    payload_path.unlink()
    return {
        "source_tag": tag,
        "role": role,
        "sheet": "Sheet1",
        "output_path": str(destination),
        "output_sha256": sha256(destination),
        "row_count": len(rows),
        "column_count": len(headers),
        "shard_index": shard_index,
        "shard_count": shard_count,
        "writer": "artifact-tool",
        "writer_result": writer_result,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--scales", nargs="+", type=int, choices=SCALES, default=list(SCALES))
    parser.add_argument("--indexes-only", action="store_true")
    args = parser.parse_args()
    raw_manifest = json.loads((PREPARED / "raw_input_manifest.json").read_text(encoding="utf-8"))
    identity = json.loads((PREPARED / "source_identity_index.json").read_text(encoding="utf-8"))
    scales = nested_scale_ids([str(value) for value in identity["candidate_ids"]])
    identity_by_id = {str(record["source_id"]): record for record in identity["records"]}

    by_tag: dict[str, list[dict[str, object]]] = {}
    for entry in raw_manifest["workbooks"]:
        by_tag.setdefault(str(entry["source_tag"]), []).append(entry)
    source_values: dict[str, tuple[list[object], list[list[object]]]] = {}
    for tag, entries in by_tag.items():
        headers: list[object] | None = None
        rows: list[list[object]] = []
        for entry in sorted(entries, key=lambda value: int(value["shard_index"])):
            values = read_values(Path(entry["output_path"]), str(entry["sheet"]))
            if headers is None:
                headers = values[0]
            elif values[0] != headers:
                raise ValueError(f"inconsistent headers across {tag} shards")
            rows.extend(values[1:])
        assert headers is not None
        source_values[tag] = (headers, rows)

    b102_headers, b102_rows = source_values["B102-2002"]
    row_by_source_id = {
        str(record["source_id"]): b102_rows[int(record["source_row"]) - 2]
        for record in identity["records"]
    }
    source_specs = {source.source_tag: source for source in SOURCES}
    for scale in args.scales:
        selected_ids = scales[scale]
        scale_dir = OUT / str(scale)
        scale_dir.mkdir(parents=True, exist_ok=True)
        selected_records = [identity_by_id[value] for value in selected_ids]
        selected_entities = {str(record["pseudonymous_entity_id"]) for record in selected_records}
        selected_b102 = [row_by_source_id[value] for value in selected_ids]

        related: dict[str, list[list[object]]] = {"B102-2002": selected_b102}
        for tag in ("B101-2001", "B101-2002"):
            headers, rows = source_values[tag]
            related[tag] = rows if scale == 100 else [row for row in rows if entity_for(headers, row) in selected_entities]

        entries = []
        existing_manifest_path = scale_dir / "scale_input_manifest.json"
        if args.indexes_only:
            if not existing_manifest_path.is_file():
                raise FileNotFoundError(existing_manifest_path)
            entries = json.loads(existing_manifest_path.read_text(encoding="utf-8"))["workbooks"]
        elif scale == 100:
            for base_entry in raw_manifest["workbooks"]:
                source = Path(base_entry["output_path"])
                destination = scale_dir / source.name
                shutil.copy2(source, destination)
                entries.append({
                    **{key: base_entry[key] for key in (
                        "source_tag", "role", "sheet", "row_count", "column_count", "shard_index", "shard_count"
                    )},
                    "output_path": str(destination),
                    "output_sha256": sha256(destination),
                    "writer": "physical_copy_of_prepared_full_input",
                })
        else:
            for tag in ("B102-2002", "B101-2001", "B101-2002"):
                spec = source_specs[tag]
                headers, _ = source_values[tag]
                partitions = shard_rows(related[tag], spec.shard_count)
                for index, rows in enumerate(partitions, start=1):
                    destination = scale_dir / output_name(spec, index)
                    entries.append(write_artifact(
                        destination,
                        tag=tag,
                        role=spec.role,
                        headers=headers,
                        rows=rows,
                        shard_index=index,
                        shard_count=spec.shard_count,
                    ))
        index_path = scale_dir / "source_identity_index.json"
        if scale == 100:
            shutil.copy2(PREPARED / "source_identity_index.json", index_path)
        else:
            physical_records = [
                {**record, "original_source_row": record["source_row"], "source_row": physical_row}
                for physical_row, record in enumerate(selected_records, start=2)
            ]
            subset_index = {
                **{key: value for key, value in identity.items() if key not in {"candidate_ids", "records"}},
                "scale_percent": scale,
                "candidate_ids": selected_ids,
                "records": physical_records,
            }
            index_path.write_text(json.dumps(subset_index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        manifest = {
            "version": "2.0.0",
            "scale_percent": scale,
            "candidate_count": len(selected_ids),
            "candidate_entity_count": len(selected_entities),
            "row_counts_by_source": {tag: len(rows) for tag, rows in related.items()},
            "workbooks": entries,
            "source_identity_index_path": str(index_path),
            "source_identity_index_sha256": sha256(index_path),
        }
        (scale_dir / "scale_input_manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    summaries = []
    for scale in SCALES:
        manifest_path = OUT / str(scale) / "scale_input_manifest.json"
        if manifest_path.is_file():
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            row_counts = manifest.get("row_counts_by_source")
            if row_counts is None:
                row_counts = {
                    tag: sum(int(entry["row_count"]) for entry in manifest["workbooks"] if entry["source_tag"] == tag)
                    for tag in ("B102-2002", "B101-2001", "B101-2002")
                }
            summaries.append({"scale": scale, "candidate_count": manifest["candidate_count"],
                              "row_counts": row_counts})
    top = {"version": "2.0.0", "scales": summaries}
    (OUT / "scale_generation_manifest.json").write_text(json.dumps(top, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(top, ensure_ascii=False))


if __name__ == "__main__":
    main()
