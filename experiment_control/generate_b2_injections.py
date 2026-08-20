#!/usr/bin/env python3
"""Create the target-neutral 20+6+4 Experiment-B2 input package."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from prepare_raw_inputs import sha256, write_payload


ROOT = Path(__file__).resolve().parents[1]
PREPARED = ROOT / "inputs" / "prepared"
OUT = ROOT / "inputs" / "experiment_b" / "B2"
ENGINE_DIR = ROOT / "plugin" / "fixed-combustion-inventory" / "engine"
if str(ENGINE_DIR) not in sys.path:
    sys.path.insert(0, str(ENGINE_DIR))
from source_domain import classify_source
POSITIVE_ROOTS = (
    "ACTIVITY_MISSING",
    "POLLUTANT_PARAMETER_MISSING",
    "SOURCE_RELATION_MISSING",
    "HARD_CONSTRAINT_CONFLICT",
)
POSITIVE_COUNTS = {
    "ACTIVITY_MISSING": 3,
    "POLLUTANT_PARAMETER_MISSING": 3,
    "SOURCE_RELATION_MISSING": 2,
    "HARD_CONSTRAINT_CONFLICT": 2,
}
# These three industrial parameter cases were independently pre-screened
# against the frozen B0 reference: all source-level SO2/PM10/PM25/BC/OC totals
# are calculated before injection.  Pinning the approved trio keeps the
# injection isolated from unrelated pre-existing blockers.
INDUSTRIAL_PARAMETER_B0_READY = (
    "SRC-RAW-D8333CABBDCF22B6AB37",
    "SRC-RAW-48EC88DCAB198382370E",
    "SRC-RAW-FF1045ABC23B668A936E",
)


def mapped_values(path: Path, raw_field: str) -> set[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {str(row[raw_field]).strip() for row in csv.DictReader(handle) if row.get("mapping_status") == "mapped"}


MAPPED_FUELS = mapped_values(ROOT / "rules" / "frozen" / "v1.0.1" / "mappings" / "source_fuel_aliases.csv", "raw_value")
MAPPED_COMBUSTION = mapped_values(
    ROOT / "rules" / "frozen" / "v1.0.1" / "mappings" / "source_combustion_aliases.csv",
    "raw_value",
)


def number(value: object) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(str(value).strip().replace(",", ""))
    except ValueError:
        return None


def stable_order(records: list[dict[str, object]], label: str) -> list[dict[str, object]]:
    return sorted(
        records,
        key=lambda record: hashlib.sha256(f"B2-v2|{label}|{record['source_id']}".encode()).hexdigest(),
    )


def fuel_candidates(row: dict[str, object]) -> list[dict[str, object]]:
    result = []
    for slot in ("一", "二"):
        amount = number(row.get(f"燃料{slot}消耗量"))
        fuel = str(row.get(f"燃料{slot}类型") or "").strip()
        unit = str(row.get(f"燃料{slot}消耗量单位") or "").strip()
        if amount is not None and amount > 0 and fuel and unit:
            result.append({"slot": slot, "amount": amount, "fuel": fuel, "unit": unit})
    return result


def coal_fuels(row: dict[str, object]) -> list[dict[str, object]]:
    return [fuel for fuel in fuel_candidates(row) if "煤" in str(fuel["fuel"]) or "焦炭" in str(fuel["fuel"])]


def preexisting_reason(row: dict[str, object]) -> str | None:
    equipment = str(row.get("工业锅炉类型") or row.get("电站锅炉/燃气轮机类型") or "").strip()
    combustion = str(row.get("工业锅炉燃烧方式") or row.get("电站锅炉燃烧方式") or "").strip()
    if equipment and not combustion:
        return "combustion_technology_missing_before_injection"
    other = number(row.get("其他燃料消耗总量（吨标准煤）"))
    if other is not None and other > 0 and not fuel_candidates(row):
        return "only_aggregate_other_fuel_before_injection"
    return None


def build_manifest(records: list[dict[str, object]]) -> dict[str, object]:
    counts = Counter(str(record["plan_class"]) for record in records)
    return {
        "version": "2.0.0",
        "design": "20_positive_injections_plus_6_negative_controls_plus_4_preexisting_problems",
        "identity_contract": "neutral_source_identity_v1",
        "planned_record_count": len(records),
        "class_counts": {
            "positive_injection": counts["positive_injection"],
            "negative_control": counts["negative_control"],
            "preexisting_problem": counts["preexisting_problem"],
        },
        "records": records,
    }


def build_class_audit(records: list[dict[str, object]], class_by_id: dict[str, str]) -> dict[str, object]:
    audited = [{"source_id": record["source_id"], "source_class": class_by_id[str(record["source_id"])],
                "plan_class": record["plan_class"], "root_cause": record["root_cause"]} for record in records]
    counts = {
        source_class: {
            plan_class: sum(item["source_class"] == source_class and item["plan_class"] == plan_class for item in audited)
            for plan_class in ("positive_injection", "negative_control", "preexisting_problem")
        }
        for source_class in ("INDUSTRIAL", "POWER")
    }
    expected = {"positive_injection": 10, "negative_control": 3, "preexisting_problem": 2}
    if any(counts[source_class] != expected for source_class in counts):
        raise ValueError(f"B2 class balance failed: {counts}")
    return {"version": "2.0.0", "purpose": "evaluation_side_balance_audit", "class_counts": counts, "records": audited}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    base_manifest = json.loads((PREPARED / "raw_input_manifest.json").read_text(encoding="utf-8"))
    identity = json.loads((PREPARED / "source_identity_index.json").read_text(encoding="utf-8"))
    source_by_row = {int(record["source_row"]): str(record["source_id"]) for record in identity["records"]}
    b102_entry = next(entry for entry in base_manifest["workbooks"] if entry["source_tag"] == "B102-2002")
    b102_path = Path(b102_entry["output_path"])
    workbook = load_workbook(b102_path, read_only=True, data_only=True)
    iterator = workbook[b102_entry["sheet"]].iter_rows(values_only=True)
    headers = list(next(iterator))
    records = []
    values = [headers]
    for source_row, raw in enumerate(iterator, start=2):
        row_values = list(raw[: len(headers)]) + [None] * max(0, len(headers) - len(raw))
        row = {str(headers[index]).strip(): row_values[index] for index in range(len(headers))}
        decision = classify_source(
            industry_code=row.get("行业类别代码"),
            industrial_equipment=row.get("工业锅炉类型"),
            power_equipment=row.get("电站锅炉/燃气轮机类型"),
        )
        records.append({"source_id": source_by_row[source_row], "source_row": source_row, "row": row,
                        "values": row_values, "source_class": decision.target})
        values.append(row_values)
    workbook.close()

    used: set[str] = set()
    plan: list[dict[str, object]] = []
    for source_class in ("INDUSTRIAL", "POWER"):
        preexisting = [record | {"reason": preexisting_reason(record["row"])} for record in records
                       if record["source_class"] == source_class]
        preexisting = [record for record in preexisting if record["reason"]]
        if len(preexisting) < 2:
            raise RuntimeError(f"not enough preexisting B2 candidates for {source_class}")
        for record in stable_order(preexisting, f"preexisting-{source_class}")[:2]:
            used.add(str(record["source_id"]))
            plan.append({
                "plan_class": "preexisting_problem",
                "root_cause": "PREEXISTING_INPUT_PROBLEM",
                "source_id": record["source_id"],
                "source_tag": "B102-2002",
                "source_row": record["source_row"],
                "existing_issue_reason": record["reason"],
                "changes": [],
            })

    def eligible(record: dict[str, object], root: str) -> bool:
        if str(record["source_id"]) in used or preexisting_reason(record["row"]):
            return False
        row = record["row"]
        fuels = fuel_candidates(row)
        coal = coal_fuels(row)
        if not fuels or any(str(fuel["fuel"]) not in MAPPED_FUELS for fuel in fuels):
            return False
        if root == "ACTIVITY_MISSING":
            return bool(fuels)
        if root in {"POLLUTANT_PARAMETER_MISSING", "HARD_CONSTRAINT_CONFLICT"}:
            if not coal:
                return False
            other_tce = number(row.get("其他燃料消耗总量（吨标准煤）"))
            if len(fuels) != 1 or (other_tce is not None and other_tce > 0):
                return False
            combustion = str(row.get("工业锅炉燃烧方式") or row.get("电站锅炉燃烧方式") or "").strip()
            slot = coal[0]["slot"]
            sulfur = row.get(f"燃料{slot}平均收到基含硫量")
            ash = row.get(f"燃料{slot}平均收到基灰分（%）")
            return combustion in MAPPED_COMBUSTION and number(sulfur) is not None and number(ash) is not None
        if root == "SOURCE_RELATION_MISSING":
            return bool(fuels and str(row.get("排放口编号") or "").strip())
        return False

    for source_class in ("INDUSTRIAL", "POWER"):
        for root in POSITIVE_ROOTS:
            if source_class == "INDUSTRIAL" and root == "POLLUTANT_PARAMETER_MISSING":
                record_by_id = {str(record["source_id"]): record for record in records}
                candidates = [record_by_id[source_id] for source_id in INDUSTRIAL_PARAMETER_B0_READY]
            else:
                candidates = stable_order(
                    [record for record in records if record["source_class"] == source_class and eligible(record, root)],
                    f"{source_class}-{root}",
                )
            required = POSITIVE_COUNTS[root]
            if len(candidates) < required:
                raise RuntimeError(f"not enough eligible B2 candidates for {source_class}/{root}")
            for record in candidates[:required]:
                used.add(str(record["source_id"]))
                row = record["row"]
                fuels = fuel_candidates(row)
                coal = coal_fuels(row)
                if root == "ACTIVITY_MISSING":
                    changes = [(f"燃料{fuels[0]['slot']}消耗量", None)]
                elif root == "POLLUTANT_PARAMETER_MISSING":
                    slot = coal[0]["slot"]
                    changes = [(f"燃料{slot}平均收到基含硫量", None), (f"燃料{slot}平均收到基灰分（%）", None)]
                elif root == "SOURCE_RELATION_MISSING":
                    changes = [("排放口编号", None)]
                else:
                    changes = [(f"燃料{coal[0]['slot']}平均收到基灰分（%）", -5)]
                applied = []
                row_values = values[int(record["source_row"]) - 1]
                for field, after in changes:
                    column = headers.index(field)
                    before = row_values[column]
                    row_values[column] = after
                    applied.append({"field": field, "before": before, "after": after})
                plan.append({
                    "plan_class": "positive_injection",
                    "root_cause": root,
                    "source_id": record["source_id"],
                    "source_tag": "B102-2002",
                    "source_row": record["source_row"],
                    "changes": applied,
                })

    for source_class in ("INDUSTRIAL", "POWER"):
        controls = [record for record in records if record["source_class"] == source_class
                    and str(record["source_id"]) not in used and not preexisting_reason(record["row"])
                    and fuel_candidates(record["row"])]
        if len(controls) < 3:
            raise RuntimeError(f"not enough negative controls for {source_class}")
        for record in stable_order(controls, f"negative-control-{source_class}")[:3]:
            used.add(str(record["source_id"]))
            plan.append({
                "plan_class": "negative_control",
                "root_cause": "NO_INJECTION_CONTROL",
                "source_id": record["source_id"],
                "source_tag": "B102-2002",
                "source_row": record["source_row"],
                "changes": [],
            })

    manifest = build_manifest(plan)
    from validate_input_variants import validate_injection_manifest
    validate_injection_manifest(manifest)
    class_audit = build_class_audit(plan, {str(record["source_id"]): str(record["source_class"]) for record in records})

    payload = {
        "source_tag": "B102-2002",
        "role": "device_fuel_base",
        "schema_mode": "controlled_missing_conflict_injection",
        "sheet_name": "Sheet1",
        "output_path": str(OUT / b102_path.name),
        "values": values,
        "variant": "B2",
        "cover_sheet": False,
    }
    payload_path = OUT / ".b102-b2.payload.json"
    payload_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    write_payload(payload_path)
    payload_path.unlink()
    for entry in base_manifest["workbooks"]:
        source = Path(entry["output_path"])
        if entry["source_tag"] != "B102-2002":
            shutil.copy2(source, OUT / source.name)
    shutil.copy2(PREPARED / "source_identity_index.json", OUT / "source_identity_index.json")
    (OUT / "injection_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "evaluation_class_audit.json").write_text(json.dumps(class_audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    variant_manifest = {
        "version": "2.0.0",
        "variant": "B2",
        "workbooks": [{"filename": Path(entry["output_path"]).name, "sha256": sha256(OUT / Path(entry["output_path"]).name)} for entry in base_manifest["workbooks"]],
        "source_identity_index_sha256": sha256(OUT / "source_identity_index.json"),
        "injection_manifest_path": str(OUT / "injection_manifest.json"),
        "injection_manifest_not_part_of_run_package": True,
    }
    (OUT / "variant_manifest.json").write_text(json.dumps(variant_manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUT), "planned_record_count": 30, "class_counts": manifest["class_counts"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
