#!/usr/bin/env python3

from __future__ import annotations

import csv
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[3]
MODULE_PATH = ROOT / "baseline" / "simple_deterministic" / "simple_inventory.py"
SPEC = importlib.util.spec_from_file_location("simple_inventory", MODULE_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class SimpleDeterministicBaselineTests(unittest.TestCase):
    def test_baseline_runtime_does_not_read_full_rule_package(self) -> None:
        source = MODULE_PATH.read_text(encoding="utf-8")
        self.assertNotIn("rules/frozen", source)
        self.assertNotIn("calculation_items", source)
        self.assertNotIn("control_calculation_admission", source)

    def test_basic_unit_conversion(self) -> None:
        self.assertEqual(MODULE.normalize_unit(2.0, "吨"), (2000.0, "kg"))
        self.assertEqual(MODULE.normalize_unit(2.0, "万立方米"), (20000.0, "m3"))
        self.assertIsNone(MODULE.normalize_unit(2.0, "吨标准煤"))

    def test_coal_sulfur_and_particle_formulas(self) -> None:
        coal = {
            "sulfur_to_bottom_ash": "0.15", "ash_to_bottom_ash": "0.85",
            "pm25_fraction": "0.07", "pm10_fraction": "0.20",
            "bc_fraction_of_pm25": "0.2", "oc_fraction_of_pm25": "0.04",
        }
        fuel = {"sulfur": 0.39, "sulfur_unit": "%", "ash": 15}
        sulfur = MODULE.factor_value({"mode": "coal_sulfur_balance"}, coal, fuel, None, "SO2")
        pm10 = MODULE.factor_value({"mode": "coal_particle_balance"}, coal, fuel, None, "PM10")
        pm25 = MODULE.factor_value({"mode": "coal_particle_balance"}, coal, fuel, None, "PM2.5")
        self.assertAlmostEqual(sulfur[0], 6.63)
        self.assertAlmostEqual(pm10[0], 4.5)
        self.assertAlmostEqual(pm25[0], 1.575)

    def test_fixed_lookup_has_natural_gas_and_no_nh3_factor(self) -> None:
        lookup = MODULE.FixedLookups()
        self.assertEqual(lookup.fixed_fuel("天然气"), ("天然气", "mapped"))
        keys = [key for key in lookup.factors if key[2] == "天然气"]
        self.assertTrue(keys)
        self.assertFalse(any(key[4] == "NH3" for key in keys))

    def test_4412_uses_electricity_production_department(self) -> None:
        self.assertEqual(MODULE.department("POWER", "4412"), "电力生产")

    def test_raw_candidate_classification_is_target_neutral(self) -> None:
        self.assertEqual(
            ("INDUSTRIAL", "采矿业和制造业"),
            MODULE.classify_candidate("2720", "燃气锅炉", ""),
        )
        self.assertEqual(
            ("POWER", "电力生产"),
            MODULE.classify_candidate("4412", "燃气锅炉", ""),
        )
        self.assertEqual(
            ("POWER", "热力生产和供应"),
            MODULE.classify_candidate("4430", "", "燃煤锅炉"),
        )
        self.assertEqual(("EXCLUDE", ""), MODULE.classify_candidate("4419", "", ""))

    def test_identity_lookup_uses_business_key_not_physical_subset_row(self) -> None:
        row_key = MODULE.identity_lookup_key({
            "统一社会信用代码": "CREDIT-ABC",
            "组织机构代码": "",
            "填报单位详细名称": "ENT-X",
            "序号": 12.0,
            "统计年份": 2022,
        })
        index_key = MODULE.identity_lookup_key({
            "pseudonymous_entity_id": "CREDIT-ABC",
            "sequence": "12",
            "inventory_year": "2022.0",
            "source_row": 999,
        })
        self.assertEqual(row_key, index_key)

    def test_six_raw_workbooks_are_classified_calculated_and_exported(self) -> None:
        def write_book(path: Path, headers: list[str], rows: list[list[object]]) -> None:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            workbook.save(path)

        with tempfile.TemporaryDirectory() as temporary:
            run_dir = Path(temporary) / "run"
            inputs = run_dir / "inputs"
            outputs = run_dir / "outputs"
            inputs.mkdir(parents=True)
            outputs.mkdir()
            b102_headers = [
                "统计年份", "组织机构代码", "统一社会信用代码", "填报单位详细名称",
                "行政区代码", "行政区名称", "行业类别代码", "行业类别名称", "序号",
                "电站锅炉/燃气轮机类型", "对应机组装机容量（万千瓦）", "电站锅炉燃烧方式",
                "工业锅炉类型", "工业锅炉燃烧方式",
                "燃料一类型", "燃料一消耗量", "燃料一消耗量单位", "燃料一平均收到基含硫量",
                "燃料一平均收到基含硫量单位", "燃料一平均收到基灰分（%）",
                "燃料二类型", "燃料二消耗量", "燃料二消耗量单位", "燃料二平均收到基含硫量",
                "燃料二平均收到基含硫量单位", "燃料二平均收到基灰分（%）",
                "其他燃料消耗总量（吨标准煤）", "排放口编号",
            ]
            b102_rows = [
                [2022, "ORG-I", "CREDIT-I", "ENT-I", "440100", "工业区", "2720", "制造业", 1,
                 "", "", "", "燃气锅炉", "室燃炉", "天然气", 2.0, "万立方米", "", "", "",
                 "", "", "", "", "", "", "", "DA001"],
                # 4412 deliberately has only the industrial-equipment column;
                # source classification must still route it to POWER and D-EP.
                [2022, "ORG-P", "CREDIT-P", "ENT-P", "440200", "电力区", "4412", "电力生产", 2,
                 "", "", "", "燃气锅炉", "室燃炉", "天然气", 1.0, "万立方米", "", "", "",
                 "", "", "", "", "", "", "", "DA002"],
                [2022, "ORG-X", "CREDIT-X", "ENT-X", "440300", "范围外", "4419", "其他电力", 3,
                 "", "", "", "", "", "", "", "", "", "", "",
                 "", "", "", "", "", "", "", "DA003"],
            ]
            b102 = inputs / "基102-2002_原始结构脱敏输入.xlsx"
            write_book(b102, b102_headers, b102_rows)

            master_headers = ["统计年份", "组织机构代码", "统一社会信用代码", "填报单位详细名称"]
            master_files = []
            for shard, row in enumerate((
                [2022, "ORG-I", "CREDIT-I", "ENT-I"],
                [2022, "ORG-P", "CREDIT-P", "ENT-P"],
                [2022, "ORG-X", "CREDIT-X", "ENT-X"],
                [2022, "ORG-Z", "CREDIT-Z", "ENT-Z"],
            ), start=1):
                path = inputs / f"基101-2001_原始结构脱敏输入_分片{shard:02d}-of04.xlsx"
                write_book(path, master_headers, [row])
                master_files.append(path)

            control = inputs / "基101-2002_原始结构脱敏输入.xlsx"
            write_book(
                control,
                ["统计年份", "组织机构代码", "统一社会信用代码", "填报单位详细名称", "对应的排放口代码", "处理工艺名称", "去除效率（%）"],
                [[2022, "ORG-P", "CREDIT-P", "ENT-P", "DZGL", "选择性催化还原法（SCR）", 50]],
            )
            ids = ["SRC-RAW-00000000000000000001", "SRC-RAW-00000000000000000002", "SRC-RAW-00000000000000000003"]
            identity = inputs / "source_identity_index.json"
            identity.write_text(json.dumps({
                "version": "1.0.0",
                "candidate_ids": ids,
                "records": [
                    {
                        "source_id": ids[index],
                        "source_row": index + 2,
                        "pseudonymous_entity_id": f"CREDIT-{'IPX'[index]}",
                        "sequence": index + 1,
                        "inventory_year": 2022,
                    }
                    for index in range(3)
                ],
            }), encoding="utf-8")
            manifest_inputs = [
                {"tag": "B102-2002", "role": "device_fuel_base", "path": str(b102), "sheet": "Sheet1"},
                *[{"tag": "B101-2001", "role": "enterprise_master", "path": str(path), "sheet": "Sheet1"} for path in master_files],
                {"tag": "B101-2002", "role": "control_facility_base", "path": str(control), "sheet": "Sheet1"},
            ]
            (run_dir / "run_manifest.json").write_text(json.dumps({
                "run_id": "A-POW-DET-R1", "experiment_method": "deterministic_program",
                "target": "POWER", "inputs": manifest_inputs,
                "source_identity_index_path": str(identity),
            }), encoding="utf-8")

            summary = MODULE.execute_run(run_dir)
            self.assertEqual(3, summary["candidate_source_count"])
            self.assertEqual(1, summary["target_source_count"])
            self.assertEqual(4, summary["enterprise_master_rows"])
            for name in (
                "source_decisions.csv", "source_pollutant_totals.csv", "generic_inventory_items.csv",
                "exceptions.csv", "简单确定性脚本目标清单.xlsx",
            ):
                self.assertTrue((outputs / name).is_file(), name)

            with (outputs / "source_decisions.csv").open(encoding="utf-8-sig", newline="") as handle:
                decisions = list(csv.DictReader(handle))
            self.assertEqual(["INDUSTRIAL", "POWER", "EXCLUDE"], [row["actual_target"] for row in decisions])

            with (outputs / "source_pollutant_totals.csv").open(encoding="utf-8-sig", newline="") as handle:
                totals = list(csv.DictReader(handle))
            nox = next(row for row in totals if row["source_id"] == ids[1] and row["pollutant"] == "NOx")
            self.assertEqual("calculated", nox["status"])
            self.assertAlmostEqual(0.041, float(nox["generation_t"]))
            self.assertAlmostEqual(0.0205, float(nox["emission_t"]))
            parameters = json.loads(nox["parameter_record_json"])
            self.assertEqual(4.1, parameters[0]["value"])

            workbook = load_workbook(outputs / "简单确定性脚本目标清单.xlsx", read_only=True, data_only=True)
            self.assertEqual(2, workbook["目标清单"].max_row)
            workbook.close()

    @unittest.skipUnless((ROOT / "inputs" / "prepared" / "raw_input_manifest.json").is_file(), "prepared raw package absent")
    def test_real_raw_package_classification_and_4412_factor_routing(self) -> None:
        prepared = ROOT / "inputs" / "prepared"
        raw_manifest = json.loads((prepared / "raw_input_manifest.json").read_text(encoding="utf-8"))
        inputs = [{
            "tag": entry["source_tag"],
            "role": entry["role"],
            "path": entry["output_path"],
            "sheet": entry["sheet"],
        } for entry in raw_manifest["workbooks"]]
        bundle = MODULE._load_raw_bundle({
            "target": "POWER",
            "inputs": inputs,
            "source_identity_index_path": str(prepared / "source_identity_index.json"),
        }, MODULE.FixedLookups())
        counts = {
            target: sum(row["actual_target"] == target for row in bundle["decisions"])
            for target in ("INDUSTRIAL", "POWER", "EXCLUDE")
        }
        self.assertEqual({"INDUSTRIAL": 4046, "POWER": 615, "EXCLUDE": 1261}, counts)

        # This 4412 record is stored with industrial-boiler fields in Base-102.
        # It must nevertheless use D-EP natural-gas NOx=4.1, not D-ES 2.09.
        source_id = "SRC-RAW-B4BC1019B13822BA6B72"
        record = bundle["records"][source_id]
        self.assertEqual("POWER", record["actual_target"])
        self.assertEqual("电力生产", record["department"])
        rows, _ = MODULE.calculate_source(source_id, record, MODULE.FixedLookups())
        nox = next(row for row in rows if row["pollutant"] == "NOx")
        self.assertEqual(4.1, json.loads(nox["parameter_record_json"])[0]["value"])


if __name__ == "__main__":
    unittest.main()
