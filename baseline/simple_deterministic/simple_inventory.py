#!/usr/bin/env python3
"""Fixed-header, fixed-lookup and direct-calculation inventory baseline.

The baseline reads the raw Base-102 candidate table and Base-101 master /
control tables. It classifies every candidate with a fixed source-code mapping
before calculating only the requested target class. It deliberately has no
Agent, declarative rule package, admission gate, pollutant-item model, or full
calculation trace.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict, deque
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook


HERE = Path(__file__).resolve().parent
LOOKUPS = HERE / "lookups"
POLLUTANTS = ("SO2", "NOx", "CO", "VOC", "PM10", "PM2.5", "BC", "OC", "NH3")
FACTOR_POLLUTANT = {"NOx": "NOX", "VOC": "VOCS", "PM2.5": "PM25"}
POWER_CODES = {"4411", "4412", "4417"}
HEAT_CODE = "4430"
INDUSTRIAL_EQUIPMENT = {"燃煤锅炉", "燃气锅炉", "燃油锅炉", "其他锅炉"}

BASE102_REQUIRED = {
    "统计年份", "组织机构代码", "统一社会信用代码", "填报单位详细名称",
    "行政区代码", "行政区名称", "行业类别代码", "行业类别名称", "序号",
    "电站锅炉/燃气轮机类型", "对应机组装机容量（万千瓦）", "电站锅炉燃烧方式",
    "工业锅炉类型", "工业锅炉燃烧方式",
    "燃料一类型", "燃料一消耗量", "燃料一消耗量单位", "燃料一平均收到基含硫量",
    "燃料一平均收到基含硫量单位", "燃料一平均收到基灰分（%）",
    "燃料二类型", "燃料二消耗量", "燃料二消耗量单位", "燃料二平均收到基含硫量",
    "燃料二平均收到基含硫量单位", "燃料二平均收到基灰分（%）",
    "其他燃料消耗总量（吨标准煤）", "排放口编号",
}
ENTITY_REQUIRED = {"统计年份", "组织机构代码", "统一社会信用代码", "填报单位详细名称"}
CONTROL_REQUIRED = ENTITY_REQUIRED | {"对应的排放口代码", "处理工艺名称", "去除效率（%）"}

BASELINE_FIELDS = (
    "source_id", "target", "pollutant", "status", "method",
    "activity_record_json", "parameter_record_json", "control_record_json",
    "generation_t", "emission_t", "standard_reference", "reason_code",
)
GENERIC_FIELDS = (
    "source_id", "target", "pollutant", "status", "method",
    "activity_value", "activity_unit", "parameter_value", "parameter_unit",
    "control_method", "control_efficiency", "generation_t", "emission_t",
    "standard_reference", "reason_code",
)


class RawInputError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def read_csv(name: str) -> list[dict[str, str]]:
    with (LOOKUPS / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def fraction(value: Any) -> float:
    result = number(value)
    if result is None or result < 0 or result > 100:
        return 0.0
    return result / 100.0


def normalize_unit(amount: float, unit: str) -> tuple[float, str] | None:
    text = str(unit or "").strip()
    if text in {"吨", "t"}:
        return amount * 1000.0, "kg"
    if text in {"万立方米", "万m3", "万m³"}:
        return amount * 10000.0, "m3"
    return None


def _code(value: Any) -> str:
    text = str(value or "").strip().split(".", 1)[0]
    return text.zfill(4) if text.isdigit() else text


def department(target: str, industry_code: Any) -> str:
    if target == "INDUSTRIAL":
        return "采矿业和制造业"
    return "热力生产和供应" if _code(industry_code) == HEAT_CODE else "电力生产"


def classify_candidate(industry_code: Any, industrial_equipment: Any, power_equipment: Any) -> tuple[str, str]:
    """Classify one Base-102 candidate with a fixed target-neutral mapping."""
    code = _code(industry_code)
    if code in POWER_CODES:
        return "POWER", "电力生产"
    if code == HEAT_CODE:
        return "POWER", "热力生产和供应"
    try:
        numeric = int(code)
    except ValueError:
        numeric = -1
    equipment = str(industrial_equipment or power_equipment or "").strip()
    if 600 <= numeric <= 4399 and equipment in INDUSTRIAL_EQUIPMENT:
        return "INDUSTRIAL", "采矿业和制造业"
    return "EXCLUDE", ""


def _entity_id(row: Mapping[str, Any]) -> str:
    for header in ("统一社会信用代码", "组织机构代码", "填报单位详细名称"):
        value = str(row.get(header) or "").strip()
        if value:
            return value
    return ""


def _identity_scalar(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except ValueError:
            pass
    return text


def identity_lookup_key(row: Mapping[str, Any]) -> tuple[str, str, str]:
    entity = str(row.get("pseudonymous_entity_id") or _entity_id(row)).strip()
    sequence = row.get("sequence") if "sequence" in row else row.get("序号")
    year = row.get("inventory_year") if "inventory_year" in row else row.get("统计年份")
    return entity, _identity_scalar(sequence), _identity_scalar(year)


def _read_workbook(spec: Mapping[str, Any], required: set[str]) -> list[tuple[int, dict[str, Any]]]:
    path = Path(str(spec["path"]))
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = str(spec.get("sheet") or "Sheet1")
        if sheet_name not in workbook.sheetnames:
            raise RawInputError("FIXED_INPUT_SCHEMA_MISMATCH", f"missing sheet {sheet_name}: {path.name}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration as exc:
            raise RawInputError("FIXED_INPUT_SCHEMA_MISMATCH", f"empty workbook: {path.name}") from exc
        missing = sorted(required - set(headers))
        if missing:
            raise RawInputError("FIXED_INPUT_SCHEMA_MISMATCH", f"{path.name} missing exact headers: {', '.join(missing)}")
        first_index: dict[str, int] = {}
        for index, header in enumerate(headers):
            first_index.setdefault(header, index)
        rows = []
        for excel_row, values in enumerate(iterator, start=2):
            padded = tuple(values) + (None,) * max(0, len(headers) - len(values))
            if not any(value not in (None, "") for value in padded):
                continue
            rows.append((excel_row, {header: padded[index] for header, index in first_index.items()}))
        return rows
    finally:
        workbook.close()


def _inputs_by_role(manifest: Mapping[str, Any]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for spec in manifest.get("inputs", []):
        grouped[str(spec.get("role") or "")].append(dict(spec))
    if len(grouped["device_fuel_base"]) != 1:
        raise RawInputError("FIXED_INPUT_SCHEMA_MISMATCH", "expected exactly one Base-102 device/fuel workbook")
    if not grouped["enterprise_master"]:
        raise RawInputError("FIXED_INPUT_SCHEMA_MISMATCH", "missing Base-101 enterprise master workbook shards")
    if len(grouped["control_facility_base"]) != 1:
        raise RawInputError("FIXED_INPUT_SCHEMA_MISMATCH", "expected exactly one Base-101 control workbook")
    return grouped


def _load_identity_index(manifest: Mapping[str, Any]) -> tuple[list[str], dict[tuple[str, str, str], dict[str, Any]]]:
    path = Path(str(manifest.get("source_identity_index_path") or ""))
    if not path.is_file():
        raise RawInputError("FIXED_IDENTITY_INDEX_MISSING", "source identity index is missing")
    payload = json.loads(path.read_text(encoding="utf-8"))
    candidate_ids = [str(value) for value in payload.get("candidate_ids", [])]
    if not candidate_ids or len(candidate_ids) != len(set(candidate_ids)):
        raise RawInputError("FIXED_IDENTITY_INDEX_INVALID", "candidate source IDs are empty or duplicated")
    records = payload.get("records") or []
    by_key = {identity_lookup_key(row): dict(row) for row in records}
    if len(by_key) != len(records):
        raise RawInputError("FIXED_IDENTITY_INDEX_INVALID", "identity index has duplicate business keys")
    return candidate_ids, by_key


class FixedLookups:
    def __init__(self) -> None:
        self.fuels = {row["raw_value"]: row for row in read_csv("fixed_fuel_mapping.csv")}
        self.combustion = {row["raw_value"]: row for row in read_csv("fixed_combustion_mapping.csv")}
        factors = read_csv("fixed_industrial_factors.csv") + read_csv("fixed_power_factors.csv")
        self.factors = {
            (row["sector"], row["department"], row["fuel"], row["combustion_technology"], row["pollutant"]): row
            for row in factors
        }
        self.coal = {
            (row["sector"], row["department"], row["combustion_technology"]): row
            for row in read_csv("fixed_coal_parameters.csv")
        }
        self.control_keywords: dict[str, list[str]] = defaultdict(list)
        for row in read_csv("fixed_control_keywords.csv"):
            self.control_keywords[row["control_group"]].append(row["keyword"].upper())
        self.available_technologies: dict[tuple[str, str, str], set[str]] = defaultdict(set)
        for sector, dept, fuel, tech, _pollutant in self.factors:
            self.available_technologies[(sector, dept, fuel)].add(tech)

    def fixed_fuel(self, raw: Any) -> tuple[str | None, str]:
        row = self.fuels.get(str(raw or "").strip())
        if not row:
            return None, "unmapped"
        return (row["standard_fuel"] or None), row["mapping_status"]

    def fixed_technology(self, sector: str, dept: str, fuel: str, raw: Any) -> str | None:
        choices = self.available_technologies.get((sector, dept, fuel), set())
        if fuel == "煤炭":
            mapped = self.combustion.get(str(raw or "").strip())
            preferred = mapped.get("standard_combustion_technology") if mapped else None
        elif "燃气锅炉" in choices:
            preferred = "燃气锅炉"
        elif "燃油锅炉" in choices:
            preferred = "燃油锅炉"
        else:
            preferred = "不分技术"
        if preferred in choices:
            return preferred
        return "不分技术" if "不分技术" in choices else None

    def control_group(self, process: Any) -> set[str]:
        text = str(process or "").upper()
        return {group for group, keywords in self.control_keywords.items() if any(keyword in text for keyword in keywords)}


def _load_enterprises(specs: Sequence[Mapping[str, Any]]) -> tuple[set[str], int]:
    entities: set[str] = set()
    count = 0
    for spec in specs:
        rows = _read_workbook(spec, ENTITY_REQUIRED)
        count += len(rows)
        entities.update(_entity_id(row) for _excel_row, row in rows if _entity_id(row))
    return entities, count


def _load_controls(spec: Mapping[str, Any], lookup: FixedLookups) -> tuple[dict[tuple[str, str], dict[str, float]], int]:
    grouped: dict[tuple[str, str], dict[str, float]] = defaultdict(dict)
    rows = _read_workbook(spec, CONTROL_REQUIRED)
    for _excel_row, row in rows:
        entity = _entity_id(row)
        outlet = str(row.get("对应的排放口代码") or "").strip()
        if not entity or not outlet:
            continue
        efficiency = fraction(row.get("去除效率（%）"))
        for group in lookup.control_group(row.get("处理工艺名称")):
            grouped[(entity, outlet)][group] = max(grouped[(entity, outlet)].get(group, 0.0), efficiency)
    return dict(grouped), len(rows)


def _fuel_rows(row: Mapping[str, Any]) -> list[dict[str, Any]]:
    result = []
    for slot, prefix in (("fuel_1", "燃料一"), ("fuel_2", "燃料二")):
        result.append({
            "slot": slot, "raw": row.get(f"{prefix}类型"), "amount": row.get(f"{prefix}消耗量"),
            "unit": row.get(f"{prefix}消耗量单位"), "sulfur": row.get(f"{prefix}平均收到基含硫量"),
            "sulfur_unit": row.get(f"{prefix}平均收到基含硫量单位"),
            "ash": row.get(f"{prefix}平均收到基灰分（%）"),
        })
    result.append({
        "slot": "other_tce", "raw": "其他燃料（吨标准煤）",
        "amount": row.get("其他燃料消耗总量（吨标准煤）"), "unit": "吨标准煤",
        "sulfur": None, "sulfur_unit": None, "ash": None,
    })
    return result


def factor_value(row: dict[str, str], coal: dict[str, str] | None, fuel: dict[str, Any], capacity_mw: float | None, pollutant: str) -> tuple[float, str] | None:
    mode = row["mode"]
    if mode == "constant":
        value = number(row["value"])
        return (value, row["unit"]) if value is not None else None
    if mode == "capacity_lookup":
        if capacity_mw is None:
            return None
        return (8.96 if capacity_mw <= 100 else 8.19 if capacity_mw < 300 else 7.21), "g/kg燃料"
    if not coal:
        return None
    if mode == "coal_sulfur_balance":
        sulfur = number(fuel["sulfur"])
        bottom = number(coal["sulfur_to_bottom_ash"])
        if sulfur is None or bottom is None or str(fuel["sulfur_unit"] or "").strip() != "%":
            return None
        return 20.0 * sulfur * (1.0 - bottom), "g/kg燃料"
    if mode == "coal_particle_balance":
        ash = number(fuel["ash"])
        bottom = number(coal["ash_to_bottom_ash"])
        if ash is None or bottom is None:
            return None
        fine = number(coal["pm25_fraction"])
        coarse = number(coal["pm10_fraction"])
        if pollutant == "PM10" and coarse is not None:
            return 10.0 * ash * (1.0 - bottom) * coarse, "g/kg燃料"
        if pollutant == "PM2.5" and fine is not None:
            return 10.0 * ash * (1.0 - bottom) * fine, "g/kg燃料"
        if pollutant in {"BC", "OC"} and fine is not None:
            share = number(coal["bc_fraction_of_pm25" if pollutant == "BC" else "oc_fraction_of_pm25"])
            if share is not None:
                return 10.0 * ash * (1.0 - bottom) * fine * share, "g/kg燃料"
    return None


def _control_record(record: Mapping[str, Any], pollutant: str) -> dict[str, float | None]:
    efficiencies = record["control_efficiencies"]
    if pollutant == "SO2":
        return {"efficiency": efficiencies.get("SO2", 0.0), "fine_efficiency": None, "coarse_efficiency": None}
    if pollutant == "NOx":
        return {"efficiency": efficiencies.get("NOx", 0.0), "fine_efficiency": None, "coarse_efficiency": None}
    if pollutant == "PM10":
        value = efficiencies.get("PM", 0.0)
        return {"efficiency": None, "fine_efficiency": value, "coarse_efficiency": value}
    if pollutant in {"PM2.5", "BC", "OC"}:
        return {"efficiency": efficiencies.get("PM", 0.0), "fine_efficiency": None, "coarse_efficiency": None}
    return {"efficiency": 0.0, "fine_efficiency": None, "coarse_efficiency": None}


def _load_raw_bundle(manifest: Mapping[str, Any], lookup: FixedLookups) -> dict[str, Any]:
    grouped = _inputs_by_role(manifest)
    all_ids, identity_by_key = _load_identity_index(manifest)
    selected = set(manifest.get("candidate_ids") or all_ids)
    if not selected.issubset(set(all_ids)):
        raise RawInputError("FIXED_IDENTITY_INDEX_INVALID", "candidate subset is outside the identity index")
    enterprises, enterprise_rows = _load_enterprises(grouped["enterprise_master"])
    control_map, control_rows = _load_controls(grouped["control_facility_base"][0], lookup)
    b102_rows = _read_workbook(grouped["device_fuel_base"][0], BASE102_REQUIRED)
    records: dict[str, dict[str, Any]] = {}
    decisions = []
    for excel_row, row in b102_rows:
        identity = identity_by_key.get(identity_lookup_key(row))
        if not identity:
            raise RawInputError("FIXED_IDENTITY_INDEX_INVALID", f"Base-102 row {excel_row} lacks a source identity business key")
        source_id = str(identity["source_id"])
        if source_id not in selected:
            continue
        entity = str(identity.get("pseudonymous_entity_id") or _entity_id(row))
        actual_target, dept = classify_candidate(row.get("行业类别代码"), row.get("工业锅炉类型"), row.get("电站锅炉/燃气轮机类型"))
        outlet_category = "DZGL" if actual_target == "POWER" else "GYGL"
        raw_capacity = number(row.get("对应机组装机容量（万千瓦）"))
        if actual_target == "POWER":
            equipment = row.get("电站锅炉/燃气轮机类型") or row.get("工业锅炉类型")
            combustion = row.get("电站锅炉燃烧方式") or row.get("工业锅炉燃烧方式")
        else:
            equipment = row.get("工业锅炉类型") or row.get("电站锅炉/燃气轮机类型")
            combustion = row.get("工业锅炉燃烧方式") or row.get("电站锅炉燃烧方式")
        records[source_id] = {
            "source_id": source_id, "entity_id": entity, "raw_excel_row": excel_row, "row": row,
            "actual_target": actual_target, "department": dept, "equipment": equipment,
            "combustion": combustion, "capacity_mw": raw_capacity * 10.0 if raw_capacity is not None else None,
            "control_efficiencies": control_map.get((entity, outlet_category), {}),
            "enterprise_master_present": entity in enterprises,
        }
        decisions.append({
            "source_id": source_id, "actual_target": actual_target, "department": dept,
            "decision_reason": "fixed_industry_and_equipment_mapping",
            "selected_for_run_target": actual_target == manifest.get("target"),
            "enterprise_master_present": entity in enterprises,
        })
    missing = selected - set(records)
    if missing:
        raise RawInputError("FIXED_INPUT_ROW_MISSING", f"{len(missing)} selected candidate rows are missing")
    return {"records": records, "decisions": decisions, "candidate_count": len(selected), "enterprise_rows": enterprise_rows, "control_rows": control_rows}


def calculate_source(source_id: str, record: dict[str, Any], lookup: FixedLookups) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    target = record["actual_target"]
    sector = "工业源" if target == "INDUSTRIAL" else "电力热力源"
    dept = record["department"]
    by_pollutant: dict[str, list[dict[str, Any]]] = defaultdict(list)
    simple_reasons: set[str] = set()
    for fuel in _fuel_rows(record["row"]):
        amount = number(fuel["amount"])
        if amount in (None, 0.0):
            continue
        if amount < 0 or fuel["slot"] == "other_tce":
            for pollutant in POLLUTANTS[:-1]:
                by_pollutant[pollutant].append({"status": "information_insufficient", "reason": "DIRECT_INPUT_OR_LOOKUP_MISSING"})
            simple_reasons.add("DIRECT_INPUT_OR_LOOKUP_MISSING")
            continue
        standard_fuel, map_status = lookup.fixed_fuel(fuel["raw"])
        if map_status == "outside_fossil_scope":
            for pollutant in POLLUTANTS[:-1]:
                by_pollutant[pollutant].append({"status": "not_involved", "reason": "OUTSIDE_FIXED_LOOKUP_SCOPE"})
            continue
        normalized = normalize_unit(amount, str(fuel["unit"] or ""))
        if not standard_fuel or not normalized:
            for pollutant in POLLUTANTS[:-1]:
                by_pollutant[pollutant].append({"status": "information_insufficient", "reason": "DIRECT_INPUT_OR_LOOKUP_MISSING"})
            simple_reasons.add("DIRECT_INPUT_OR_LOOKUP_MISSING")
            continue
        activity, activity_unit = normalized
        technology = lookup.fixed_technology(sector, dept, standard_fuel, record["combustion"])
        coal = lookup.coal.get((sector, dept, technology or "")) if standard_fuel == "煤炭" else None
        for pollutant in POLLUTANTS[:-1]:
            factor_row = lookup.factors.get((sector, dept, standard_fuel, technology or "", FACTOR_POLLUTANT.get(pollutant, pollutant)))
            resolved = factor_value(factor_row, coal, fuel, record["capacity_mw"], pollutant) if factor_row else None
            if not factor_row or not resolved:
                by_pollutant[pollutant].append({"status": "information_insufficient", "reason": "DIRECT_INPUT_OR_LOOKUP_MISSING"})
                simple_reasons.add("DIRECT_INPUT_OR_LOOKUP_MISSING")
                continue
            factor, factor_unit = resolved
            generation = activity * factor / 1_000_000.0
            control = _control_record(record, pollutant)
            if pollutant == "PM10":
                pm25_row = lookup.factors.get((sector, dept, standard_fuel, technology or "", "PM25"))
                pm25 = factor_value(pm25_row, coal, fuel, record["capacity_mw"], "PM2.5") if pm25_row else None
                if not pm25:
                    by_pollutant[pollutant].append({"status": "information_insufficient", "reason": "DIRECT_INPUT_OR_LOOKUP_MISSING"})
                    simple_reasons.add("DIRECT_INPUT_OR_LOOKUP_MISSING")
                    continue
                fine_generation = activity * pm25[0] / 1_000_000.0
                emission = fine_generation * (1.0 - float(control["fine_efficiency"]))
                emission += (generation - fine_generation) * (1.0 - float(control["coarse_efficiency"]))
            else:
                emission = generation * (1.0 - float(control["efficiency"]))
            by_pollutant[pollutant].append({
                "status": "calculated", "activity": activity, "activity_unit": activity_unit,
                "factor": factor, "factor_unit": factor_unit, "control": control,
                "generation_t": generation, "emission_t": emission,
                "method": factor_row["mode"], "standard_reference": factor_row["source_section"],
            })
    by_pollutant["NH3"].append({"status": "not_involved", "reason": "NO_FIXED_FUEL_FACTOR"})

    rows = []
    for pollutant in POLLUTANTS:
        parts = by_pollutant.get(pollutant, [])
        statuses = {part["status"] for part in parts}
        status = "information_insufficient" if "information_insufficient" in statuses else "calculated" if "calculated" in statuses else "not_involved"
        calculated = [part for part in parts if part["status"] == "calculated"]
        reason = next((part.get("reason", "") for part in parts if part["status"] != "calculated"), "")
        rows.append({
            "source_id": source_id, "target": target, "pollutant": pollutant, "status": status,
            "method": ";".join(sorted({part["method"] for part in calculated})),
            "activity_record_json": json.dumps([{"value": part["activity"], "unit": part["activity_unit"]} for part in calculated], ensure_ascii=False, separators=(",", ":")),
            "parameter_record_json": json.dumps([{"value": part["factor"], "unit": part["factor_unit"]} for part in calculated], ensure_ascii=False, separators=(",", ":")),
            "control_record_json": json.dumps([part["control"] for part in calculated], ensure_ascii=False, separators=(",", ":")),
            "generation_t": sum(part["generation_t"] for part in calculated) if status == "calculated" else "",
            "emission_t": sum(part["emission_t"] for part in calculated) if status == "calculated" else "",
            "standard_reference": ";".join(sorted({part["standard_reference"] for part in calculated})),
            "reason_code": reason,
        })
    return rows, {"source_id": source_id, "actual_target": target, "reason_codes": sorted(simple_reasons)}


def _write_csv(path: Path, fields: Sequence[str], rows: Iterable[Mapping[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_outputs(run_dir: Path, rows: list[dict[str, Any]], unresolved: list[dict[str, Any]], decisions: list[dict[str, Any]], metadata: Mapping[str, Any]) -> dict[str, Any]:
    output_dir = run_dir / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "source_decisions.csv", ("source_id", "actual_target", "department", "decision_reason", "selected_for_run_target", "enterprise_master_present"), decisions)
    _write_csv(output_dir / "baseline_inventory_items.csv", BASELINE_FIELDS, rows)
    _write_csv(output_dir / "source_pollutant_totals.csv", BASELINE_FIELDS, rows)
    generic = [{
        "source_id": row["source_id"], "target": row["target"], "pollutant": row["pollutant"],
        "status": row["status"], "method": row["method"], "activity_value": "", "activity_unit": "",
        "parameter_value": "", "parameter_unit": "", "control_method": "", "control_efficiency": "",
        "generation_t": row["generation_t"], "emission_t": row["emission_t"],
        "standard_reference": row["standard_reference"], "reason_code": row["reason_code"],
    } for row in rows]
    _write_csv(output_dir / "generic_inventory_items.csv", GENERIC_FIELDS, generic)
    exception_rows = [{
        "source_id": item.get("source_id", ""), "actual_target": item.get("actual_target", ""),
        "reason_code": reason,
        "description": "fixed direct input or lookup could not produce all requested totals",
    } for item in unresolved for reason in item.get("reason_codes", [])]
    _write_csv(output_dir / "exceptions.csv", ("source_id", "actual_target", "reason_code", "description"), exception_rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "目标清单"
    sheet.append(["源ID", "源类"] + [f"{p}_状态" for p in POLLUTANTS] + [f"{p}_排放量(t)" for p in POLLUTANTS])
    grouped_rows: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in rows:
        grouped_rows[row["source_id"]][row["pollutant"]] = row
    for source_id, values in grouped_rows.items():
        target = next(iter(values.values()))["target"]
        sheet.append([source_id, target] + [values[p]["status"] for p in POLLUTANTS] + [values[p]["emission_t"] for p in POLLUTANTS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook_path = output_dir / "简单确定性脚本目标清单.xlsx"
    workbook.save(workbook_path)

    summary = {
        "baseline_definition": "fixed_mapping_plus_fixed_lookup_plus_direct_calculation",
        "candidate_source_count": int(metadata.get("candidate_count", len(decisions))),
        "target_source_count": len(grouped_rows),
        "enterprise_master_rows": int(metadata.get("enterprise_rows", 0)),
        "control_facility_rows": int(metadata.get("control_rows", 0)),
        "source_decision_counts": {target: sum(item.get("actual_target") == target for item in decisions) for target in ("INDUSTRIAL", "POWER", "EXCLUDE", "UNRESOLVED")},
        "normalized_rows": len(rows),
        "status_counts": {status: sum(row["status"] == status for row in rows) for status in ("calculated", "not_involved", "information_insufficient")},
        "exception_rows": len(exception_rows),
        "complete_process_record": False,
        "source_classification_performed": True,
    }
    (output_dir / "baseline_run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output_dir / "machine_qa_simple_deterministic.json").write_text(json.dumps({
        "formula_error_count": 0, "external_link_count": 0, "workbook": workbook_path.name,
        "source_decision_rows": len(decisions), "source_pollutant_rows": len(rows),
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def _identity_ids_for_failure(manifest: Mapping[str, Any]) -> list[str]:
    try:
        ids, _ = _load_identity_index(manifest)
    except RawInputError:
        return []
    return [str(value) for value in (manifest.get("candidate_ids") or ids)]


def execute_run(run_dir: Path) -> dict[str, Any]:
    run_dir = run_dir.resolve()
    manifest = json.loads((run_dir / "run_manifest.json").read_text(encoding="utf-8"))
    if manifest.get("experiment_method") != "deterministic_program":
        raise ValueError("simple baseline only accepts deterministic_program")
    lookup = FixedLookups()
    try:
        bundle = _load_raw_bundle(manifest, lookup)
    except RawInputError as exc:
        ids = _identity_ids_for_failure(manifest)
        decisions = [{
            "source_id": source_id, "actual_target": "UNRESOLVED", "department": "",
            "decision_reason": exc.code, "selected_for_run_target": False, "enterprise_master_present": "",
        } for source_id in ids]
        unresolved = [{"source_id": source_id, "actual_target": "UNRESOLVED", "reason_codes": [exc.code]} for source_id in ids]
        return write_outputs(run_dir, [], unresolved, decisions, {"candidate_count": len(ids)})

    target = str(manifest.get("target") or "")
    rows: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []
    for decision in bundle["decisions"]:
        if decision["actual_target"] != target:
            continue
        source_rows, unresolved_row = calculate_source(decision["source_id"], bundle["records"][decision["source_id"]], lookup)
        rows.extend(source_rows)
        unresolved.append(unresolved_row)
    return write_outputs(run_dir, rows, unresolved, bundle["decisions"], bundle)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("run_package", type=Path)
    args = parser.parse_args()
    print(json.dumps(execute_run(args.run_package), ensure_ascii=False))


if __name__ == "__main__":
    main()
