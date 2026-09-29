from __future__ import annotations

import argparse
import csv
import json
import random
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PREPARED = ROOT / 'inputs/prepared'
REFERENCE = ROOT / 'reference/frozen/v2.0.0'
NODE = Path('/Users/wushuo/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node')
CLASSES = ('INDUSTRIAL', 'POWER')
ROOTS = ('POLLUTANT_PARAMETER_MISSING', 'SOURCE_RELATION_MISSING',
         'ACTIVITY_MISSING', 'HARD_CONSTRAINT_CONFLICT', 'NO_ANOMALY')


def rows(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def table(path, records):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows({k: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v
                         for k, v in r.items()} for r in records)


def workbook_values(path):
    w = load_workbook(path, read_only=True, data_only=True)
    s = w['Sheet1']
    s.reset_dimensions()
    values = [list(r) for r in s.iter_rows(values_only=True)]
    w.close()
    return values


def build():
    identity = json.loads((PREPARED / 'source_identity_index.json').read_text())
    by_row = {r['source_row']: r for r in identity['records']}
    manifest = json.loads((PREPARED / 'raw_input_manifest.json').read_text())
    dev = json.loads((ROOT / 'inputs/experiment_b/B2/injection_manifest.json').read_text())['records']
    excluded = {r['source_id'] for r in dev}
    excluded |= {r['source_id'] for r in rows(REFERENCE / 'expected_exceptions.csv')
                 if r['variant'] == 'B0' and r['requires_user_judgment'] == 'true'}
    excluded |= {r['源ID'] for r in rows(ROOT.parent / '审计报告/规则包与参照包逐条核查_20260907/受影响条目定位/01_涉及源记录_326条.csv')}
    sources = {r['source_id']: r for r in rows(REFERENCE / 'expected_sources.csv') if r['variant'] == 'B0'}
    totals = defaultdict(dict)
    for r in rows(REFERENCE / 'expected_source_pollutant_totals.csv'):
        if r['variant'] == 'B0':
            totals[r['source_id']][r['pollutant']] = r
    aliases = {r['raw_value']: r['standard_fuel'] for r in rows(ROOT / 'rules/frozen/v1.0.1/mappings/source_fuel_aliases.csv')
               if r['mapping_status'] == 'mapped'}
    device_entry = next(e for e in manifest['workbooks'] if e['source_tag'] == 'B102-2002')
    control_entry = next(e for e in manifest['workbooks'] if e['source_tag'] == 'B101-2002')
    control_values = workbook_values(Path(control_entry['output_path']))
    controls = defaultdict(set)
    for values in control_values[1:]:
        control = dict(zip(control_values[0], values))
        entity = control['统一社会信用代码'] or control['组织机构代码']
        if control['处理工艺名称']:
            controls[entity].add(control['处理工艺名称'])
    original = workbook_values(Path(device_entry['output_path']))
    headers = original[0]
    outlets = defaultdict(set)
    records = []
    for n, values in enumerate(original[1:], 2):
        ident = by_row[n]
        r = dict(zip(headers, values))
        outlets[ident['pseudonymous_entity_id']].add(r['排放口编号'])
        records.append(dict(identity=ident, row=r, values=values, source=sources[ident['source_id']]))
    eligible = []
    for r in records:
        sid = r['identity']['source_id']
        v = r['row']
        if sid in excluded or r['source']['expected_target'] not in CLASSES:
            continue
        if len(totals[sid]) != 9 or any(t['expected_status'] not in ('calculated', 'not_involved') for t in totals[sid].values()):
            continue
        fuels = []
        for slot in ('一', '二'):
            amount = v[f'燃料{slot}消耗量']
            fuel = v[f'燃料{slot}类型']
            if amount not in (None, '', 0):
                amount = float(amount)
                if amount == 0:
                    continue
                if amount < 0 or fuel not in aliases:
                    break
                fuels.append(dict(slot=slot, amount=amount, fuel=fuel, standard_fuel=aliases[fuel]))
        else:
            if fuels and float(v.get('其他燃料消耗总量（吨标准煤）') or 0) == 0 and all(totals[sid][p]['expected_status'] == 'calculated' for p in ('SO2','NOx','CO','VOC','PM10','PM2.5','BC','OC')):
                r['fuels'] = fuels
                eligible.append(r)
    rng = random.Random(20260929)
    rng.shuffle(eligible)
    used = set()
    entity_uses = Counter()
    selected = []
    for target in CLASSES:
        for root in ROOTS:
            pool = [r for r in eligible if r['source']['expected_target'] == target and r['identity']['source_id'] not in used]
            if root == 'POLLUTANT_PARAMETER_MISSING':
                pool = [r for r in pool if any(f['standard_fuel'] == '煤炭' and
                        r['row'][f"燃料{f['slot']}平均收到基含硫量"] is not None and
                        r['row'][f"燃料{f['slot']}平均收到基灰分（%）"] is not None for f in r['fuels'])]
            if root == 'SOURCE_RELATION_MISSING':
                pool = [r for r in pool if len(outlets[r['identity']['pseudonymous_entity_id']] - {None, ''}) > 1]
            if len(pool) < 30:
                raise ValueError(f'{target}/{root}: {len(pool)} candidates')
            strata_used = Counter()
            for j in range(30):
                def stratum(r):
                    v = r['row']
                    return (r['fuels'][0]['fuel'], v.get('工业锅炉燃烧方式') or v.get('电站锅炉燃烧方式'),
                            tuple(sorted(controls[r['identity']['pseudonymous_entity_id']])))
                r = min(pool, key=lambda r: (strata_used[stratum(r)], entity_uses[r['identity']['pseudonymous_entity_id']]))
                pool.remove(r)
                strata_used[stratum(r)] += 1
                entity_uses[r['identity']['pseudonymous_entity_id']] += 1
                sid = r['identity']['source_id']
                used.add(sid)
                v = r['row']
                fuel = next((f for f in r['fuels'] if f['standard_fuel'] == '煤炭'), r['fuels'][0]) if root == 'POLLUTANT_PARAMETER_MISSING' else r['fuels'][0]
                slot = fuel['slot']
                field, after, affected, detail = '', '', [], 'NO_ANOMALY'
                if root == 'POLLUTANT_PARAMETER_MISSING':
                    sulfur = j % 2 == 0
                    field = f'燃料{slot}平均收到基含硫量' if sulfur else f'燃料{slot}平均收到基灰分（%）'
                    after = None
                    affected = ['SO2'] if sulfur else ['PM10', 'PM2.5', 'BC', 'OC']
                    detail = 'SULFUR_MISSING' if sulfur else 'ASH_MISSING'
                elif root in ('ACTIVITY_MISSING', 'HARD_CONSTRAINT_CONFLICT'):
                    field = f'燃料{slot}消耗量'
                    after = None if root == 'ACTIVITY_MISSING' else -abs(fuel['amount'])
                    affected = ['SO2', 'NOx', 'CO', 'VOC', 'PM10', 'PM2.5', 'BC', 'OC']
                    if fuel['standard_fuel'] == '煤炭' and 'NH3_SCR_SNCR_SLIP' in json.loads(totals[sid]['NH3']['expected_reason_codes']):
                        affected.append('NH3')
                    detail = 'FUEL_AMOUNT_MISSING' if after is None else 'FUEL_AMOUNT_NEGATIVE'
                elif root == 'SOURCE_RELATION_MISSING':
                    field, after = '排放口编号', None
                    affected = list(totals[sid])
                    detail = 'OUTLET_LINK_MISSING'
                selected.append(dict(source_id=sid, source_class=target,
                    entity_id=r['identity']['pseudonymous_entity_id'], original_row=r['identity']['source_row'],
                    root_cause=root, detail_cause=detail, field=field,
                    before=v[field] if field else None, after=after,
                    affected_pollutants=affected, fuel=fuel['fuel'], fuel_slot=slot,
                    combustion_technology=v.get('工业锅炉燃烧方式') or v.get('电站锅炉燃烧方式'),
                    enterprise_control_processes=sorted(controls[r['identity']['pseudonymous_entity_id']]),
                    original_outlet=v['排放口编号']))
    selected.sort(key=lambda r: r['original_row'])
    by_id = {r['source_id']: r for r in selected}
    chosen_records = [r for r in records if r['identity']['source_id'] in by_id]
    selected_entities = {r['entity_id'] for r in selected}
    identity_new = {k: v for k, v in identity.items() if k not in ('records', 'candidate_ids')}
    identity_new['records'] = []
    output_values = [headers]
    for n, r in enumerate(chosen_records, 2):
        label = by_id[r['identity']['source_id']]
        values = list(r['values'])
        if label['field']:
            values[headers.index(label['field'])] = label['after']
        output_values.append(values)
        identity_new['records'].append(r['identity'] | {'source_row': n})
        label['input_row'] = n
    identity_new['candidate_ids'] = [r['source_id'] for r in identity_new['records']]
    dump(HERE / 'inputs/source_identity_index.json', identity_new)
    payloads = []
    for entry in manifest['workbooks']:
        values = output_values if entry['source_tag'] == 'B102-2002' else workbook_values(Path(entry['output_path']))
        if entry['source_tag'] != 'B102-2002':
            h = values[0]
            credit, org = h.index('统一社会信用代码'), h.index('组织机构代码')
            values = [h] + [v for v in values[1:] if v[credit] in selected_entities or v[org] in selected_entities]
        name = Path(entry['output_path']).name
        payload = dict(source_tag=entry['source_tag'], schema_mode='raw_structure', sheet_name='Sheet1',
                       values=values, output_path=str(HERE / 'inputs' / name))
        dest = HERE / 'preparation' / (name + '.json')
        dump(dest, payload)
        payloads.append(dict(payload=str(dest), output=payload['output_path'], rows=len(values)-1, source_tag=entry['source_tag']))
    dump(HERE / 'preparation/workbooks.json', payloads)
    table(HERE / 'adjudication/objects.csv', selected)
    dump(HERE / 'adjudication/objects.json', selected)
    table(HERE / 'adjudication/baseline_pollutants.csv', [t for sid in by_id for t in totals[sid].values()])
    dump(HERE / 'development/objects.json', [r for r in dev if r['plan_class'] != 'preexisting_problem'])
    dump(HERE / 'diagnostics/objects.json', [r for r in dev if r['plan_class'] == 'preexisting_problem'])
    summary = dict(objects=len(selected), entities=len(selected_entities),
                   counts={t: dict(Counter(r['root_cause'] for r in selected if r['source_class']==t)) for t in CLASSES},
                   fuels={t: dict(Counter(r['fuel'] for r in selected if r['source_class']==t)) for t in CLASSES},
                   workbooks=payloads)
    dump(HERE / 'preparation/summary.json', summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def write():
    for entry in json.loads((HERE / 'preparation/workbooks.json').read_text()):
        subprocess.run([str(NODE), str(HERE / 'write_inputs.mjs'), entry['payload']], check=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=('build', 'write'))
    args = parser.parse_args()
    (build if args.action == 'build' else write)()
