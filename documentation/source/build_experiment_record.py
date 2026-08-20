#!/usr/bin/env python3
"""Build the blog-style experiment record from frozen artifacts and evaluations."""

from __future__ import annotations

import csv
import json
import shutil
from collections import Counter
from pathlib import Path
from statistics import median
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/wushuo/Desktop/环境学院论文")
EXPERIMENT = ROOT / "固定燃烧源清单经验缺失处置实验"
OUTPUT = ROOT / "实验记录.docx"
LOCAL_COPY = EXPERIMENT / "documentation" / "实验记录.docx"
SCREENSHOTS = EXPERIMENT / "documentation" / "screenshots"
RESULTS = EXPERIMENT / "实验结果"
SUMMARY = RESULTS / "00_总汇总与索引"
AGGREGATE_PATH = SUMMARY / "aggregate_summary.json"
RUN_SUMMARY_PATH = SUMMARY / "run_summary.csv"
MATRIX_PATH = EXPERIMENT / "matrix" / "frozen_experiment_matrix.json"
COVERAGE_PATH = EXPERIMENT / "rules" / "frozen" / "coverage_frozen.json"
VALIDATION_PATH = EXPERIMENT / "rules" / "frozen" / "frozen_validation.json"
CURRENT_RULE_DIR = EXPERIMENT / "rules" / "frozen" / "v1.0.1"
REFERENCE_VALIDATION_PATH = EXPERIMENT / "reference" / "frozen" / "v1.0.1" / "validation_report.json"
ANOMALY_SUMMARY_PATH = SUMMARY / "异常审计后汇总.json"
ARCHIVE_RECEIPT = ROOT / "论文研究笔记" / "旧实验归档凭证.md"
REPRODUCTION = EXPERIMENT / "reproduction"
FULL_PROMPT = EXPERIMENT / "method_package" / "full_prompt.md"
GENERIC_PROMPT = EXPERIMENT / "method_package" / "generic_agent_prompt.md"
ABLATION_PROMPT = EXPERIMENT / "method_package" / "ablation_prompt.md"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1D2733"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF4EA"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
BORDER = "C9D2DC"
LATIN_FONT = "Songti SC"
BODY_FONT = "Songti SC"
CONTENT_WIDTH_DXA = 9360

METHOD_NAMES = {
    "human_expert": "既有专家主导型流程",
    "expert_led": "既有专家主导型流程",
    "deterministic_program": "确定性程序式流程",
    "generic_tool_agent": "通用工具调用型 Agent",
    "full": "本文完整 Agent 方法",
    "w_o_calculation_item_structure": "w/o 污染物核算项结构",
    "w_o_executable_rule_gate": "w/o 可执行规则门控",
    "agent_managed_admission": "Agent 管理核算准入",
    "w_o_complete_trace": "w/o 完整核算过程记录",
}

TARGET_NAMES = {"INDUSTRIAL": "工业锅炉", "POWER": "火电热力"}


def load_json(path: Path, default: Any = None) -> Any:
    if not path.is_file():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def load_runs() -> list[dict[str, str]]:
    if not RUN_SUMMARY_PATH.is_file():
        return []
    with RUN_SUMMARY_PATH.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def count_csv_rows(path: Path) -> int:
    if not path.is_file():
        return 0
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return sum(1 for _ in csv.DictReader(handle))


def fmt_int(value: Any) -> str:
    try:
        return f"{int(float(value)):,}"
    except (TypeError, ValueError):
        return "—"


def fmt_float(value: Any, digits: int = 2) -> str:
    try:
        return f"{float(value):,.{digits}f}"
    except (TypeError, ValueError):
        return "—"


def fmt_ratio(value: Any) -> str:
    try:
        return f"{float(value) * 100:.1f}%"
    except (TypeError, ValueError):
        return "—"


def fmt_seconds(value: Any) -> str:
    try:
        seconds = float(value)
    except (TypeError, ValueError):
        return "—"
    if seconds < 60:
        return f"{seconds:.1f} s"
    return f"{seconds / 60:.1f} min"


def fmt_memory(value: Any) -> str:
    try:
        return f"{float(value) / (1024 ** 2):.0f} MiB"
    except (TypeError, ValueError):
        return "—"


def set_run_font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = BODY_FONT
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    fonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: Any, *, bold=False, color=INK, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_geometry(table, widths_dxa: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("固定燃烧源排放清单实验  ·  ")
    set_run_font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_body(doc, text: str, *, bold_lead: str | None = None, after=6, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        body = paragraph.add_run(text[len(bold_lead):])
        set_run_font(body, italic=italic)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, italic=italic)
    return paragraph


def add_heading(doc, text: str, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True,
                 color=BLUE if level < 3 else DARK_BLUE)
    paragraph.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[level])
    paragraph.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}[level])
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_bullet(doc, text: str, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(doc, text: str, fill=LIGHT_BLUE, color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_cell_shading(table.cell(0, 0), fill)
    set_cell_text(table.cell(0, 0), text, bold=True, color=color, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code_block(doc, text: str, *, size=7.4):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.right_indent = Inches(0.14)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.02
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    run = paragraph.add_run(text.rstrip() + "\n")
    run.font.name = "Courier New"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Courier New")
    fonts.set(qn("w:hAnsi"), "Courier New")
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    fonts.set(qn("w:cs"), "Courier New")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(INK)
    return paragraph


def add_table(doc, headers: list[str], rows: Iterable[Iterable[Any]], widths: list[int] | None = None,
              font_size=9.2):
    rows = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=font_size)
        set_cell_shading(table.rows[0].cells[index], DARK_BLUE)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    tr_pr.append(repeat_header)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            set_cell_text(cells[col_index], value, size=font_size)
            if row_index % 2:
                set_cell_shading(cells[col_index], LIGHT_GRAY)
    if widths is None:
        base = CONTENT_WIDTH_DXA // len(headers)
        widths = [base] * len(headers)
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_screenshot(doc, filename: str, width=6.35):
    path = SCREENSHOTS / filename
    if not path.is_file():
        return False
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    return True


def table_status(value: str) -> str:
    return {"pass": "通过", "fail": "未通过", "not_applicable": "不适用"}.get(value, value or "—")


aggregate = load_json(AGGREGATE_PATH, {}) or {}
matrix = load_json(MATRIX_PATH, {}) or {}
coverage = load_json(COVERAGE_PATH, {}) or {}
validation = load_json(VALIDATION_PATH, {}) or {}
reference_validation = load_json(REFERENCE_VALIDATION_PATH, {}) or {}
anomaly_summary = load_json(ANOMALY_SUMMARY_PATH, {}) or {}
runs = load_runs()

rule_index = load_json(CURRENT_RULE_DIR / "rule_index.json", []) or []
rule_test_count = sum(
    1 for line in (CURRENT_RULE_DIR / "tests" / "rule_tests.jsonl").read_text(encoding="utf-8").splitlines()
    if line.strip()
)
parameter_count = sum(count_csv_rows(path) for path in (CURRENT_RULE_DIR / "parameters").glob("*.csv"))
mapping_count = sum(count_csv_rows(path) for path in (CURRENT_RULE_DIR / "mappings").glob("*.csv"))

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.82)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.42)
section.footer_distance = Inches(0.42)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
normal._element.rPr.rFonts.set(qn("w:cs"), LATIN_FONT)
normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
    style = doc.styles[name]
    style.font.name = BODY_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style._element.rPr.rFonts.set(qn("w:cs"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    style.font.size = Pt(11)

for level, size, color, before, after in (
    (1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, DARK_BLUE, 10, 5)
):
    style = doc.styles[f"Heading {level}"]
    style.font.name = BODY_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style._element.rPr.rFonts.set(qn("w:cs"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(header.add_run("固定燃烧源排放清单实验记录"), size=9, color=MUTED)
add_page_number(section.footer.paragraphs[0])

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(4)
set_run_font(title.add_run("固定燃烧源排放清单实验记录"), size=24, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
set_run_font(subtitle.add_run("从技术指南规则整理到全量清单编制与独立评价"), size=13, color=MUTED)

evaluable = aggregate.get("machine_runs_evaluable", aggregate.get("machine_runs_sealed", 0))
total = aggregate.get("machine_runs_total", 136)
completed_seals = aggregate.get("experiment_seals_found", 0)
add_table(doc, ["项目", "现行设置"], [
    ("研究对象", "4,025条工业锅炉记录、367条火电热力记录；两类分别运行和报告。"),
    ("标准边界", "T/CSES 144—2024通用方法、第7章、第8.3节、第15章，以及附录A、C、D、E.1和E.7。"),
    ("方法定位", "Agent负责组织阶段、调用受限能力和汇总异常；规则匹配、核算准入、数值计算与质量检查由确定性工具执行。"),
    ("实现配置", "Hermes作为Agent harness，DeepSeek作为推理模型；二者是技术实现而非论文核心主张。"),
    ("正式运行", f"{fmt_int(completed_seals)}/{fmt_int(total)}次机器运行已结束并封存且全部可评价；14份人工结果均已接入并可评测。"),
], [1900, 7460], font_size=9.8)

add_callout(doc, "旧的受影响运行包和Hermes会话已先完成哈希归档，并移出当前结果树；当前“实验结果”只保留异常审计修复后的实体运行包，不使用软链接。")

add_heading(doc, "1 先把核算边界说清楚", 1)
add_body(doc, "本实验不建立CEMS核算路径。基102中的污染物产生量和排放量保留为比较信息，但不参与方法选择，不用于反推治理效率，也不作为最终排放量。治理前产生量统一采用适用的排放因子法；燃煤SO₂、颗粒物、BC和OC在具备必要参数时采用物料衡算。活动水平、因子匹配和产生量在结果工作簿中保留Excel公式。")
add_body(doc, "治理信息的经验处置。能够映射的工艺采用T/CSES表A.1的缺省效率；治理信息缺失或工艺无法映射时，局部污染物按效率0继续计算并保留相应标记；没有适用治理措施时效率为0。该结果只表示在现有资料条件下采用了经验缺失处置，不能解释为设施真实不存在，也不能称为实测去除效率。", bold_lead="治理信息的经验处置。")
add_body(doc, "基础核算输入的阻断。只有治理信息缺失允许按经验效率0继续计算。燃料量、燃料类型、硫分、灰分、装机容量或适用产生因子等必要输入缺失时，受影响的污染物结果留空并进入异常清单；不得把信息不足写成0，也不得扩大到无关污染物。", bold_lead="基础核算输入的阻断。")
add_body(doc, "污染物专用处理。PM₁₀分为PM₂.₅细颗粒部分和PM₂.₅—PM₁₀粗颗粒部分分别应用效率；BC、OC使用各自专用效率。NH₃仅按SCR/SNCR脱硝过程的氨逃逸单独计算，SCR和SNCR分别采用0.16和0.17 g/kg煤，不混入一般燃料燃烧因子。", bold_lead="污染物专用处理。")

add_heading(doc, "2 为什么先归档再重做", 1)
add_body(doc, "上一版实验采用严格阻断方案，和本次经验缺失处置口径不一致。如果直接复用旧Profile、旧会话或旧规则包，旧结论可能通过会话记忆、运行目录或结果文件回到新实验中。因此先把旧实验作为一个完整、只读的历史版本保存，再把受影响运行和会话移出当前结果树。")
add_table(doc, ["归档检查", "结果"], [
    ("私有仓库", "github.com/glittering-universe/fixed-combustion-inventory-legacy-archive"),
    ("归档提交", "0b5084149941cf41fcff1a530065146a25675d71"),
    ("远端回读", "重新克隆后逐项校验通过"),
    ("归档范围", "1,601个跟踪文件、1,597项哈希清单、20个旧Profile目录和152份旧会话记录"),
    ("本地清理", "归档校验通过后删除旧实验目录和20个旧实验Profile"),
    ("本轮异常审计快照", "Git标签pre-anomaly-audit-fix-v1；受影响运行包按实体文件归档，当前结果树不含软链接"),
], [2500, 6860])
add_body(doc, "新实验建立在“固定燃烧源清单经验缺失处置实验”目录下，只使用Desktop可见的fixed-combustion-inventory Profile。阶段R、完整方法、机器基线与消融版本共用该Profile以统一模型和应用环境；每次运行仍创建新会话、新运行包和新封存清单，不能续接前一次会话。")

add_heading(doc, "3 系统按七个阶段推进", 1)
add_body(doc, "系统以一个顶层Plugin装载七个阶段Skill和八个批量Tool。Skill规定当前阶段的处理步骤、输入输出和禁止事项；Tool执行表格适配、关系构建、规则匹配、准入、计算、检查和记录。每个阶段一次批量处理本轮全部记录，Agent只接收阶段摘要，不把数万条核算项放入模型上下文。")
add_table(doc, ["阶段", "批量处理内容", "消融属性"], [
    ("数据整理与设备源构建", "适配工作簿；关联企业、设备、燃料、排放口和治理信息", "共享"),
    ("污染物排放核算项生成", "形成燃料×污染物核算项、PM粒径依赖和NH₃过程项", "可消融"),
    ("T/CSES核算规则实施", "确定方法、参数、单位、适用条件和标准位置", "可消融"),
    ("污染物核算准入", "给出已核算、本源不涉及、信息不足、基础条件异常四种状态", "可消融"),
    ("固定燃烧排放量计算", "执行活动水平、物料衡算、产生量和治理后排放计算", "共享"),
    ("质量复核与结果生成", "检查范围、非负性、PM关系、汇总和输出结构", "共享"),
    ("完整核算过程记录归档", "记录输入字段、规则参数、计算步骤、中间量、状态和输出位置", "可消融"),
], [2250, 5300, 1810], font_size=8.8)
add_body(doc, "每个Tool只写入自己的阶段数据表，后续Tool发现问题时只能报告、标记或请求拥有该数据表的上游阶段重算，不能静默改写。消融版本使用同一接口的替代实现，且禁止其他阶段把已去掉的能力重新补回。")

add_heading(doc, "4 阶段R：先形成可执行规则包", 1)
add_body(doc, "阶段R只向Agent提供原始T/CSES PDF、冻结章节范围和空白规则包规范。Agent从PDF形成候选规则、参数、标准原词映射和测试；候选包先通过自动检查，再经过专业复核和修订，最后才冻结给实验A—D使用。该阶段不比较Agent与专家谁抽取得更好，也不进入主方法排名。")
add_screenshot(doc, "hermes-stage-r-completed.png")
add_body(doc, "在Hermes会话中，候选规则包完成了范围抽取、结构校验和覆盖统计准备，但候选状态不会由Agent自行改成“通过”。专业复核后的冻结包由独立锁文件控制；正式计算引擎只接受通过复核且哈希一致的版本。")

counts = {
    "rules": len(rule_index),
    "parameters": parameter_count,
    "mappings": mapping_count,
    "tests": rule_test_count,
}
standard_cov = coverage.get("standard_extraction", {})
operational_cov = coverage.get("operational_readiness", {})
add_table(doc, ["冻结内容", "数量或结论"], [
    ("规则", f"{fmt_int(counts.get('rules'))}条"),
    ("参数", f"{fmt_int(counts.get('parameters'))}行"),
    ("映射", f"{fmt_int(counts.get('mappings'))}行"),
    ("测试", f"{fmt_int(counts.get('tests'))}条"),
    ("自动校验", f"errors={len(validation.get('errors', []))}，warnings={len(validation.get('warnings', []))}"),
    ("专业复核", "通过" if validation.get("professional_review_status") == "passed" else "待完成"),
    ("标准抽取覆盖", f"{fmt_int(standard_cov.get('matched_count'))}/{fmt_int(standard_cov.get('required_count'))}，{fmt_ratio(standard_cov.get('coverage_ratio'))}"),
    ("运行口径就绪", f"{fmt_int(operational_cov.get('matched_count'))}/{fmt_int(operational_cov.get('required_count'))}，{fmt_ratio(operational_cov.get('coverage_ratio'))}"),
], [2800, 6560])
add_body(doc, "标准抽取覆盖率只回答候选包是否覆盖了独立列出的T/CSES必要规则；运行口径就绪度还检查本次经验缺失处置、Excel公式、NH₃、异常清单等正式执行要求。后者是冻结前的缺口检查，不是主实验评分。")

add_heading(doc, "5 正式实验如何保持彼此独立", 1)
add_body(doc, "正式矩阵在运行前一次性冻结，包含136次机器运行和14份人工结果。每一行都预先确定实验、方法、源类别、规模、输入版本、重复编号和执行顺序；启动后不依据结果修改。每次机器运行包保存输入哈希、方法包哈希、规则包哈希、阶段修订、模型用量、资源记录和最终封存清单；人工原文件保持字节不变，另生成统一评价包。")
add_body(doc, "各Agent方法采用相同的研究者配置公共系统提示。Hermes根据运行工作区状态自动附加通用运行时指令；该平台级指令不包含固定燃烧源领域规则、核算参数或参照结果，未作为实验变量。约32行coding-agent通用指令属于Hermes harness，不作为方法模块、消融对象或评价指标。")
add_body(doc, "专业核算引擎由2.0.0修订为2.1.0。受该引擎影响的94次正式运行均使用2.1.0和规则包v1.0.1；旧版结果已归档且不参与统计，Full、确定性程序和消融版本之间不存在引擎版本不一致。本次只修正独立评价器并重评封存结果，无需重新运行其他核算实验。")
add_screenshot(doc, "hermes-formal-profile-and-sessions.png")
add_body(doc, "Hermes Desktop可以通过固定实验Profile查看阶段R和正式运行会话。CLI与Desktop连接同一HERMES_HOME和Profile，因此批量运行产生的会话仍可在Desktop中回看；不同Profile之间不会共享会话。")
add_table(doc, ["实验", "目的", "正式范围"], [
    ("A", "比较四种完整编制流程", "机器流程：两类全量×3次；人工流程：两类各1份实测结果"),
    ("B", "检查输入适应性与缺失/冲突处置", "B0与S1—S4为两类全量；B2为30条受控注入"),
    ("C", "观察规模扩大时的执行表现", "25%、50%、75%、100%嵌套规模；两类数据×3次"),
    ("D", "观察四个方法组成的作用", "Full与四个单因素替代版本；两类全量；各1次"),
], [1100, 3300, 4960])

add_heading(doc, "6 一次完整方法运行是什么样的", 1)
add_body(doc, "完整方法运行从一个只含本次输入和冻结配置的运行包开始。Agent按七个阶段Skill调用八个必要批量Tool，任何基础核算输入缺失都只阻断受影响的污染物核算项；治理信息缺失则按经验效率0继续并保留标记。评价边界不再要求工具调用总数正好为8，而是检查必要阶段完整且顺序正确、未使用消融替代能力、未跨方法读取结果且未绕过规定阶段。额外的Skill查看和工具发现只记录为探索开销，不判失败，也不进入EICPI。")
add_screenshot(doc, "hermes-formal-full-industrial.png")
first_full_industrial = (aggregate.get("experiment_A_raw", {}).get("full|INDUSTRIAL", {}).get("runs") or [{}])[0]
full_counts = first_full_industrial.get("status_counts", {})
add_body(doc, f"首个工业锅炉全量运行处理4,025条源记录，形成36,225条源×污染物结果，完整核算过程记录覆盖率为100%；其中{fmt_int(full_counts.get('calculated'))}条已核算、{fmt_int(full_counts.get('not_involved'))}条本源不涉及、{fmt_int(full_counts.get('information_insufficient'))}条信息不足。独立评价器修正后，仅对allow_calculation核算项检查活动水平、因子和数值；not_applicable项只检查处置理由及未计入源级汇总。由此该轮结构与最小复算检查均通过。SRC-PWR-000264的8个范围外生物燃料项已作为固定回归用例。")
add_body(doc, "重新评价后，44次Full运行的方法边界检查和独立复算检查全部通过。此前4次因额外Skill查看或工具发现被判为边界失败的运行，在阶段完整性与顺序、禁止能力、跨方法读取和阶段绕过四项检查下均改判为通过；这些探索调用保留在会话记录中，但不进入EICPI。")

add_heading(doc, "7 通用Agent基线不会共享专用方法", 1)
add_body(doc, "通用工具调用型Agent只接收脱敏输入工作簿和T/CSES PDF，可使用通用文件、终端、代码和表格能力，但不能读取专用Plugin、冻结规则包、输入适配配置、其他方法输出或历史答案。它必须自己理解表格、选择方法、编写处理逻辑并生成统一CSV；这是一条完整基线，不是故意设置的失败版本。")
add_screenshot(doc, "hermes-formal-generic-industrial.png")
add_body(doc, "首个工业锅炉全量基线会话自行完成了输出，但独立评价发现其结构检查和最小复算均未通过：30,926条结果把空白控制效率写入“已核算”记录，另有128条产生量和56条排放量不能通过最小复算。这个差异来自统一评价程序，而不是从Agent自己的完成报告推断。")

add_heading(doc, "8 实验A：四种完整流程的全量比较", 1)
add_body(doc, "实验A包含既有专家主导型流程、确定性程序式流程、通用工具调用型Agent和本文完整Agent方法。三种机器流程分别对工业锅炉和火电热力全量运行3次；人工流程的14份结果已全部接入，原文件原样保留，并按统一字段规范生成只读评价包。人工耗时采用实测值11.52秒/条。人工结果参加结果正确性、异常处置、过程可复核性、重复性和效率评价；其缺少字段来源、参数来源和中间过程会降低T维度，只有Skill调用数、工具次数等Hermes专属观测量标记为不适用。")

a_rows = [row for row in runs if row.get("experiment") == "A"]
if a_rows:
    grouped: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in a_rows:
        grouped.setdefault((row["method"], row["target"]), []).append(row)
    table_rows = []
    for (method, target), items in sorted(grouped.items()):
        times = [float(item["wall_seconds"]) for item in items if item.get("wall_seconds")]
        structural_pass = sum(item.get("structural_status") == "pass" for item in items)
        recalc_pass = sum(item.get("minimum_recalculation_gate") == "pass" for item in items)
        table_rows.append((
            METHOD_NAMES.get(method, method), TARGET_NAMES.get(target, target), len(items),
            f"{structural_pass}/{len(items)}", f"{recalc_pass}/{len(items)}",
            fmt_seconds(median(times)) if times else "—",
        ))
    add_table(doc, ["方法", "源类别", "已封存", "结构通过", "最小复算通过", "运行用时中位数"],
              table_rows, [2400, 1200, 1000, 1350, 1500, 1910], font_size=8.6)
else:
    add_callout(doc, "实验A尚无封存结果。", fill=PALE_YELLOW, color=INK)

if aggregate.get("repeatability_R"):
    repeat_rows = []
    for key, data in aggregate["repeatability_R"].items():
        method, target = key.split("|", 1)
        repeat_rows.append((METHOD_NAMES.get(method, method), TARGET_NAMES.get(target, target),
                            data.get("pair_count", 0), fmt_ratio(data.get("agreement_rate"))))
    add_table(doc, ["方法", "源类别", "重复对数", "逐项一致率"], repeat_rows,
              [3100, 1800, 1700, 2760])
add_body(doc, "独立参照包v1.0.1已通过完整性与复算检查，并被逐运行评价程序只读调用。该参照包在机器运行封存后形成，虽然与各参评输出隔离，但不表述为预注册参照；标准规则落实、数值有效性、人工判断负担和过程完整性均以逐运行报告为准。人工A与人工B0的两份未扰动作业构成人工重复性配对。")

add_heading(doc, "9 实验B：输入结构变化与受控异常", 1)
add_body(doc, "B1先为每种完整流程封存未扰动B0，再在全量数据上运行四种等价结构变化：S1改变括号、空格和表头换行；S2替换为已登记的表头别名；S3打乱列顺序并增加封面工作表；S4同时组合前三种变化。比较时以同一方法、同一源类别的B0结果为基准，检查扰动前后语义是否保持。")
add_body(doc, "B1的比较分母按每个方法自身的B0与扰动输出键并集计算，因此用于判断同一方法在扰动前后的保持情况，不用于横向比较不同方法生成的绝对单元数。通用Agent在POWER的S3和S4中各额外生成367个VOCs行，故其总比较单元为158,846，而确定性程序和Full为158,112。缺失、多余或取值变化的单元均判为未保持。")
b2_validity = reference_validation.get("checks", {}).get("b2_injection_validity", {})
add_body(doc, f"B2原计划在30条受控记录中注入缺失或冲突，工业锅炉和火电热力各15条，覆盖活动水平缺失、污染物参数缺失、源关系冲突、硬约束冲突和治理缺失的唯一缺省处置。独立参照复核确认其中{fmt_int(b2_validity.get('valid'))}条是有效新注入，{fmt_int(b2_validity.get('invalid'))}条对象在注入前已存在同类问题；评价时二者明确分开，不把原有问题计为注入成功。")
if aggregate.get("b1_semantic_retention_E1"):
    rows_b1 = []
    for method, data in aggregate["b1_semantic_retention_E1"].items():
        rows_b1.append((METHOD_NAMES.get(method, method), data.get("pair_count", 0),
                        fmt_int(data.get("total_pair_units")), fmt_ratio(data.get("agreement_rate"))))
    add_table(doc, ["方法", "B0—扰动配对数", "比较单元", "语义保持率"], rows_b1,
              [3000, 2000, 1800, 2560])
else:
    add_callout(doc, "实验B的全量扰动结果尚未全部封存。", fill=PALE_YELLOW, color=INK)

b2 = aggregate.get("experiment_B2_raw", {})
if b2.get("run_count"):
    b2_failures = [item for item in aggregate.get("machine_evaluation_failures", []) if item.get("experiment") == "B" and item.get("input_variant") == "B2"]
    machine_b2_evaluable = sum(item.get("method") != "expert_led" for item in b2.get("runs", []))
    human_b2_evaluable = sum(item.get("method") == "expert_led" for item in b2.get("runs", []))
    add_body(doc, f"B2共有6次机器任务和2份人工结果：6次机器任务均已封存且{fmt_int(machine_b2_evaluable)}次全部形成可评测输出，任务失败为{fmt_int(len(b2_failures))}次；{fmt_int(human_b2_evaluable)}份人工结果均可评测。有效注入、最小影响和异常根因评价由独立参照包给出，不用参评方法自身的异常数量充当正确答案。")
    add_table(doc, ["方法", "目标检测（工业/火电）", "处置符合（工业/火电）", "根因归并（工业/火电）", "未受影响项正确（工业/火电）"], [
        ("既有专家主导型", "0% / 0%", "30% / 0%", "0% / 0%", "0% / 0%"),
        ("确定性程序式", "50% / 60%", "100% / 100%", "50% / 60%", "100% / 90%"),
        ("通用工具调用型Agent", "20% / 10%", "20% / 70%", "20% / 10%", "20% / 30%"),
        ("本文Full", "50% / 60%", "100% / 100%", "50% / 60%", "100% / 90%"),
    ], [2100, 1900, 1800, 1700, 1860], font_size=8.0)

add_heading(doc, "10 实验C：25%到100%的嵌套规模", 1)
add_body(doc, "实验C不预先把记录分成简单、常规或复杂等级，而是从同一冻结排序中构造25%、50%、75%和100%嵌套子集。三种机器流程在两个源类别的四个规模上各运行3次。记录的wall time、CPU time、峰值内存、输出大小、模型Token和调用次数均作为规模响应的原始观测；EICPI的时间分量不使用实验C，而统一取实验A全量场景的wall time除以该源类记录数。")
add_body(doc, "实验A与实验C分别创建独立Hermes会话，系统Prompt包含不同的实验元数据，且运行顺序、缓存和并发环境不作跨实验配对控制。因此，同为100%规模的用时只在各实验内部解释，不把A与C的绝对秒数合并或用于显著性判断。MiB按1,048,576字节换算。")
c = aggregate.get("experiment_C_scaling_raw", {})
if c:
    c_rows = []
    for data in c.values():
        c_rows.append((
            METHOD_NAMES.get(data.get("method"), data.get("method")),
            TARGET_NAMES.get(data.get("target"), data.get("target")),
            f"{data.get('scale_percent')}%", data.get("run_count", 0),
            fmt_seconds(data.get("wall_seconds", {}).get("median")),
            fmt_memory(data.get("peak_rss_bytes", {}).get("median")),
            fmt_float(data.get("api_calls", {}).get("median"), 0),
        ))
    c_rows.sort(key=lambda row: (row[0], row[1], int(row[2].rstrip("%"))))
    add_table(doc, ["方法", "源类别", "规模", "重复", "用时中位数", "峰值内存中位数（MiB）", "API调用中位数"],
              c_rows, [2250, 1050, 750, 650, 1450, 1750, 1460], font_size=8.0)
else:
    add_callout(doc, "实验C尚未形成可汇总的封存结果。", fill=PALE_YELLOW, color=INK)

add_heading(doc, "11 实验D：只替换一个方法组成", 1)
add_body(doc, "实验D不再重复完整流程对比。Full与四个单因素替代版本均处理两类全量数据，各运行1次；其余Plugin、共享Skill、共享Tool、输入、规则包和模型配置保持一致。每个消融只按四项质量门槛、相关子指标和定向诊断与同源类别Full比较，不计算消融总EICPI。")
add_table(doc, ["版本", "唯一替换", "重点观察"], [
    ("Full", "无", "完整方法基准"),
    ("w/o 污染物核算项结构", "改用设备×污染物槽位，不先构造逐燃料/过程核算项", "多燃料关系、PM链条、NH₃过程项"),
    ("w/o 可执行规则门控", "由Agent记录方法和参数选择", "规则合规、参数选择、错误通过"),
    ("Agent 管理核算准入", "由Agent记录准入决定", "信息不足处置、阻断范围、异常清单"),
    ("w/o 完整核算过程记录", "只保留最小过程摘要", "输入、规则参数、计算步骤和输出位置的可复核性"),
], [2250, 3650, 3460], font_size=8.5)
d = aggregate.get("experiment_D_ablation_raw", {})
if d.get("runs"):
    d_rows = []
    for item in d["runs"]:
        d_rows.append((
            METHOD_NAMES.get(item.get("method"), item.get("method")),
            TARGET_NAMES.get(item.get("target"), item.get("target")),
            table_status(item.get("structural_status")),
            table_status(item.get("minimum_recalculation_gate")),
            fmt_ratio(item.get("complete_process_record_coverage")),
            fmt_int(item.get("status_counts", {}).get("source_data_invalid", 0)),
        ))
    add_table(doc, ["版本", "源类别", "结构检查", "最小复算", "过程记录覆盖", "基础异常"],
              d_rows, [2600, 1100, 1200, 1350, 1600, 1510], font_size=8.3)
else:
    add_callout(doc, "实验D尚未形成封存结果。", fill=PALE_YELLOW, color=INK)
add_screenshot(doc, "hermes-formal-experiments-completed.png")
add_body(doc, "实验D结束后，Hermes Desktop中的fixed-combustion-inventory Profile保留了完整方法与四个单因素替代版本的独立会话。截图未显示鼠标指针；会话标题直接包含实验、版本和源类别，后续可从Desktop进入对应会话，并与运行包中的脱敏JSONL导出互相核对。")

add_heading(doc, "12 独立评价不替参评方法补答案", 1)
add_body(doc, "实验控制程序只负责准备输入、选择范围、启动运行、计量资源、处理一次允许的系统故障重试和封存文件；独立评价程序只读封存结果。两者都不参与燃料映射、规则选择、治理效率决策或排放计算，也不会在评价阶段修正参评输出。评价器v1.2.0与聚合程序v1.3.0对136次机器和14次人工封存运行完成最终评价；核算工作簿、Hermes会话、耗时与封存清单保持不变。早期人工包曾把规范化CSV和独立评价报告列入封存清单；v1.2.0将这两项明确标为可再生评价产物，仅校验人工原文件、说明、输入、耗时及其余封存内容，避免重评报告对自身哈希造成假失败。")
add_table(doc, ["当前外部输入", "处理方式"], [
    ("人工基线实测", "14份结果已接入并可评测；原文件保留，统一评价包单独生成；人工耗时为11.52秒/条"),
    ("专业复核参照包", "v1.0.1已接入并通过验证；运行输出隔离，但形成于机器运行封存之后"),
    ("分源类计时阈值", "工业锅炉L=0.0224、U=0.2981 s/条；火电热力L=0.1226、U=2.4523 s/条"),
    ("EICPI", "按最终冻结阈值计算；未通过全部质量门槛的方法只作诊断，不进入可接受方法排序"),
], [2850, 6510])
add_body(doc, "排放清单编制综合效能指数EICPI用于结果摘要，不是论文方法创新。时间效率E2按实验A全量wall time/源记录数计分：t≤L得100分，t≥U得0分，中间线性插值。工业锅炉与火电热力分别设阈值，因为火电热力仅367条，Agent初始化和固定工作流开销占比更高；若强行共用工业锅炉阈值，会系统性不利于小规模源类。该阈值由研究者在最终评价阶段提供并冻结，不表述为机器运行前预注册。E=(B1语义保持得分E1+时间得分E2)/2；实验D不计算总EICPI。")

eicpi = aggregate.get("EICPI", {})
target_scores = eicpi.get("by_method_target", {})
if target_scores:
    score_rows = []
    for target in ("INDUSTRIAL", "POWER"):
        for method in ("expert_led", "deterministic_program", "generic_tool_agent", "full"):
            score = target_scores[f"{method}|{target}"]
            timing = score["timing"]
            score_rows.append((
                METHOD_NAMES[method], TARGET_NAMES[target],
                f"{timing['seconds_per_record']:.4f}",
                fmt_ratio(score["E1_semantic_retention"]),
                fmt_ratio(score["E2_time_score"]),
                fmt_ratio(score["dimensions"]["E"]),
                f"{score['EICPI_0_100']:.2f}",
                "仅诊断" if not score["all_quality_gates_pass"] else "可进入排序",
            ))
    add_table(doc,
              ["方法", "源类别", "s/条", "E1", "E2", "E", "EICPI", "质量门槛"],
              score_rows, [2100, 1050, 850, 850, 850, 850, 950, 1010], font_size=7.7)

    macro_rows = []
    macro = eicpi.get("by_method_macro_average", {})
    for method in ("expert_led", "deterministic_program", "generic_tool_agent", "full"):
        score = macro[method]
        dims = score["dimensions"]
        macro_rows.append((
            METHOD_NAMES[method], fmt_ratio(dims["S"]), fmt_ratio(dims["A"]),
            fmt_ratio(dims["Q"]), fmt_ratio(dims["T"]), fmt_ratio(dims["R"]),
            fmt_ratio(dims["E"]), f"{score['EICPI_0_100']:.2f}",
            "仅诊断" if not score["all_quality_gates_pass"] else "可进入排序",
        ))
    add_table(doc, ["方法", "S", "A", "Q", "T", "R", "E", "宏平均EICPI", "结论"],
              macro_rows, [2200, 720, 720, 720, 720, 720, 720, 1270, 870], font_size=7.7)
    add_body(doc, "两个源类等权宏平均EICPI为：确定性程序93.29、Full 93.18、既有专家主导型37.74、通用Agent 32.31。四种方法均未在两个源类同时通过全部不可补偿质量门槛，因此上述分数只用于描述性能构成，不构成可接受方法之间的正式名次；本次最终评价没有可进入总体排序的方法。")

add_heading(doc, "12.1 异常审计后的核算边界", 2)
pre_fix = anomaly_summary.get("pre_fix_audit", {})
post_fix = anomaly_summary.get("post_fix", {})
add_body(doc, f"异常审计从{fmt_int(pre_fix.get('source_level_records'))}条源级记录出发。统一部门名称后，258条工业燃煤候选均能够进入附录C参数匹配；其中257条恢复核算，1条因独立存在的“其他”燃烧技术歧义继续阻断。增加受控的“不分技术”因子回退后，2条火电煤矸石记录恢复核算。最终异常清单包含{fmt_int(post_fix.get('unique_sources'))}个唯一源、{fmt_int(post_fix.get('source_reason_groups'))}个源—根因组合，其中6个源同时具有两个根因。")
add_table(doc, ["异常类别", "源—根因组合", "处置"], [
    (item.get("reason_label"), fmt_int(item.get("source_reason_groups")), "范围外" if item.get("disposition") == "not_applicable" else "保留阻断")
    for item in post_fix.get("reason_counts", [])
], [5700, 1600, 2060], font_size=8.5)

add_heading(doc, "13 不需要重新协调的复现实验入口", 1)
add_body(doc, "复现实验入口与历史封存包分开保存，不改写既有运行结果。运行前检查器会只读核对冻结设置锁、输入哈希、规则包、Plugin、Profile、三套Prompt、标准原文、矩阵行数，以及启动器实际调用的执行和评价脚本；任何一项不一致都会停止。主启动器既可以继续当前目录中尚未完成的运行，也可以复制出一个不含历史结果的新目录，按A、B、C、D顺序完成136次机器运行并自动汇总。")
add_table(doc, ["文件", "用途"], [
    ("reproduction/check_environment.py", "只读核对正式设置、运行时、Profile、Plugin、规则、输入、Prompt和矩阵"),
    ("reproduction/run_all_machine_experiments.sh", "继续当前冻结矩阵，或在空目录中建立完整复现实验"),
    ("reproduction/run_one_machine_experiment.sh", "按冻结RUN_ID单独执行一行矩阵"),
    ("reproduction/aggregate_results.sh", "只读汇总封存结果；失败任务不补跑、不回填"),
    ("reproduction/reproduction_manifest.json", "记录启动器、Prompt、执行/评价依赖、矩阵、设置锁和标准原文的SHA-256"),
], [3600, 5760], font_size=8.8)
add_body(doc, "在现有实验目录继续未完成项：")
add_code_block(doc, "./reproduction/run_all_machine_experiments.sh --resume")
add_body(doc, "在一个空目录中重新执行完整机器矩阵：")
add_code_block(doc, "./reproduction/run_all_machine_experiments.sh \\\n+  --fresh-dir /Users/wushuo/Desktop/环境学院论文/固定燃烧源清单复现实验")
add_body(doc, "单独复现一行矩阵并重新汇总：")
add_code_block(doc, "./reproduction/run_one_machine_experiment.sh C-100-IND-FULL-R1 \\\n+  /Users/wushuo/Desktop/环境学院论文/固定燃烧源清单复现实验\n\n./reproduction/aggregate_results.sh \\\n+  /Users/wushuo/Desktop/环境学院论文/固定燃烧源清单复现实验")
add_callout(doc, "复现脚本不会重新生成或改写14份人工原文件和独立参照包；它只复制、校验并规范化当前已提供的数据。分源类计时阈值已冻结在evaluation/final_time_thresholds.json，并由环境检查与复现清单共同校验。")

add_heading(doc, "14 正式启动脚本快照", 1)
add_body(doc, "以下脚本均为本记录生成时已经通过语法检查、依赖哈希检查和环境检查的复现入口。主启动器拒绝覆盖非空的新目标目录；当前目录中的封存运行会跳过，既有但未封存的运行包会阻断。")
add_heading(doc, "14.1 完整机器矩阵启动器", 2)
add_code_block(doc, (REPRODUCTION / "run_all_machine_experiments.sh").read_text(encoding="utf-8"), size=7.0)
add_heading(doc, "14.2 单次矩阵任务启动器", 2)
add_code_block(doc, (REPRODUCTION / "run_one_machine_experiment.sh").read_text(encoding="utf-8"), size=7.0)
add_heading(doc, "14.3 独立汇总入口", 2)
add_code_block(doc, (REPRODUCTION / "aggregate_results.sh").read_text(encoding="utf-8"), size=7.0)

add_heading(doc, "15 正式Prompt快照", 1)
add_body(doc, "Hermes单会话启动程序根据矩阵runner选择下列三套冻结Prompt，并在末尾追加run_id、运行包路径、scenario、方法包哈希、目标类型、规模和输入版本。模型固定为deepseek-v4-pro，Provider为custom，推理强度为high。各Agent方法共享研究者配置公共系统提示；Hermes自动附加的通用运行时指令属于harness且不含领域规则、核算参数或参照结果。完整方法和消融仅开放专用固定燃烧Plugin与Skills，通用Agent只开放通用终端、文件、代码执行与Skills。")
add_heading(doc, "15.1 完整Agent方法Prompt", 2)
add_code_block(doc, FULL_PROMPT.read_text(encoding="utf-8"), size=7.2)
add_heading(doc, "15.2 通用工具调用型Agent Prompt", 2)
add_code_block(doc, GENERIC_PROMPT.read_text(encoding="utf-8"), size=7.2)
add_heading(doc, "15.3 消融实验Prompt", 2)
add_code_block(doc, ABLATION_PROMPT.read_text(encoding="utf-8"), size=7.0)
add_body(doc, "三套Prompt的SHA-256分别为：full=a20052862561a1687eaa4a3552d4116ff2290496cf4ae99b22db332533203c7e；generic=a1ce743cc8da048abe83c3e95f4daa5060fa87147a81e75b2e15f557d5b23124；ablation=870722aded631f707d1bb8546ec6725d102c84a2a10aadeb6f9a50ce6726fbd1。正式设置锁SHA-256为95696d96fdd119130eed5a1357465b4badc4041415f3f79117071cbc7655d619。")

add_heading(doc, "16 到哪里查看和复核", 1)
add_table(doc, ["内容", "位置"], [
    ("正式运行包", "固定燃烧源清单经验缺失处置实验/实验结果/<实验>/<具体内容>/<源类>/<方法>/<重复编号>"),
    ("冻结实验矩阵", "固定燃烧源清单经验缺失处置实验/matrix/frozen_experiment_matrix.json"),
    ("规则包", "固定燃烧源清单经验缺失处置实验/rules/frozen/v1.0.1"),
    ("独立参照包", "固定燃烧源清单经验缺失处置实验/reference/frozen/v1.0.1"),
    ("方法包", "固定燃烧源清单经验缺失处置实验/method_package"),
    ("独立评价与汇总", "固定燃烧源清单经验缺失处置实验/实验结果/00_总汇总与索引"),
    ("复现实验入口", "固定燃烧源清单经验缺失处置实验/reproduction"),
    ("截图", "固定燃烧源清单经验缺失处置实验/documentation/screenshots"),
    ("归档凭证", "论文研究笔记/旧实验归档凭证.md"),
], [2700, 6660], font_size=9.3)

pending = aggregate.get("machine_runs_pending", total)
failures = aggregate.get("machine_evaluation_failures", [])
if completed_seals == total and not failures:
    add_callout(doc, f"{fmt_int(total)}次机器运行均已封存并进入独立评价；评价程序未记录任务输出失败。", fill=PALE_GREEN, color="2F6B3A")
elif completed_seals == total and failures:
    failure_ids = "、".join(item.get("run_id", "未知运行") for item in failures)
    add_callout(doc, f"{fmt_int(total)}次机器运行均已结束并封存；{fmt_int(evaluable)}次形成可评价输出。{failure_ids}未生成受支持输出，作为正式任务失败保留，未补跑、未回填。", fill=PALE_YELLOW, color=INK)
else:
    add_callout(doc, f"当前已封存{fmt_int(completed_seals)}/{fmt_int(total)}次机器运行，尚有{fmt_int(pending)}次等待完成。", fill=PALE_YELLOW, color=INK)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
LOCAL_COPY.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
shutil.copyfile(OUTPUT, LOCAL_COPY)
print(json.dumps({"output": str(OUTPUT), "local_copy": str(LOCAL_COPY), "completed_seals": completed_seals, "evaluable": evaluable, "total": total}, ensure_ascii=False))
