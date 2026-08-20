from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = [
    "STSong",
    "Songti SC",
    "Heiti SC",
    "Arial Unicode MS",
    "Hiragino Sans GB",
    ".Hiragino Sans GB Interface",
    "GB18030 Bitmap",
    "PingFang SC",
    "SimSun",
    "Microsoft YaHei",
]

doc = Document()
for font in fonts:
    p = doc.add_paragraph()
    run = p.add_run(f"{font}: 固定燃烧源排放清单实验记录，工业锅炉与火电热力。")
    run.font.name = font
    run.font.size = Pt(15)
    fonts_el = run._element.get_or_add_rPr().rFonts
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts_el.set(qn(f"w:{attr}"), font)
    fonts_el.set(qn("w:hint"), "eastAsia")
doc.save("documentation/source/font_probe.docx")
