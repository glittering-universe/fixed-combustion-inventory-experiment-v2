#!/usr/bin/env python3
"""Build a beginner-facing, CLI-only fixed-combustion experiment manual."""

from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/wushuo/Desktop/环境学院论文")
EXPERIMENT = ROOT / "固定燃烧源清单经验缺失处置实验"
TUTORIAL = ROOT / "教程临时运行" / "固定燃烧源清单实验手册"
RUNS = TUTORIAL / "tutorial_runs"
SCREENSHOTS = EXPERIMENT / "documentation" / "tutorial_screenshots"
OUTPUT = ROOT / "固定燃烧源清单实验手册.docx"
LOCAL_COPY = EXPERIMENT / "documentation" / "固定燃烧源清单实验手册.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1D2733"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF4EA"
PALE_YELLOW = "FFF4CE"
WHITE = "FFFFFF"
BORDER = "C9D2DC"
LATIN_FONT = "Arial"
# Arial Unicode MS is available in the macOS/LibreOffice render path used for
# final QA and contains the full Chinese glyph set required by this manual.
BODY_FONT = "Arial Unicode MS"
CONTENT_WIDTH_DXA = 9360


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def set_run_font(run, size=11, bold=False, color=INK, italic=False, mono=False):
    run.font.name = "Courier New" if mono else BODY_FONT
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Courier New" if mono else LATIN_FONT)
    fonts.set(qn("w:hAnsi"), "Courier New" if mono else LATIN_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    fonts.set(qn("w:cs"), "Courier New" if mono else LATIN_FONT)
    fonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.5, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_geometry(table, widths_dxa: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_body(doc, text: str, *, after=6, italic=False, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        set_run_font(p.add_run(text), italic=italic)
    return p


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_run_font(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)
    p.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[level])
    p.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}[level])
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.18
    set_run_font(p.add_run(text))
    return p


def add_number(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    set_run_font(p.add_run(text))
    return p


def add_callout(doc, text: str, fill=LIGHT_BLUE, color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_cell_shading(table.cell(0, 0), fill)
    set_cell_text(table.cell(0, 0), text, bold=True, color=color, size=10.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(doc, text: str, *, size=7.4):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = False
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    set_run_font(p.add_run(text.rstrip() + "\n"), size=size, mono=True)
    return p


def add_table(doc, headers: list[str], rows: Iterable[Iterable], widths: list[int], font_size=9.2):
    rows = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=font_size, center=True)
        set_cell_shading(table.rows[0].cells[i], DARK_BLUE)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for r_index, values in enumerate(rows):
        cells = table.add_row().cells
        for c_index, value in enumerate(values):
            set_cell_text(cells[c_index], value, size=font_size, center=c_index == 0 and len(headers) > 2)
            if r_index % 2:
                set_cell_shading(cells[c_index], LIGHT_GRAY)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot(doc, name: str, width=6.35):
    path = SCREENSHOTS / name
    if not path.is_file():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(width))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("固定燃烧源清单实验手册  ·  "), size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def command(run_id: str, method: str, scenario: str, target: str, scale: str, extra: str = "") -> str:
    value = (
        "python3 \"$EXPERIMENT_ROOT/tutorial/run_tutorial_case.py\" \\\n"
        "  --root \"$EXPERIMENT_ROOT\" \\\n"
        f"  --run-id {run_id} \\\n"
        f"  --method {method} \\\n"
        f"  --scenario {scenario} \\\n"
        f"  --target {target} \\\n"
        f"  --scale {scale}"
    )
    if extra:
        value += " \\\n  " + extra
    return value


def report_status(run_id: str) -> tuple[str, str, str, str]:
    report = load(RUNS / run_id / "evaluation" / "reference_independent_report.json")
    checks = report["structural_and_arithmetic_checks"]
    return (
        run_id,
        report["sealed_integrity"]["status"],
        checks["status"],
        report["reference_package"]["status"],
    )


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.82)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.42)
section.footer_distance = Inches(0.42)

normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
normal._element.rPr.rFonts.set(qn("w:cs"), LATIN_FONT)
normal._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name in ("List Bullet", "List Bullet 2", "List Number"):
    style = doc.styles[name]
    style.font.name = BODY_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style.font.size = Pt(11)

for level, size, color, before, after in ((1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, DARK_BLUE, 10, 5)):
    style = doc.styles[f"Heading {level}"]
    style.font.name = BODY_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
set_run_font(header.add_run("固定燃烧源清单实验手册"), size=9, color=MUTED)
add_page_number(section.footer.paragraphs[0])

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(4)
set_run_font(title.add_run("固定燃烧源清单实验手册"), size=24, bold=True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
set_run_font(subtitle.add_run("面向只熟悉基本电脑与 Excel 操作的读者·macOS·Hermes CLI"), size=12.5, color=MUTED)
add_callout(doc, "这次演示不是正式实验。", fill=PALE_YELLOW, color=INK)
add_body(doc, "本手册默认 Hermes CLI 已安装，从 API 配置开始。所有操作只通过 CLI 完成，不使用 Hermes Desktop。手册中的命令都以 macOS 为准，且每次只执行一条命令。")

add_heading(doc, "1 先认识这个实验", 1)
add_body(doc, "本实验把工业锅炉和火电热力原始表格转换为按设备、燃料和污染物组织的排放清单。Agent 负责组织阶段和调用能力；规则包、核算项结构、准入控制、确定性计算和完整过程记录约束实际编制行为。")
add_table(doc, ["名称", "说明"], [
    ("Profile", "一套独立的 Hermes 配置、密钥、Plugin 和会话空间。本手册固定使用 fixed-combustion-inventory-tutorial。"),
    ("Plugin", "固定燃烧清单的顶层扩展包，包含阶段 Skill 和可执行 Tool。"),
    ("Skill", "规定某一阶段怎样处理、需要什么输入、应输出什么。"),
    ("Tool", "执行读取表格、构造核算项、匹配规则、计算和导出等确定操作。"),
    ("规则包", "已复核并冻结的 T/CSES 144—2024 可执行规则、参数和映射。"),
    ("运行包", "某一次任务的输入、Prompt、运行记录、输出和封存文件的集合。"),
    ("封存", "通过 experiment_seal.json 记录文件哈希和运行信息，表明本次运行已结束。"),
    ("not_computable", "缺少外部参照或预先冻结阈值时，某项指标不可计算；不是 0 分。"),
], [2000, 7360], font_size=9.3)
add_body(doc, "路径是文件或文件夹在电脑中的位置。手册使用 <实验材料目录> 这类占位符时，需要替换为你自己的路径；后续命令统一使用 $SOURCE_PACKAGE 和 $EXPERIMENT_ROOT，避免反复输入长路径。")
add_body(doc, "Prompt 是发给 Agent 的任务说明，不是执行命令。读者不需手工修改长 Prompt；启动器会根据命令参数填入运行编号、源类、规模和输出目录。")

add_heading(doc, "2 建立完全隔离的教程环境", 1)
add_heading(doc, "2.1 设置路径变量", 2)
add_body(doc, "把第一条命令中的占位符替换为实验材料包所在目录。第二条命令创建教程工作区的统一路径。")
add_code(doc, 'export SOURCE_PACKAGE="<实验材料目录>"')
add_code(doc, 'export EXPERIMENT_ROOT="$HOME/Documents/fixed-combustion-inventory-tutorial"')
add_code(doc, 'mkdir -p "$EXPERIMENT_ROOT"')
add_heading(doc, "2.2 复制教程所需文件", 2)
add_body(doc, "所有内容都使用实体副本，不使用软链接。教程运行只写入 $EXPERIMENT_ROOT。")
for folder in ("experiment_control", "evaluation", "method_package", "plugin", "reproduction", "rules", "tutorial"):
    add_code(doc, f'cp -R "$SOURCE_PACKAGE/{folder}" "$EXPERIMENT_ROOT/"')
add_code(doc, 'mkdir -p "$EXPERIMENT_ROOT/inputs/prepared" "$EXPERIMENT_ROOT/inputs/experiment_b" "$EXPERIMENT_ROOT/tutorial_runs"')
add_code(doc, 'cp "$SOURCE_PACKAGE/inputs/prepared/"*.xlsx "$EXPERIMENT_ROOT/inputs/prepared/"')
add_code(doc, 'cp "$SOURCE_PACKAGE/inputs/prepared/scope_catalog.json" "$EXPERIMENT_ROOT/inputs/prepared/"')
add_code(doc, 'cp -R "$SOURCE_PACKAGE/inputs/experiment_b/." "$EXPERIMENT_ROOT/inputs/experiment_b/"')
add_code(doc, 'cp "$SOURCE_PACKAGE/1 城市大气污染源排放清单编制技术指南 T_CSES 144-2024.pdf" "$EXPERIMENT_ROOT/"')

add_heading(doc, "3 创建临时 Profile 并配置 API", 1)
add_heading(doc, "3.1 创建空白 Profile", 2)
add_code(doc, 'hermes profile create fixed-combustion-inventory-tutorial --description "固定燃烧源清单教程临时演示" --no-skills')
add_body(doc, "成功时应看到 Profile 的创建位置和同名启动器。这是一个空白 Profile，不从正式 Profile 克隆设置、会话或 Plugin。")
add_heading(doc, "3.2 写入模型和 Provider 配置", 2)
for key, value in (
    ("model.default", "deepseek-v4-pro"),
    ("model.provider", "custom"),
    ("model.base_url", "https://one-hub.hycx-gd.cn/v1"),
    ("model.key_env", "HERMES_CUSTOM_CUSTOM_API_KEY"),
    ("providers.custom.name", "环院"),
    ("providers.custom.base_url", "https://one-hub.hycx-gd.cn/v1"),
    ("providers.custom.model", "deepseek-v4-pro"),
    ("providers.custom.key_env", "HERMES_CUSTOM_CUSTOM_API_KEY"),
    ("agent.max_turns", "80"),
    ("terminal.cwd", '"$EXPERIMENT_ROOT"'),
):
    add_code(doc, f"hermes -p fixed-combustion-inventory-tutorial config set {key} {value}")
add_heading(doc, "3.3 安全写入 API 密钥", 2)
add_body(doc, "执行第一条后粘贴密钥并按回车；输入过程不会显示字符。第二条将密钥写入临时 Profile 的 .env，第三条清除当前 shell 变量。")
add_code(doc, 'read -s "HERMES_API_KEY?请粘贴 API 密钥（输入不会显示）: "')
add_code(doc, 'printf "\\nHERMES_CUSTOM_CUSTOM_API_KEY=%s\\n" "$HERMES_API_KEY" > "$HOME/.hermes/profiles/fixed-combustion-inventory-tutorial/.env"')
add_code(doc, 'unset HERMES_API_KEY')
add_callout(doc, "不要把真实 API 密钥写进 Word、Prompt、截图或实验结果。")

add_heading(doc, "4 接入实体 Plugin 副本", 1)
add_code(doc, 'mkdir -p "$HOME/.hermes/profiles/fixed-combustion-inventory-tutorial/plugins"')
add_code(doc, 'cp -R "$EXPERIMENT_ROOT/plugin/fixed-combustion-inventory" "$HOME/.hermes/profiles/fixed-combustion-inventory-tutorial/plugins/"')
add_code(doc, 'hermes -p fixed-combustion-inventory-tutorial plugins enable fixed-combustion-inventory --no-allow-tool-override')
add_code(doc, 'hermes -p fixed-combustion-inventory-tutorial plugins list --enabled --plain')
add_body(doc, "成功时列表中应出现 enabled 和 fixed-combustion-inventory。--no-allow-tool-override 表示 Plugin 不允许覆盖 Hermes 内置工具。")

add_heading(doc, "5 运行前统一检查", 1)
add_body(doc, "在第一次调用 Agent 前，统一检查 Profile、模型、Provider、密钥文件、Plugin 实体副本、规则包、两类输入和输出边界。")
add_code(doc, 'python3 "$EXPERIMENT_ROOT/tutorial/check_tutorial_environment.py" --root "$EXPERIMENT_ROOT"')
add_screenshot(doc, "01_preflight.png")
add_body(doc, "只有顶层 status 为 pass，且所有子项均为 pass 时才继续。检查器只判断密钥文件是否存在，不输出密钥内容。")

add_heading(doc, "6 启动器如何工作", 1)
add_body(doc, "所有演示都使用同一个单次运行启动器。它根据命令参数创建运行包，然后调用确定性程序或 Hermes CLI。每次任务使用唯一 run_id；如果同名运行目录已存在，启动器会拒绝覆盖。")
add_table(doc, ["参数", "用途", "例子"], [
    ("--run-id", "本次运行的唯一名称", "TUTORIAL-A-POWER-FULL"),
    ("--method", "完整方法、基线或单因素版本", "full"),
    ("--scenario", "工作流组装方式", "full"),
    ("--target", "工业锅炉或火电热力", "POWER"),
    ("--scale", "单记录或25%/50%/75%/100%", "25"),
    ("--scope-id", "只在 single 时指定记录", "SRC-PWR-000078"),
    ("--variant", "原始结构或扰动版本", "B0 / S4 / B2"),
], [1600, 4600, 3160], font_size=9.0)

add_heading(doc, "7 先跑通一条记录", 1)
add_body(doc, "这一步只用来确认环境能够完成整条链路，不进入实验 A—D 统计。使用一条火电热力记录和 Full 方法。")
add_code(doc, command("TUTORIAL-SINGLE-POWER-FULL", "full", "full", "POWER", "single", "--scope-id SRC-PWR-000078"))
add_screenshot(doc, "02_single_run.png")
add_body(doc, "命令运行时可能一段时间没有新文字。只要命令尚未返回、错误日志没有报错，就应继续等待。成功时最后会显示 status=pass、run_id、run_package 和 sealed_files。")
add_heading(doc, "7.1 查看运行包", 2)
add_code(doc, 'find "$EXPERIMENT_ROOT/tutorial_runs/TUTORIAL-SINGLE-POWER-FULL" -maxdepth 2 -type f | sort')
add_screenshot(doc, "07_output_tree.png")
add_heading(doc, "7.2 四层验收", 2)
add_number(doc, "进程检查：命令正常结束，execution_metrics.json 中 return_code 为 0。")
add_number(doc, "封存检查：experiment_seal.json 和 sealed_manifest.json 存在。")
add_number(doc, "结果检查：目标 Excel、最终异常清单.csv 和 quality_summary.json 存在。")
add_number(doc, "过程检查：完整计算过程、Hermes 会话导出和用量记录存在。")
add_callout(doc, "Agent 的文字回复不是实验成功的唯一依据；实体结果和封存记录才是最终判断依据。")

add_heading(doc, "8 实验 A：完整编制流程对比", 1)
add_body(doc, "正式实验要求两类全量数据都运行，机器流程各重复3次，人工基线由真实参与者给出。教程临时演示只在火电热力25%范围各运行1次三种机器流程，并补充1条工业锅炉 Full 任务。")
add_heading(doc, "8.1 Full 方法", 2)
add_code(doc, command("TUTORIAL-A-POWER-FULL", "full", "full", "POWER", "25"))
add_heading(doc, "8.2 确定性程序式流程", 2)
add_code(doc, command("TUTORIAL-A-POWER-DET", "deterministic_program", "full", "POWER", "25"))
add_heading(doc, "8.3 通用工具调用型 Agent", 2)
add_code(doc, command("TUTORIAL-A-POWER-GENERIC", "generic_tool_agent", "full", "POWER", "25"))
add_heading(doc, "8.4 工业锅炉单记录补充", 2)
add_code(doc, command("TUTORIAL-A-INDUSTRIAL-FULL-SINGLE", "full", "full", "INDUSTRIAL", "single", "--scope-id SRC-IND-000229"))
add_heading(doc, "8.5 分别进行独立评价", 2)
for rid in ("TUTORIAL-A-POWER-FULL", "TUTORIAL-A-POWER-DET", "TUTORIAL-A-POWER-GENERIC"):
    add_code(doc, f'python3 "$EXPERIMENT_ROOT/evaluation/evaluate_sealed_run.py" "$EXPERIMENT_ROOT/tutorial_runs/{rid}" --output "$EXPERIMENT_ROOT/tutorial_runs/{rid}/evaluation/reference_independent_report.json"')
add_screenshot(doc, "03_experiment_a.png")
add_table(doc, ["临时运行", "封存完整性", "结构与最小复算", "专业参照"], [report_status(rid) for rid in ("TUTORIAL-A-POWER-FULL", "TUTORIAL-A-POWER-DET", "TUTORIAL-A-POWER-GENERIC")], [3300, 1800, 2350, 1910], font_size=8.8)
add_body(doc, "无专业参照包时，规则落实准确性和最终排放量准确性保持 not_computable。这不影响检查封存、结构、数值基本约束和资源记录。")

add_heading(doc, "9 实验 B：输入变化和异常处置", 1)
add_body(doc, "B1考察表头形式、已登记别名、列顺序和封面工作表变化前后，同一方法的语义结果是否保持。B2在受控记录中注入缺失或冲突，检查异常是否被识别，未受影响部分是否保持。正式实验 B 包含既有专家主导型流程与机器流程的对比；临时演示只展示可复现的机器处理。")
add_heading(doc, "9.1 未扰动 B0", 2)
add_code(doc, command("TUTORIAL-B-B0-POWER-DET", "deterministic_program", "full", "POWER", "25", "--variant B0"))
add_heading(doc, "9.2 组合结构扰动 S4", 2)
add_code(doc, command("TUTORIAL-B-S4-POWER-DET", "deterministic_program", "full", "POWER", "25", "--variant S4"))
add_heading(doc, "9.3 受控缺失与冲突 B2", 2)
add_code(doc, command("TUTORIAL-B-B2-POWER-DET", "deterministic_program", "full", "POWER", "25", "--variant B2"))
add_heading(doc, "9.4 规范化并比较 B0 与 S4", 2)
for rid in ("TUTORIAL-B-B0-POWER-DET", "TUTORIAL-B-S4-POWER-DET"):
    add_code(doc, f'python3 "$EXPERIMENT_ROOT/evaluation/normalize_run_output.py" "$EXPERIMENT_ROOT/tutorial_runs/{rid}" --output "$EXPERIMENT_ROOT/tutorial_runs/{rid}/evaluation/normalized_items.csv"')
add_code(doc, 'python3 "$EXPERIMENT_ROOT/evaluation/compare_normalized_runs.py" "$EXPERIMENT_ROOT/tutorial_runs/TUTORIAL-B-B0-POWER-DET/evaluation/normalized_items.csv" "$EXPERIMENT_ROOT/tutorial_runs/TUTORIAL-B-S4-POWER-DET/evaluation/normalized_items.csv" --output "$EXPERIMENT_ROOT/tutorial/B_B0_vs_S4_comparison.json"')
add_code(doc, 'python3 "$EXPERIMENT_ROOT/evaluation/evaluate_sealed_run.py" "$EXPERIMENT_ROOT/tutorial_runs/TUTORIAL-B-B2-POWER-DET" --output "$EXPERIMENT_ROOT/tutorial_runs/TUTORIAL-B-B2-POWER-DET/evaluation/reference_independent_report.json"')
add_screenshot(doc, "04_experiment_b.png")
add_body(doc, "本次临时演示中，B0与S4的828个比较单元全部保持；B2运行形成13个异常根因组。这些数字只用于说明检查方法，不替代正式全量实验结果。")

add_heading(doc, "10 实验 C：规模扩展", 1)
add_body(doc, "正式实验对25%、50%、75%和100%四个嵌套规模运行三种机器流程，且包含全量规模。临时演示使用确定性程序依次执行四个火电热力规模，用来展示规模参数和资源记录的读取方式。")
for scale in (25, 50, 75, 100):
    add_heading(doc, f"10.{(25, 50, 75, 100).index(scale) + 1}  {scale}% 规模", 2)
    add_code(doc, command(f"TUTORIAL-C-{scale}-POWER-DET", "deterministic_program", "full", "POWER", str(scale)))
add_screenshot(doc, "05_experiment_c.png")
add_body(doc, "每个运行的 logs/execution_metrics.json 保存 wall_seconds、CPU时间、峰值内存、输入和输出字节数。只有事前冻结且对所有方法独立的时间阈值才能进入合成评分；否则只报告原始资源观测。")

add_heading(doc, "11 实验 D：单因素方法组成分析", 1)
add_body(doc, "实验 D 与完整流程对比不同：每个版本只替换一个方法组成，其他输入、模型、Plugin、共享阶段和评价程序保持不变。临时演示使用与第7节相同的单记录。")
d_specs = (
    ("TUTORIAL-D-WO-ITEM", "w_o_calculation_item_structure", "w_o_calculation_item_structure", "去掉污染物核算项结构"),
    ("TUTORIAL-D-WO-RULE", "w_o_executable_rule_gate", "w_o_executable_rule_gate", "去掉可执行规则门控"),
    ("TUTORIAL-D-AGENT-ADMISSION", "agent_managed_admission", "agent_managed_admission", "Agent管理核算准入"),
    ("TUTORIAL-D-WO-TRACE", "w_o_complete_trace", "w_o_complete_trace", "去掉完整核算过程记录"),
)
for index, (rid, method, scenario, label) in enumerate(d_specs, start=1):
    add_heading(doc, f"11.{index}  {label}", 2)
    add_code(doc, command(rid, method, scenario, "POWER", "single", "--scope-id SRC-PWR-000078"))
    add_code(doc, f'python3 "$EXPERIMENT_ROOT/evaluation/evaluate_sealed_run.py" "$EXPERIMENT_ROOT/tutorial_runs/{rid}" --output "$EXPERIMENT_ROOT/tutorial_runs/{rid}/evaluation/reference_independent_report.json"')
add_screenshot(doc, "06_experiment_d.png")
add_body(doc, "单记录演示中，w/o核算项结构和w/o可执行规则门控的结构与最小复算检查未通过；Agent管理准入和w/o完整过程记录的结构检查通过。正式实验D必须对两类全量数据各运行1次，且不计算总EICPI。")

add_heading(doc, "12 如何看懂结果目录", 1)
add_table(doc, ["文件或目录", "作用", "验收要点"], [
    ("run_manifest.json", "本次方法、源类、规模、输入和哈希", "与命令参数一致"),
    ("outputs/*.xlsx", "核算项和源级结果", "无公式错误和外部链接"),
    ("outputs/最终异常清单.csv", "少量需要用户处理的根因组", "不是逐单元格重复报错"),
    ("outputs/quality_summary.json", "输出结构和工作簿检查摘要", "error_count、formula_error_count"),
    ("trace/*.jsonl", "完整或最小核算过程", "是否符合本次方法版本"),
    ("logs/execution_metrics.json", "时间、CPU、内存、I/O、会话和用量", "return_code=0"),
    ("session_export/session.jsonl", "脱敏后的 Hermes CLI 会话", "无密钥和敏感数据"),
    ("experiment_seal.json", "运行文件哈希和字节数", "存在且与实体文件一致"),
    ("evaluation/reference_independent_report.json", "只读独立评价", "区分pass、fail和not_computable"),
], [2900, 3600, 2860], font_size=8.8)

add_heading(doc, "13 常见问题和判断方法", 1)
add_table(doc, ["现象", "先检查什么", "怎样处理"], [
    ("401或403", "API密钥、base_url、Provider", "重新使用隐藏输入写入临时Profile，不要截图密钥"),
    ("429", "API频率或配额", "等待后重新创建一个新run_id，不覆盖旧运行包"),
    ("Plugin工具不存在", "plugins list与Plugin实体目录", "确认已完整复制并显式enable，不使用软链接"),
    ("长时间无输出", "命令是否仍在运行、stderr是否有错误", "不要因为文字暂停就强制结束；通用Agent解析PDF可能较慢"),
    ("有Excel但无封存", "execution_metrics、session_export、experiment_seal", "本次仍判为未完成，先查明收尾失败原因"),
    ("同名run_id被拒绝", "tutorial_runs中是否已有同名目录", "使用新run_id；不手动覆盖旧记录"),
    ("not_computable", "是否缺专业参照包或事前阈值", "保留not_computable，不改成0，也不用内部检查替代"),
], [2100, 3320, 3940], font_size=8.7)

add_heading(doc, "14 完整运行前检查器源码", 1)
add_body(doc, "下面是本手册实际使用的完整脚本。正常跟做时直接使用实验材料包中的文件，不需要从 Word 重新输入。")
add_code(doc, (TUTORIAL / "tutorial" / "check_tutorial_environment.py").read_text(encoding="utf-8"), size=6.9)

add_heading(doc, "15 完整单次运行启动器源码", 1)
add_code(doc, (TUTORIAL / "tutorial" / "run_tutorial_case.py").read_text(encoding="utf-8"), size=6.9)

add_heading(doc, "16 完整 Prompt", 1)
add_body(doc, "Prompt在启动时由程序追加运行编号、运行包路径、方法哈希、源类、规模和输入版本。读者不手工修改以下正文。")
add_heading(doc, "16.1 Full 方法 Prompt", 2)
add_code(doc, (TUTORIAL / "method_package" / "full_prompt.md").read_text(encoding="utf-8"), size=7.0)
add_heading(doc, "16.2 通用工具调用型 Agent Prompt", 2)
add_code(doc, (TUTORIAL / "method_package" / "generic_agent_prompt.md").read_text(encoding="utf-8"), size=7.0)
add_heading(doc, "16.3 实验 D Prompt", 2)
add_code(doc, (TUTORIAL / "method_package" / "ablation_prompt.md").read_text(encoding="utf-8"), size=6.8)

add_heading(doc, "17 完成检查", 1)
add_bullet(doc, "所有命令使用 fixed-combustion-inventory-tutorial Profile，没有使用正式 Profile。")
add_bullet(doc, "Plugin、规则、输入、脚本和输出都是实体副本，不是软链接或硬链接。")
add_bullet(doc, "每个命令使用唯一run_id，每个运行包具有experiment_seal.json。")
add_bullet(doc, "目标Excel、异常清单、质量摘要、过程记录、会话导出和用量记录齐全。")
add_bullet(doc, "独立评价只读封存结果，没有在评价阶段修正参评方法的输出。")
add_bullet(doc, "人工基线、专业参照包和独立时间阈值没有由教程伪造；缺失时保持not_computable。")
add_callout(doc, "只有同时通过运行、封存、结果和过程四层检查，才能把一次任务记为成功。", fill=PALE_GREEN, color="2F6B3A")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
LOCAL_COPY.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
shutil.copyfile(OUTPUT, LOCAL_COPY)
print(json.dumps({"output": str(OUTPUT), "local_copy": str(LOCAL_COPY)}, ensure_ascii=False))
