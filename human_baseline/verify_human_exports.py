"""Reconcile corrected human exports with preserved workbook cells, read-only."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
# These addresses were independently read from the two preserved original
# headers. They are audit expectations for these immutable templates, not the
# production reader's selection logic.
POLLUTANTS = ("SO2", "NOx", "CO", "PM10", "PM2.5", "BC", "OC", "VOC", "NH3")
AUDITED_COLUMNS = {
    "INDUSTRIAL": {
        "generation": dict(zip(POLLUTANTS, ("BX", "BY", "BZ", "CA", "CB", "CC", "CD", "CE", "CF"))),
        "emission": dict(zip(POLLUTANTS, ("CN", "CO", "CP", "CQ", "CR", "CS", "CT", "CU", "CV"))),
    },
    "POWER": {
        "generation": dict(zip(POLLUTANTS, ("CA", "CB", "CC", "CD", "CE", "CF", "CG", "CH", "CI"))),
        "emission": dict(zip(POLLUTANTS, ("CQ", "CR", "CS", "CT", "CU", "CV", "CW", "CX", "CY"))),
    },
}


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader(); w.writerows(rows)


def numeric(value):
    if value is None or value == "":
        return None
    try:
        n = Decimal(str(value))
        return n if n.is_finite() else None
    except Exception:
        return None


def header_pollutant(value):
    s = "".join(unicodedata.normalize("NFKC", str(value)).split()).upper()
    name, unit = s.split("(", 1)
    if unit not in ("T)", "吨)"):
        raise ValueError(f"unexpected original result unit: {value}")
    return {"NOX": "NOx", "VOCS": "VOC", "PM25": "PM2.5"}.get(name, name)


def verify(normalized_root: Path, report_root: Path):
    index = json.loads((normalized_root / "normalization_index.json").read_text())
    original_root = Path(index["preserved_original_root"])
    report_root.mkdir(parents=True, exist_ok=True)
    column_rows, coverage_rows, mismatches = [], [], []
    cell_count = 0
    old_root = ROOT / "human_baseline" / "normalized_v2"
    for package in index["packages"]:
        wb_path = original_root / package["workbook"]
        target = package["target"]
        records = read_csv(normalized_root / package["output"] / "calculation_totals.csv")
        actual = {(int(r["human_workbook_row"]), r["pollutant"]): r for r in records}
        if len(actual) != len(records):
            raise ValueError(f"duplicate workbook-row/pollutant observation: {wb_path}")
        old_path = old_root / package["output"] / "calculation_totals.csv"
        old_records = read_csv(old_path) if old_path.exists() else []
        old = {(int(r["human_workbook_row"]), r["pollutant"]): r for r in old_records}
        w = load_workbook(wb_path, read_only=True, data_only=True)
        s = w.worksheets[0]; s.reset_dimensions()
        it = s.iter_rows(values_only=True); headers = tuple(next(it))
        source_rows = {row_no: tuple(values) for row_no, values in enumerate(it, 2)
                       if len(values) > 3 and values[0] not in (None, "") and values[3] not in (None, "")}
        sheet_name = s.title; w.close()
        expected_keys = {(row_no, p) for row_no in source_rows for p in POLLUTANTS}
        row_coverage_ok = set(actual) == expected_keys
        hash_ok = digest(wb_path) == index["original_workbook_hashes"][package["workbook"]]
        year_counts = Counter(str(values[0]) for values in source_rows.values())
        notes = json.loads((normalized_root / package["output"] / "normalization_notes.json").read_text())
        for role in ("generation", "emission"):
            for pollutant, column in AUDITED_COLUMNS[target][role].items():
                pos = column_index_from_string(column) - 1
                if header_pollutant(headers[pos]) != pollutant:
                    raise ValueError(f"independent header expectation failed: {wb_path}:{column}1")
                stats = Counter(); source_sum = Decimal(0); export_sum = Decimal(0); old_sum = Decimal(0)
                for row_no, values in source_rows.items():
                    cell_count += 1
                    raw = values[pos] if pos < len(values) else None
                    expected_value = numeric(raw)
                    r = actual.get((row_no, pollutant))
                    got = r.get(role + "_t", "") if r else ""
                    observed_value = numeric(got)
                    stats["source_number" if expected_value is not None else "source_blank" if raw in (None, "") else "source_non_numeric"] += 1
                    stats["export_number" if observed_value is not None else "export_blank"] += 1
                    if expected_value is not None: source_sum += expected_value
                    if observed_value is not None: export_sum += observed_value
                    old_value = numeric(old.get((row_no, pollutant), {}).get(role + "_t"))
                    if old_value is not None: old_sum += old_value
                    if old_value != observed_value: stats["changed_from_prior_export"] += 1
                    ok = r is not None and expected_value == observed_value
                    ok = ok and (expected_value is not None or got == "")
                    ok = ok and r.get(role + "_cell") == f"{column}{row_no}"
                    ok = ok and r.get(role + "_header") == str(headers[pos])
                    ok = ok and r.get("human_sheet") == sheet_name
                    kind = "number" if expected_value is not None else "blank" if raw in (None, "") else "non_numeric"
                    ok = ok and r.get(role + "_value_kind") == kind
                    stats["matched" if ok else "mismatched"] += 1
                    if not ok:
                        mismatches.append({"workbook":package["workbook"],"sheet":sheet_name,"cell":f"{column}{row_no}","pollutant":pollutant,"role":role,"source_value":str(raw),"export_value":got})
                column_rows.append({
                    "workbook": package["workbook"], "target":target,"sheet":sheet_name,"role":role,"pollutant":pollutant,
                    "original_column":column,"original_header":headers[pos],"rows":len(source_rows),
                    "source_numeric_cells":stats["source_number"],"source_blank_cells":stats["source_blank"],
                    "source_non_numeric_cells":stats["source_non_numeric"],"export_numeric_cells":stats["export_number"],"export_blank_cells":stats["export_blank"],
                    "source_numeric_sum":str(source_sum),"export_numeric_sum":str(export_sum),"prior_export_numeric_sum":str(old_sum),
                    "changed_values_from_prior_export":stats["changed_from_prior_export"],"matched_cells":stats["matched"],"mismatched_cells":stats["mismatched"],
                    "sum_equal":source_sum==export_sum,
                })
        source_decisions = normalized_root / package["output"] / "source_decisions.csv"
        prior_decisions = old_root / package["output"] / "source_decisions.csv"
        coverage_rows.append({"workbook":package["workbook"],"source_valid_rows":len(source_rows),"expected_source_pollutant_rows":len(expected_keys),
                              "exported_rows":len(records),"complete_unique_row_coverage":row_coverage_ok,"original_workbook_hash_unchanged":hash_ok,
                              "scope_decisions_byte_identical_to_prior":prior_decisions.exists() and digest(source_decisions)==digest(prior_decisions),
                              "inventory_year_counts":json.dumps(dict(year_counts),ensure_ascii=False),
                              "scope_source_count":notes["human_scope_include_count"]})
        print(json.dumps({"checked_workbook":package["workbook"],"cells":len(source_rows)*18},ensure_ascii=False),flush=True)
    write_csv(report_root/"column_reconciliation.csv",column_rows,list(column_rows[0]))
    write_csv(report_root/"row_coverage.csv",coverage_rows,list(coverage_rows[0]))
    write_csv(report_root/"cell_mismatches.csv",mismatches,["workbook","sheet","cell","pollutant","role","source_value","export_value"])
    summary={"normalization_contract":index["schema_version"],"workbook_count":len(coverage_rows),"column_checks":len(column_rows),"source_cell_checks":cell_count,
             "mismatched_cells":len(mismatches),"all_column_sums_equal":all(r["sum_equal"] for r in column_rows),
             "all_source_rows_covered_once":all(r["complete_unique_row_coverage"] for r in coverage_rows),
             "all_original_workbook_hashes_unchanged":all(r["original_workbook_hash_unchanged"] for r in coverage_rows),
             "all_scope_decisions_unchanged":all(r["scope_decisions_byte_identical_to_prior"] for r in coverage_rows),
             "changed_generation_or_emission_values":sum(r["changed_values_from_prior_export"] for r in column_rows),
             "verification_basis":"Independent original-header column addresses and workbook cached values; no emissions recalculated; adapter-generated cell trace distinguished from historical human process records."}
    summary["status"]="pass" if not mismatches and all(summary[k] for k in ("all_column_sums_equal","all_source_rows_covered_once","all_original_workbook_hashes_unchanged","all_scope_decisions_unchanged")) else "fail"
    (report_root/"verification_report.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    if summary["status"]!="pass":raise RuntimeError("human export reconciliation failed")
    return summary


if __name__=="__main__":
    parser=argparse.ArgumentParser();parser.add_argument("--normalized-root",type=Path,default=ROOT/"human_baseline/normalized_v2_1");parser.add_argument("--report-root",type=Path,default=ROOT/"human_baseline/reading_audit_v2_1")
    args=parser.parse_args();verify(args.normalized_root.resolve(),args.report_root.resolve())
