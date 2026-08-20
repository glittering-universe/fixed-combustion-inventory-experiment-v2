#!/usr/bin/env python3
"""Normalize preserved expert-led workbooks against the raw Base-102 universe.

This adapter observes only what the supplied workbooks contain.  It never
copies fields from the professional reference package and never reconstructs
unreported activity, parameter, control, rule, or intermediate-calculation
records.  Those fields are explicitly exported as ``NA``.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
import sys
from collections import defaultdict, deque
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from experiment_control.source_identity import (
    identity_from_base102_row,
    pseudonymous_entity_id,
)


WORKSPACE = ROOT.parent
DEFAULT_ORIGINAL_ROOT = ROOT / "human_baseline" / "original_packages"
DEFAULT_OUTPUT_ROOT = ROOT / "human_baseline" / "normalized_v2"
DEFAULT_RAW_102 = (
    WORKSPACE
    / "01 环境统计数据"
    / "2022年基表查询工业企业锅炉_燃气轮机污染物和温室气体排放及治理情况(基102表)2023102002.xlsx"
)
DEFAULT_RAW_101_ENTERPRISE = (
    WORKSPACE
    / "01 环境统计数据"
    / "2022年基表查询工业企业污染物和温室气体排放及治理情况(基101表)2023102001.xlsx"
)
DEFAULT_RAW_101_CONTROL = (
    WORKSPACE
    / "01 环境统计数据"
    / "2022年基表查询工业企业污染物和温室气体排放及治理情况(基101表)2023102002.xlsx"
)

TARGET_COUNTS = {"INDUSTRIAL": 4160, "POWER": 426}
POLLUTANTS = ("SO2", "NOx", "CO", "VOC", "PM10", "PM2.5", "BC", "OC", "NH3")
SECONDS_PER_RECORD = 11.52
NA = "NA"

SOURCE_DECISION_FIELDS = (
    "source_id",
    "target",
    "human_scope_decision",
    "human_workbook_row",
    "mapping_method",
    "raw_base102_row",
    "raw_base102_sequence",
    "inventory_year",
    "pseudonymous_entity_id",
    "industry_code",
    "target_equipment_value",
    "enterprise_master_present",
    "enterprise_control_row_count",
)
CALCULATION_FIELDS = (
    "observation_id",
    "source_id",
    "target",
    "source_decision_role",
    "human_workbook_row",
    "pollutant",
    "status",
    "method",
    "generation_t",
    "emission_t",
    "activity_value",
    "activity_unit",
    "parameter_value",
    "parameter_unit",
    "control_method",
    "control_efficiency",
    "rule_id",
    "standard_reference",
    "intermediate_calculation",
    "reason_code",
)
EXCEPTION_FIELDS = (
    "exception_id",
    "source_id",
    "observation_id",
    "target",
    "human_workbook_row",
    "pollutant",
    "category",
    "description",
)


@dataclass(frozen=True)
class RawCandidate:
    source_id: str
    pseudonymous_entity_id: str
    raw_excel_row: int
    sequence: str
    inventory_year: str
    industry_code: str
    industrial_equipment: str
    power_equipment: str
    row_values: tuple[Any, ...]

    def features(self, target: str) -> dict[str, str]:
        return _raw_features(self.row_values, target, self.pseudonymous_entity_id)


@dataclass(frozen=True)
class RawContext:
    candidates: tuple[RawCandidate, ...]
    enterprise_rows: int
    control_rows: int
    enterprise_entity_ids: frozenset[str]
    control_counts: Mapping[str, int]

    @property
    def by_source_id(self) -> dict[str, RawCandidate]:
        return {item.source_id: item for item in self.candidates}


@dataclass(frozen=True)
class MappedHumanRow:
    human_excel_row: int
    source_id: str
    raw_excel_row: int
    mapping_method: str
    is_scope_row: bool
    row_values: tuple[Any, ...]


@dataclass(frozen=True)
class WorkbookMapping:
    workbook: Path
    target: str
    rows: tuple[MappedHumanRow, ...]
    scope_source_ids: tuple[str, ...]
    mapping_exceptions: tuple[dict[str, Any], ...]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: object) -> str:
    if value is None:
        return ""
    return "".join(str(value).split()).upper()


def _scalar(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: object) -> str:
    if value in (None, ""):
        return ""
    try:
        return format(Decimal(str(value)).normalize(), "f")
    except InvalidOperation:
        return _text(value)


def _industry(value: object) -> str:
    text = _scalar(value).split(".")[0]
    return text.zfill(4) if text.isdigit() else _text(value)


def _finite(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _entity_id(values: Sequence[Any]) -> str:
    return pseudonymous_entity_id(credit=values[2], organization=values[1], company=values[3])


def _human_entity_id(values: Sequence[Any]) -> str:
    """Ignore the historical reader's numeric-75 phantom in blank ID cells."""
    placeholders = {"", "75", "NA", "N/A", "NONE", "NULL"}
    credit = None if _text(values[2]) in placeholders else values[2]
    organization = None if _text(values[1]) in placeholders else values[1]
    return pseudonymous_entity_id(credit=credit, organization=organization, company=values[3])


def _pad(values: Sequence[Any], length: int = 110) -> tuple[Any, ...]:
    return tuple(values) + (None,) * max(0, length - len(values))


def _raw_features(values: Sequence[Any], target: str, entity_id: str) -> dict[str, str]:
    row = _pad(values)
    if target == "INDUSTRIAL":
        equipment, combustion, rated, low_nox = 23, 27, 28, 31
    else:
        equipment, combustion, rated, low_nox = 12, 16, 17, 20
    return {
        "entity": entity_id,
        "year": _scalar(row[0]),
        "industry": _industry(row[6]),
        "equipment": _text(row[equipment]),
        "combustion": _text(row[combustion]),
        "rated": _number(row[rated]),
        "low_nox": _text(row[low_nox]),
        "fuel1_type": _text(row[35]),
        "fuel1_amount": _number(row[36]),
        "fuel1_unit": _text(row[37]),
        "fuel2_type": _text(row[50]),
        "fuel2_amount": _number(row[51]),
        "fuel2_unit": _text(row[52]),
        "other_amount": _number(row[64]),
        "outlet": _text(row[65]),
    }


def _human_features(values: Sequence[Any], target: str) -> dict[str, str]:
    row = _pad(values)
    if target == "INDUSTRIAL":
        equipment, combustion, rated, low_nox = 10, 11, 12, 13
        fuel1_type, fuel1_amount, fuel1_unit = 14, 15, 16
        fuel2_type, fuel2_amount, fuel2_unit = 30, 31, 32
        other_amount, outlet = 46, 57
    else:
        equipment, combustion, rated, low_nox = 10, 12, 13, 14
        fuel1_type, fuel1_amount, fuel1_unit = 16, 17, 18
        fuel2_type, fuel2_amount, fuel2_unit = 32, 33, 34
        other_amount, outlet = 48, 59
    return {
        "entity": _human_entity_id(row),
        "year": _scalar(row[0]),
        "industry": _industry(row[7]),
        "equipment": _text(row[equipment]),
        "combustion": _text(row[combustion]),
        "rated": _number(row[rated]),
        "low_nox": _text(row[low_nox]),
        "fuel1_type": _text(row[fuel1_type]),
        "fuel1_amount": _number(row[fuel1_amount]),
        "fuel1_unit": _text(row[fuel1_unit]),
        "fuel2_type": _text(row[fuel2_type]),
        "fuel2_amount": _number(row[fuel2_amount]),
        "fuel2_unit": _text(row[fuel2_unit]),
        "other_amount": _number(row[other_amount]),
        "outlet": _text(row[outlet]),
    }


def _signature(features: Mapping[str, str]) -> tuple[str, ...]:
    return tuple(features[key] for key in (
        "entity", "year", "industry", "equipment", "combustion", "rated", "low_nox",
        "fuel1_type", "fuel1_amount", "fuel1_unit", "fuel2_type", "fuel2_amount",
        "fuel2_unit", "other_amount", "outlet",
    ))


_MATCH_WEIGHTS = {
    "industry": 1,
    "equipment": 4,
    "combustion": 2,
    "rated": 2,
    "low_nox": 1,
    "fuel1_type": 3,
    "fuel1_amount": 4,
    "fuel1_unit": 2,
    "fuel2_type": 1,
    "fuel2_amount": 2,
    "fuel2_unit": 1,
    "other_amount": 1,
    "outlet": 4,
}


def _fallback_candidate(
    human_features: Mapping[str, str],
    candidates: Iterable[RawCandidate],
    target: str,
) -> tuple[RawCandidate | None, str]:
    ranked: list[tuple[int, int, RawCandidate]] = []
    for candidate in candidates:
        raw = candidate.features(target)
        if raw["entity"] != human_features["entity"] or raw["year"] != human_features["year"]:
            continue
        score = sum(weight for key, weight in _MATCH_WEIGHTS.items() if raw[key] == human_features[key])
        ranked.append((score, -candidate.raw_excel_row, candidate))
    if not ranked:
        return None, "no_entity_year_candidate"
    ranked.sort(reverse=True, key=lambda item: (item[0], item[1]))
    best_score = ranked[0][0]
    if best_score < 10:
        return None, f"fallback_score_below_threshold:{best_score}"
    tied = [item for item in ranked if item[0] == best_score]
    if len(tied) > 1:
        # Human rows preserve the original Base-102 order; raw-row order is the
        # deterministic final disambiguator after the semantic score.
        return tied[0][2], f"ranked_fallback_ordered_tie:{best_score}:{len(tied)}"
    return ranked[0][2], f"ranked_fallback:{best_score}"


def _read_rows(path: Path) -> tuple[list[str], list[tuple[int, tuple[Any, ...]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = [_scalar(value) for value in next(iterator)]
        rows = []
        for excel_row, values in enumerate(iterator, start=2):
            padded = _pad(tuple(values))
            if padded[0] in (None, "") or padded[3] in (None, ""):
                continue
            rows.append((excel_row, padded))
        return headers, rows
    finally:
        workbook.close()


def _load_entity_table(path: Path) -> tuple[int, frozenset[str], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    entities: set[str] = set()
    counts: dict[str, int] = defaultdict(int)
    row_count = 0
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        next(iterator)
        for values in iterator:
            row = _pad(tuple(values), 20)
            if row[0] in (None, "") or row[3] in (None, ""):
                continue
            row_count += 1
            entity = _entity_id(row)
            entities.add(entity)
            counts[entity] += 1
    finally:
        workbook.close()
    return row_count, frozenset(entities), dict(counts)


def load_raw_context(
    raw_102: Path,
    raw_101_enterprise: Path,
    raw_101_control: Path,
) -> RawContext:
    """Load the three raw tables and build the target-neutral source catalog."""
    for path in (raw_102, raw_101_enterprise, raw_101_control):
        if not path.is_file():
            raise FileNotFoundError(path)

    workbook = load_workbook(raw_102, read_only=True, data_only=True)
    candidates: list[RawCandidate] = []
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = [_scalar(value) for value in next(iterator)]
        for excel_row, values in enumerate(iterator, start=2):
            row = _pad(tuple(values))
            if row[0] in (None, "") or row[3] in (None, ""):
                continue
            record = dict(zip(headers, row))
            identity = identity_from_base102_row(record)
            candidates.append(RawCandidate(
                source_id=identity["source_id"],
                pseudonymous_entity_id=identity["pseudonymous_entity_id"],
                raw_excel_row=excel_row,
                sequence=_scalar(row[9]),
                inventory_year=_scalar(row[0]),
                industry_code=_industry(row[6]),
                industrial_equipment=_scalar(row[23]),
                power_equipment=_scalar(row[12]),
                row_values=row,
            ))
    finally:
        workbook.close()

    enterprise_rows, enterprise_entities, _ = _load_entity_table(raw_101_enterprise)
    control_rows, _, control_counts = _load_entity_table(raw_101_control)
    source_ids = [item.source_id for item in candidates]
    if len(source_ids) != len(set(source_ids)):
        raise ValueError("raw Base-102 source identity collision")
    return RawContext(
        candidates=tuple(candidates),
        enterprise_rows=enterprise_rows,
        control_rows=control_rows,
        enterprise_entity_ids=enterprise_entities,
        control_counts=control_counts,
    )


def map_human_workbook(
    workbook_path: Path,
    target: str,
    context: RawContext,
    canonical_scope_ids: Sequence[str] | None = None,
) -> WorkbookMapping:
    """Map one preserved human workbook to stable raw-source IDs.

    If ``canonical_scope_ids`` is supplied, the first 4,160/426 rows retain the
    already established expert scope order.  Additional B2 rows are mapped as
    controlled observations and do not expand the expert scope decision.
    """
    if target not in TARGET_COUNTS:
        raise ValueError(f"unsupported target: {target}")
    _, human_rows = _read_rows(workbook_path)
    expected_scope = TARGET_COUNTS[target]
    by_source = context.by_source_id
    exceptions: list[dict[str, Any]] = []
    mapped: list[MappedHumanRow] = []

    exact: dict[tuple[str, ...], deque[RawCandidate]] = defaultdict(deque)
    by_entity_year: dict[tuple[str, str], list[RawCandidate]] = defaultdict(list)
    for candidate in context.candidates:
        features = candidate.features(target)
        exact[_signature(features)].append(candidate)
        by_entity_year[(features["entity"], features["year"])].append(candidate)
    used: set[str] = set()
    canonical_order = tuple(canonical_scope_ids or ())

    for ordinal, (human_excel_row, values) in enumerate(human_rows):
        is_scope_row = ordinal < expected_scope
        features = _human_features(values, target)
        candidate: RawCandidate | None = None
        method = ""

        if canonical_scope_ids is not None and is_scope_row:
            if ordinal >= len(canonical_order):
                method = "canonical_scope_missing"
            else:
                candidate = by_source.get(canonical_order[ordinal])
                method = "canonical_expert_scope_order"
                if candidate is not None and candidate.pseudonymous_entity_id != features["entity"]:
                    candidate = None
                    method = "canonical_scope_entity_mismatch"
        else:
            queue = exact.get(_signature(features), deque())
            if canonical_scope_ids is None:
                while queue and queue[0].source_id in used:
                    queue.popleft()
            if queue:
                candidate = queue.popleft()
                method = "exact_semantic_signature_order"
            else:
                entity_pool = by_entity_year.get((features["entity"], features["year"]), ())
                available = entity_pool if canonical_scope_ids is not None else (
                    item for item in entity_pool if item.source_id not in used
                )
                candidate, method = _fallback_candidate(features, available, target)

        if candidate is None:
            local_id = f"UNMAPPED-HUMAN-{target}-{human_excel_row:06d}"
            exceptions.append({
                "source_id": "",
                "observation_id": local_id,
                "target": target,
                "human_workbook_row": human_excel_row,
                "pollutant": "",
                "category": "SOURCE_MAPPING_FAILED",
                "description": method or "no stable raw-source candidate",
            })
            mapped.append(MappedHumanRow(
                human_excel_row=human_excel_row,
                source_id="",
                raw_excel_row=0,
                mapping_method=method or "unmapped",
                is_scope_row=is_scope_row,
                row_values=values,
            ))
            continue

        used.add(candidate.source_id)
        mapped.append(MappedHumanRow(
            human_excel_row=human_excel_row,
            source_id=candidate.source_id,
            raw_excel_row=candidate.raw_excel_row,
            mapping_method=method,
            is_scope_row=is_scope_row,
            row_values=values,
        ))

    scope_rows = [row for row in mapped if row.is_scope_row and row.source_id]
    if len(human_rows) < expected_scope:
        exceptions.append({
            "source_id": "",
            "observation_id": "",
            "target": target,
            "human_workbook_row": "",
            "pollutant": "",
            "category": "HUMAN_SCOPE_ROW_COUNT_SHORT",
            "description": f"observed {len(human_rows)} rows, expected at least {expected_scope}",
        })
    if len({row.source_id for row in scope_rows}) != len(scope_rows):
        exceptions.append({
            "source_id": "",
            "observation_id": "",
            "target": target,
            "human_workbook_row": "",
            "pollutant": "",
            "category": "HUMAN_SCOPE_SOURCE_REUSE",
            "description": "two or more expert scope rows resolved to the same raw source",
        })
    if canonical_scope_ids is not None:
        scope_source_ids = tuple(row.source_id for row in scope_rows)
    else:
        scope_source_ids = tuple(row.source_id for row in scope_rows)
    return WorkbookMapping(
        workbook=workbook_path,
        target=target,
        rows=tuple(mapped),
        scope_source_ids=scope_source_ids,
        mapping_exceptions=tuple(exceptions),
    )


def _write_csv(path: Path, fieldnames: Sequence[str], rows: Iterable[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _decision_rows(mapping: WorkbookMapping, context: RawContext) -> list[dict[str, Any]]:
    included = {row.source_id: row for row in mapping.rows if row.is_scope_row and row.source_id}
    output = []
    for candidate in context.candidates:
        human = included.get(candidate.source_id)
        output.append({
            "source_id": candidate.source_id,
            "target": mapping.target,
            "human_scope_decision": "include" if human else "not_selected_for_target",
            "human_workbook_row": human.human_excel_row if human else "",
            "mapping_method": human.mapping_method if human else "absence_from_expert_target_workbook",
            "raw_base102_row": candidate.raw_excel_row,
            "raw_base102_sequence": candidate.sequence,
            "inventory_year": candidate.inventory_year,
            "pseudonymous_entity_id": candidate.pseudonymous_entity_id,
            "industry_code": candidate.industry_code,
            "target_equipment_value": (
                candidate.industrial_equipment if mapping.target == "INDUSTRIAL" else candidate.power_equipment
            ),
            "enterprise_master_present": candidate.pseudonymous_entity_id in context.enterprise_entity_ids,
            "enterprise_control_row_count": context.control_counts.get(candidate.pseudonymous_entity_id, 0),
        })
    return output


def _result_offsets(target: str) -> tuple[int, int]:
    # 0-based offsets in the preserved expert result workbooks.
    return (75, 91) if target == "INDUSTRIAL" else (78, 94)


def _calculation_and_exceptions(
    mapping: WorkbookMapping,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    generation_start, emission_start = _result_offsets(mapping.target)
    calculations: list[dict[str, Any]] = []
    exceptions: list[dict[str, Any]] = [dict(item) for item in mapping.mapping_exceptions]
    exception_number = len(exceptions)
    for row in mapping.rows:
        role = "expert_scope" if row.is_scope_row else "controlled_extra_observation"
        base_observation = row.source_id or f"UNMAPPED-HUMAN-{mapping.target}-{row.human_excel_row:06d}"
        observation_id = (
            base_observation
            if row.is_scope_row
            else f"{base_observation}::HUMAN-EXTRA-{row.human_excel_row:06d}"
        )
        values = _pad(row.row_values)
        for offset, pollutant in enumerate(POLLUTANTS):
            generation = _finite(values[generation_start + offset])
            emission = _finite(values[emission_start + offset])
            reported = generation is not None and emission is not None
            reason = "" if reported else "HUMAN_TOTAL_NOT_REPORTED"
            calculations.append({
                "observation_id": observation_id,
                "source_id": row.source_id,
                "target": mapping.target,
                "source_decision_role": role,
                "human_workbook_row": row.human_excel_row,
                "pollutant": pollutant,
                "status": "reported_total" if reported else "not_observable",
                "method": "expert_led_workbook",
                "generation_t": "" if generation is None else generation,
                "emission_t": "" if emission is None else emission,
                "activity_value": NA,
                "activity_unit": NA,
                "parameter_value": NA,
                "parameter_unit": NA,
                "control_method": NA,
                "control_efficiency": NA,
                "rule_id": NA,
                "standard_reference": NA,
                "intermediate_calculation": NA,
                "reason_code": reason,
            })
            if not reported:
                exception_number += 1
                exceptions.append({
                    "source_id": row.source_id,
                    "observation_id": observation_id,
                    "target": mapping.target,
                    "human_workbook_row": row.human_excel_row,
                    "pollutant": pollutant,
                    "category": "HUMAN_TOTAL_NOT_REPORTED",
                    "description": "generation or final emission total is not numerically observable in the supplied workbook",
                })
    for index, item in enumerate(exceptions, start=1):
        item["exception_id"] = f"HUMAN-V2-{mapping.target}-{index:06d}"
    return calculations, exceptions


def _target_for_workbook(path: Path) -> str:
    if "火电、热力" in path.name:
        return "POWER"
    if "工业锅炉" in path.name:
        return "INDUSTRIAL"
    raise ValueError(f"cannot infer target from preserved workbook: {path}")


def _baseline_workbook(workbooks: Sequence[Path], target: str) -> Path:
    matches = [
        path for path in workbooks
        if "02_实验A_端到端清单编制" in path.parts and _target_for_workbook(path) == target
    ]
    if len(matches) != 1:
        raise RuntimeError(f"expected one Experiment-A human workbook for {target}, found {len(matches)}")
    return matches[0]


def normalize_all(
    *,
    original_root: Path,
    output_root: Path,
    raw_102: Path,
    raw_101_enterprise: Path,
    raw_101_control: Path,
    replace: bool = False,
) -> dict[str, Any]:
    """Normalize all 14 packages without modifying the preserved originals."""
    original_root = original_root.resolve()
    output_root = output_root.resolve()
    if (
        output_root == original_root
        or output_root in original_root.parents
        or original_root in output_root.parents
    ):
        raise ValueError("output_root must be disjoint from the preserved original_packages tree")
    if output_root.exists():
        if not replace:
            raise FileExistsError(f"output already exists: {output_root}")
        shutil.rmtree(output_root)
    output_root.mkdir(parents=True)

    workbooks = sorted(original_root.rglob("*.xlsx"))
    if len(workbooks) != 14:
        raise RuntimeError(f"expected 14 preserved human workbooks, found {len(workbooks)}")
    before_hashes = {str(path.relative_to(original_root)): _sha256(path) for path in workbooks}
    context = load_raw_context(raw_102, raw_101_enterprise, raw_101_control)
    raw_inputs = {
        "base102": {"path": str(raw_102), "sha256": _sha256(raw_102)},
        "base101_enterprise": {"path": str(raw_101_enterprise), "sha256": _sha256(raw_101_enterprise)},
        "base101_control": {"path": str(raw_101_control), "sha256": _sha256(raw_101_control)},
    }

    canonical: dict[str, WorkbookMapping] = {}
    for target in TARGET_COUNTS:
        baseline = _baseline_workbook(workbooks, target)
        canonical[target] = map_human_workbook(baseline, target, context)
        if canonical[target].mapping_exceptions:
            raise RuntimeError(f"baseline mapping failed for {target}: {canonical[target].mapping_exceptions[:3]}")
        if len(canonical[target].scope_source_ids) != TARGET_COUNTS[target]:
            raise RuntimeError(f"baseline scope count mismatch for {target}")

    package_summaries = []
    for workbook in workbooks:
        target = _target_for_workbook(workbook)
        if workbook == canonical[target].workbook:
            mapping = canonical[target]
        else:
            mapping = map_human_workbook(
                workbook,
                target,
                context,
                canonical_scope_ids=canonical[target].scope_source_ids,
            )
        destination = output_root / workbook.relative_to(original_root).parent
        destination.mkdir(parents=True, exist_ok=True)
        decisions = _decision_rows(mapping, context)
        calculations, exceptions = _calculation_and_exceptions(mapping)
        _write_csv(destination / "source_decisions.csv", SOURCE_DECISION_FIELDS, decisions)
        _write_csv(destination / "calculation_totals.csv", CALCULATION_FIELDS, calculations)
        _write_csv(destination / "exceptions.csv", EXCEPTION_FIELDS, exceptions)

        included = sum(row["human_scope_decision"] == "include" for row in decisions)
        reported = sum(row["status"] == "reported_total" for row in calculations)
        notes = {
            "schema_version": "human-baseline-normalization-v2.0.0",
            "method": "既有专家主导型清单编制流程",
            "target": target,
            "source_workbook": str(workbook.relative_to(original_root)),
            "source_workbook_sha256": before_hashes[str(workbook.relative_to(original_root))],
            "raw_inputs": raw_inputs,
            "candidate_source_count": len(context.candidates),
            "human_scope_include_count": included,
            "human_scope_not_selected_count": len(context.candidates) - included,
            "controlled_extra_observation_count": sum(not row.is_scope_row for row in mapping.rows),
            "calculation_total_row_count": len(calculations),
            "reported_total_row_count": reported,
            "not_observable_total_row_count": len(calculations) - reported,
            "exception_count": len(exceptions),
            "process_record_policy": (
                "No reference or raw-table values are used to reconstruct human activity, parameter, control, "
                "rule, or intermediate-calculation records; unreported process fields are NA."
            ),
            "scope_policy": (
                "Presence in the preserved expert target workbook is an include decision; absence from that "
                "target workbook is not-selected-for-target. Controlled B2 extra rows do not expand the "
                "4,160/426 expert scope decisions."
            ),
            "timing": {
                "seconds_per_candidate_record": SECONDS_PER_RECORD,
                "candidate_record_count": len(context.candidates),
                "total_seconds": SECONDS_PER_RECORD * len(context.candidates),
                "basis": "extrapolated_from_user_supplied_per_candidate_rate",
            },
        }
        notes_path = destination / "normalization_notes.json"
        notes_path.write_text(json.dumps(notes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        package_summaries.append({
            "workbook": str(workbook.relative_to(original_root)),
            "output": str(destination.relative_to(output_root)),
            "target": target,
            "scope_included": included,
            "controlled_extra_observations": notes["controlled_extra_observation_count"],
            "calculation_rows": len(calculations),
            "not_observable_rows": notes["not_observable_total_row_count"],
            "exceptions": len(exceptions),
        })

    after_hashes = {str(path.relative_to(original_root)): _sha256(path) for path in workbooks}
    if after_hashes != before_hashes:
        raise RuntimeError("preserved human workbook bytes changed during normalization")
    summary = {
        "schema_version": "human-baseline-normalization-v2.0.0",
        "status": "ok",
        "package_count": len(package_summaries),
        "raw_candidate_count": len(context.candidates),
        "preserved_originals_unchanged": True,
        "original_workbook_hashes": before_hashes,
        "packages": package_summaries,
    }
    (output_root / "normalization_index.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--original-root", type=Path, default=DEFAULT_ORIGINAL_ROOT)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--raw-102", type=Path, default=DEFAULT_RAW_102)
    parser.add_argument("--raw-101-enterprise", type=Path, default=DEFAULT_RAW_101_ENTERPRISE)
    parser.add_argument("--raw-101-control", type=Path, default=DEFAULT_RAW_101_CONTROL)
    parser.add_argument("--replace", action="store_true")
    args = parser.parse_args()
    summary = normalize_all(
        original_root=args.original_root,
        output_root=args.output_root,
        raw_102=args.raw_102,
        raw_101_enterprise=args.raw_101_enterprise,
        raw_101_control=args.raw_101_control,
        replace=args.replace,
    )
    print(json.dumps({
        "status": summary["status"],
        "package_count": summary["package_count"],
        "raw_candidate_count": summary["raw_candidate_count"],
        "output_root": str(args.output_root.resolve()),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
